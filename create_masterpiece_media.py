import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pythainlp.tokenize import word_tokenize

target_dir = r"D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101\สื่อการสอน"
os.makedirs(target_dir, exist_ok=True)

ZWSP = '\u200B'

def insert_zwsp(text):
    if not text: return ""
    tokens = word_tokenize(text, engine="newmm")
    return ZWSP.join(tokens)

def set_font(run, font_name="TH Sarabun New", size=16):
    run.font.name = font_name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set('ascii', font_name)
    rFonts.set('hAnsi', font_name)
    rFonts.set('cs', font_name)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(insert_zwsp(text))
        set_font(run, size=20)
        run.bold = True
    else:
        run = p.add_run(insert_zwsp(text))
        set_font(run, size=18)
        run.bold = True

def add_paragraph(doc, text, bold=False, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(insert_zwsp(text))
    set_font(run, size=16)
    run.bold = bold

def add_table(doc, headers, rows_count):
    table = doc.add_table(rows=1 + rows_count, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(insert_zwsp(h))
        set_font(run, size=16)
        run.bold = True
    
    for row in table.rows[1:]:
        for cell in row.cells:
            p = cell.paragraphs[0]
            run = p.add_run(" ")
            set_font(run, size=16)
    return table

def create_doc(filename, title, content_builder):
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    
    add_heading(doc, title, level=1)
    doc.add_paragraph()
    
    content_builder(doc)
    
    doc.save(os.path.join(target_dir, filename))

# --- Builders ---

def build_p01_masterpiece(doc):
    add_heading(doc, "ภารกิจที่ 1: สแกนฟังเสียง Reference Pitch", level=2)
    add_paragraph(doc, "คำชี้แจง: ให้นักเรียนสแกน QR Code ด้านล่างเพื่อฟังเสียงความถี่มาตรฐาน C Major แล้วหลับตาเทียบกับเสียงคาลิมบ้าของตนเอง (ห้ามมองหน้าจอแอปพลิเคชัน)")
    
    add_table(doc, ["[   พื้นที่สำหรับแปะ QR Code เสียงมาตรฐาน (Reference Pitch)   ]"], 1)
    doc.add_paragraph()
    
    add_heading(doc, "ภารกิจที่ 2: Blind Tuning Challenge (จับผิดเสียงเพี้ยน)", level=2)
    add_paragraph(doc, "คำชี้แจง: ครูผู้สอนได้ทำการตั้งสายคาลิมบ้าของนักเรียนเพี้ยนไว้จำนวน 2 ลิ่ม ให้นักเรียนหาให้เจอโดยใช้ 'โสตทักษะ (หูฟัง)' เท่านั้น แล้วบันทึกผลลงในตาราง (ส่งคลิปเสียงขณะจูน 15 วินาทีเข้าระบบ)")
    add_table(doc, ["ลำดับ", "ลิ่มเสียงที่คาดว่าเพี้ยน (คีย์)", "ทิศทางความเพี้ยน (สูงไป / ต่ำไป)", "ผลการเฉลยจากครู (ถูกต้อง / ไม่ถูกต้อง)"], 2)

def build_p03_masterpiece(doc):
    add_paragraph(doc, "คำชี้แจง: เปลี่ยนจากการขอลายเซ็น เป็นการพิสูจน์ความแม่นยำด้วยหลักฐานเชิงประจักษ์ (Empirical Evidence)")
    doc.add_paragraph()
    add_heading(doc, "กติกาการทดสอบ (Blind Dictation)", level=2)
    add_paragraph(doc, "1. ผู้เล่น (Player) บรรเลงแพทเทิร์นจังหวะและสเกลที่จับฉลากได้")
    add_paragraph(doc, "2. ผู้ประเมิน (Peer) ต้องหันหลัง หลับตาฟัง และเขียนโน้ตจังหวะ (Rhythm Dictation) สิ่งที่ได้ยินลงในตารางให้ถูกต้อง")
    add_paragraph(doc, "3. นำผลการจดบันทึกไปเทียบเคียงกับต้นฉบับ หากตรงกัน 100% จึงจะถือว่า 'สอบผ่าน'")
    doc.add_paragraph()
    
    add_table(doc, ["รอบการบรรเลง", "จังหวะที่ผู้ประเมินได้ยินและเขียนได้ (Rhythm Dictation)", "ผลการเทียบเคียงต้นฉบับ (ตรง / ไม่ตรง)", "ลายเซ็นพยานยืนยัน (ครูผู้สอน)"], 4)

def build_p14_masterpiece(doc):
    add_paragraph(doc, "คำชี้แจง: โครงร่างนี้บังคับให้เกิดกระบวนการคิดวิเคราะห์เปรียบเทียบข้ามวัฒนธรรมเชิงลึก โดยไม่อนุญาตให้คัดลอกเนื้อหาจากอินเทอร์เน็ตมาวางโดยตรง")
    doc.add_paragraph()
    add_heading(doc, "Cross-Cultural Script Matrix (โครงสร้างบังคับความคิด)", level=2)
    add_table(doc, ["ส่วนประกอบ", "Timestamp ของเพลง (ระบุวินาทีที่จะเปิดเทียบ)", "มิติทางดนตรีที่ต่างกัน (เช่น จังหวะ, บันไดเสียง, โทนสีเสียง)", "เนื้อหาสคริปต์: อธิบายความแตกต่างเชิงสุนทรียภาพทางดนตรี", "Sound Effect / BGM Cue"], 5)

# --- Generation ---

files = [
    ("สื่อ_น1_ผ01_คู่มือจูนเสียงและดูแลรักษาคาลิมบ้า_Masterpiece.docx", "ใบกิจกรรม: Ear Training & Tuning Challenge (Masterpiece)", build_p01_masterpiece),
    ("สื่อ_น1_ผ03_Challengeไล่สเกลและจังหวะ_Masterpiece.docx", "ใบกิจกรรม: Interactive Rhythm Dictation Challenge (Masterpiece)", build_p03_masterpiece),
    ("สื่อ_น3_ผ14_โครงร่างสคริปต์Podcastข้ามวัฒนธรรม_Masterpiece.docx", "ใบกิจกรรม: Cross-Cultural Script Matrix (Masterpiece)", build_p14_masterpiece)
]

for filename, title, builder in files:
    create_doc(filename, title, builder)
