import os
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

target_dir = r"D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101\สื่อการสอน"
files = sorted([f for f in os.listdir(target_dir) if f.endswith('.docx')])

for fname in files:
    path = os.path.join(target_dir, fname)
    doc = Document(path)
    print(f"=== {fname} ===")
    for para in doc.paragraphs:
        if para.text.strip():
            print(" P:", para.text.strip()[:250])
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip()[:80] for c in row.cells)
            if row_text.strip(" |"):
                print(" T:", row_text)
    print()
