import docx
from docx.shared import Pt
from docx.oxml.ns import qn

def set_font(run, size=16, bold=False):
    run.font.name = 'TH Sarabun New'
    run.font.size = Pt(size)
    run.bold = bold
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
    r.rPr.rFonts.set(qn('w:cs'), 'TH Sarabun New')

def update_wk3():
    path = 'เอกสารแผนการสอน/ม.1 ดนตรี1 ศ21101/วก3/วก3_โครงสร้างรายวิชา_ศ21101_ดนตรี1.docx'
    doc = docx.Document(path)
    
    # 1. Update Paragraph 3 (Time & Credit)
    for p in doc.paragraphs:
        if 'ชั้นมัธยมศึกษาปีที่ 1' in p.text and 'เวลา 20' in p.text:
            p.text = 'ชั้นมัธยมศึกษาปีที่ 1  ภาคเรียนที่ 1  เวลา 20 ชั่วโมง (20 คาบเรียน รวม 1,000 นาที)  จำนวน 0.5 หน่วยกิต'
            for r in p.runs:
                set_font(r, 16, bold=False)
            break

    # 2. Update Paragraph 6 (Course description: Add โน้ตไทย)
    for p in doc.paragraphs:
        if 'ศึกษาเกี่ยวกับการอ่าน' in p.text:
            p.text = '    ศึกษาเกี่ยวกับการอ่าน เขียน และร้องโน้ตไทยและโน้ตสากลเบื้องต้น การเปรียบเทียบเสียงร้องและเสียงของเครื่องดนตรีต่างวัฒนธรรม การวิเคราะห์จัดประเภทของวงดนตรีไทยและสากล อธิบายบทบาทความสัมพันธ์และอิทธิพลของดนตรีที่มีต่อสังคมไทย รวมถึงระบุความหลากหลายขององค์ประกอบดนตรีในวัฒนธรรมต่างกัน'
            for r in p.runs:
                set_font(r, 16, bold=False)
            break

    # 3. Update Paragraph 7 (Active Learning & T-PACE description)
    for p in doc.paragraphs:
        if 'โดยใช้การออกแบบการจัดการเรียนรู้แบบอิงมาตรฐาน' in p.text or 'โดยใช้กระบวนการจัดการเรียนรู้เชิงรุก' in p.text:
            p.text = '    โดยใช้กระบวนการจัดการเรียนรู้เชิงรุก (Active Learning) ตามกรอบแนวคิดนวัตกรรม T-PACE Framework (Technology–Performance–Active–Creative–Evaluation) ที่ขับเคลื่อนด้วยการออกแบบแบบอิงมาตรฐาน (Backward Design) ควบคู่กับระบบนิเวศการเรียนรู้ดิจิทัล (Digital Learning Ecosystem) ผ่านแอปพลิเคชัน MusicSmart ที่สนับสนุนการปรับเทียบเสียง (Tuner) การประเมินเพื่อนร่วมชั้นแบบทันที (Real-time Peer Review) และการบันทึกข้อมูลเชิงประจักษ์เพื่อการวิจัยในชั้นเรียน (Classroom Action Research: CAR)'
            for r in p.runs:
                set_font(r, 16, bold=False)
            break

    # 4. Clean up duplicate indicators
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == 'ศ 2.2 ม.1/1, ม.1/2' and i > 11:
            p.text = ''
        if p.text.strip() == 'รวมทั้งหมด 11 ตัวชี้วัด' and i > 11:
            p.text = ''

    # 5. Update P30 & P31 (Score ratio & K:P:A breakdown)
    for p in doc.paragraphs:
        if 'อัตราส่วนคะแนน K:' in p.text:
            p.text = ('อัตราส่วนคะแนน K: 30  P: 50  A: 20  จำแนกเป็น:\n'
                      '  • ด้านความรู้ (K) 30 คะแนน: คะแนนเก็บในหน่วยการเรียนรู้ 26 คะแนน + แบบทดสอบวัดผลกลางภาค 2 คะแนน + แบบทดสอบวัดผลปลายภาค 2 คะแนน\n'
                      '  • ด้านทักษะกระบวนการ (P) 50 คะแนน: คะแนนเก็บปฏิบัติในหน่วยการเรียนรู้ 34 คะแนน + สอบปฏิบัติเดี่ยวกลางภาค 8 คะแนน + สอบปฏิบัติรวมวง/สุนทรียภาพปลายภาค 8 คะแนน\n'
                      '  • ด้านคุณลักษณะอันพึงประสงค์ (A) 20 คะแนน: การสังเกตพฤติกรรมการเรียนรู้ จิตพิสัย วินัย และความรับผิดชอบในการดูแลรักษาเครื่องดนตรีตลอดภาคเรียน 20 คะแนน\n'
                      '  (รวมทั้งสิ้น 100 คะแนนเต็ม สอดคล้องตามมาตรฐานดนตรีศึกษาที่มุ่งเน้นสมรรถนะทักษะการปฏิบัติร้อยละ 50)')
            for r in p.runs:
                set_font(r, 16, bold=False)
            break

    # 6. Update Table 2 (Course Structure Table: Midterm & Final breakdown)
    t2 = doc.tables[2]
    for r in t2.rows:
        if 'สอบกลางภาค' in r.cells[1].text:
            r.cells[2].text = 'ทดสอบวัดผลกลางภาค (หน่วย 1-2): ข้อสอบปรนัยวัดความรู้ (K: 2) และสอบทักษะปฏิบัติเดี่ยวคาลิมบ้า (P: 8)'
            r.cells[5].text = '10\n(K: 2, P: 8)'
            for c in r.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        set_font(run, 14)
        elif 'สอบปลายภาค' in r.cells[1].text:
            r.cells[2].text = 'ทดสอบวัดผลปลายภาค (หน่วย 3-4): ข้อสอบปรนัยวัดความรู้/การวิจารณ์ (K: 2) และสอบปฏิบัติรวมวง/นำเสนอผลงาน (P: 8)'
            r.cells[5].text = '10\n(K: 2, P: 8)'
            for c in r.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        set_font(run, 14)
        elif 'รวม' in r.cells[1].text:
            r.cells[5].text = '100\n(K: 30, P: 50, A: 20)'
            for c in r.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        set_font(run, 14, bold=True)

    # 7. Update Table 3 (Scoring Table: Clarify K/P in Midterm and Final)
    t3 = doc.tables[3]
    # Header row 1:
    t3.rows[1].cells[5].text = 'กลางภาค\n(K:2 / P:8)'
    t3.rows[1].cells[6].text = 'ปลายภาค\n(K:2 / P:8)'
    t3.rows[1].cells[7].text = 'รวมสุทธิ\n(K:30/P:50/A:20)'

    # Unit rows:
    t3.rows[2].cells[5].text = '5\n(K:1, P:4)'
    t3.rows[2].cells[6].text = '-'
    t3.rows[2].cells[7].text = '25'

    t3.rows[3].cells[5].text = '5\n(K:1, P:4)'
    t3.rows[3].cells[6].text = '-'
    t3.rows[3].cells[7].text = '25'

    t3.rows[4].cells[5].text = '-'
    t3.rows[4].cells[6].text = '5\n(K:1, P:4)'
    t3.rows[4].cells[7].text = '25'

    t3.rows[5].cells[5].text = '-'
    t3.rows[5].cells[6].text = '5\n(K:1, P:4)'
    t3.rows[5].cells[7].text = '25'

    # Total row:
    t3.rows[6].cells[1].text = '26'
    t3.rows[6].cells[2].text = '34'
    t3.rows[6].cells[3].text = '20'
    t3.rows[6].cells[4].text = '80'
    t3.rows[6].cells[5].text = '10\n(K:2, P:8)'
    t3.rows[6].cells[6].text = '10\n(K:2, P:8)'
    t3.rows[6].cells[7].text = '100\n(K:30, P:50, A:20)'

    for r in t3.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    set_font(run, 13)

    # 8. Add 4-tier Approval Block
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    if 'การตรวจพิจารณาและอนุมัติโครงสร้างรายวิชา' not in full_text:
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(16)
        p_head.paragraph_format.space_after = Pt(6)
        r = p_head.add_run('การตรวจพิจารณาและอนุมัติโครงสร้างรายวิชา')
        set_font(r, 16, bold=True)

        items = [
            ("ลงชื่อ ....................................................... ผู้จัดทำ / ครูผู้สอน\n(นายพงศธร โพธิแก้ว)\nตำแหน่ง ครู", 15),
            ("1. ความเห็นของหัวหน้ากลุ่มสาระการเรียนรู้ศิลปะ: .................................................................................\nลงชื่อ ....................................................... หัวหน้ากลุ่มสาระการเรียนรู้ศิลปะ\n(นายอุดม ไชยวุฒิ)", 14),
            ("2. ความเห็นของรองผู้อำนวยการกลุ่มบริหารงานวิชาการ: ............................................................................\nลงชื่อ ....................................................... รองผู้อำนวยการกลุ่มบริหารงานวิชาการ\n(นายกฤศ วงค์เรือง)", 14),
            ("3. การพิจารณาอนุมัติของผู้อำนวยการสถานศึกษา: ...................................................................................\nลงชื่อ ....................................................... ผู้อำนวยการโรงเรียนพะเยาพิทยาคม\n(ว่าที่ร้อยตรีญาณบดินทร์ อินเตชะ)", 14)
        ]
        for text, size in items:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_font(r, size)

    doc.save(path)
    print("WK3 updated successfully!")

if __name__ == '__main__':
    update_wk3()
