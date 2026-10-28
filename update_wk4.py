import os
import glob
import re
import docx
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

def set_font(run, size=16, bold=False):
    run.font.name = 'TH Sarabun New'
    run.font.size = Pt(size)
    run.bold = bold
    r = run._element
    if r.rPr is None:
        r.get_or_add_rPr()
    r.rPr.rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
    r.rPr.rFonts.set(qn('w:cs'), 'TH Sarabun New')

def update_wk4_doc(filepath):
    doc = docx.Document(filepath)
    modified = False

    # 1. Update CAR & PDPA placeholders if present
    for p in doc.paragraphs:
        if 'ชื่อผู้วิจัย: นาย/นาง/นางสาว' in p.text:
            p.text = p.text.replace(
                'ชื่อผู้วิจัย: นาย/นาง/นางสาว _________________________________________',
                'ชื่อผู้วิจัย: นายพงศธร โพธิแก้ว'
            )
            for r in p.runs:
                set_font(r, 16)
            modified = True

        if 'ลงชื่อ __________________________________________________ ผู้วิจัย' in p.text:
            p.text = 'ลงชื่อ ....................................................... ผู้วิจัย'
            for r in p.runs:
                set_font(r, 15)
            modified = True

        if p.text.strip() == '( __________________________________________________ )':
            p.text = '     (นายพงศธร โพธิแก้ว)'
            for r in p.runs:
                set_font(r, 15)
            modified = True

        if 'ด้วย ข้าพเจ้า นาย/นาง/นางสาว' in p.text:
            p.text = p.text.replace(
                'ด้วย ข้าพเจ้า นาย/นาง/นางสาว _____________________________________',
                'ด้วย ข้าพเจ้า นายพงศธร โพธิแก้ว'
            )
            for r in p.runs:
                set_font(r, 16)
            modified = True

    # 2. Check if 4-tier approval block exists before ภาคผนวก
    full_text = '\n'.join([p.text for p in doc.paragraphs])
    if 'การตรวจพิจารณาและอนุมัติหน่วยการเรียนรู้' not in full_text:
        target_p = None
        for p in doc.paragraphs:
            if p.text.strip() == 'ภาคผนวก':
                target_p = p
                break
        
        blocks = [
            ("การตรวจพิจารณาและอนุมัติหน่วยการเรียนรู้", 16, True),
            ("ลงชื่อ ....................................................... ผู้จัดทำ / ครูผู้สอน\n(นายพงศธร โพธิแก้ว)\nตำแหน่ง ครู", 15, False),
            ("1. ความเห็นของหัวหน้ากลุ่มสาระการเรียนรู้ศิลปะ: .................................................................................\nลงชื่อ ....................................................... หัวหน้ากลุ่มสาระการเรียนรู้ศิลปะ\n(นายอุดม ไชยวุฒิ)", 14, False),
            ("2. ความเห็นของรองผู้อำนวยการกลุ่มบริหารงานวิชาการ: ............................................................................\nลงชื่อ ....................................................... รองผู้อำนวยการกลุ่มบริหารงานวิชาการ\n(นายกฤศ วงค์เรือง)", 14, False),
            ("3. การพิจารณาอนุมัติของผู้อำนวยการสถานศึกษา: ...................................................................................\nลงชื่อ ....................................................... ผู้อำนวยการโรงเรียนพะเยาพิทยาคม\n(ว่าที่ร้อยตรีญาณบดินทร์ อินเตชะ)", 14, False),
        ]
        
        for text, size, is_bold in blocks:
            new_p = doc.add_paragraph()
            new_p.paragraph_format.space_before = Pt(8 if is_bold else 4)
            new_p.paragraph_format.space_after = Pt(6)
            new_p.paragraph_format.line_spacing = 1.15
            r = new_p.add_run(text)
            set_font(r, size, bold=is_bold)
            if target_p:
                target_p._element.addprevious(new_p._element)
            
        modified = True

    if modified:
        doc.save(filepath)
    return modified

def process_all_wk4():
    base = 'เอกสารแผนการสอน/ม.1 ดนตรี1 ศ21101'
    dirs = [
        os.path.join(base, 'วก4'),
        os.path.join(base, 'คู่มือการสอนสำหรับครู/คู่มือครู ม.1 ดนตรี1 ศ21101/วก4')
    ]
    total = 0
    updated = 0
    for d in dirs:
        if not os.path.exists(d):
            continue
        files = sorted(glob.glob(os.path.join(d, '*.docx')))
        files = [f for f in files if not os.path.basename(f).startswith('~$')]
        print(f"\nProcessing WK4 in {d} ({len(files)} files)")
        for f in files:
            total += 1
            if update_wk4_doc(f):
                updated += 1
                print(f"  ✓ Updated {os.path.basename(f)}")
            else:
                print(f"  - No changes needed: {os.path.basename(f)}")
    print(f"\nCompleted WK4: {updated}/{total} files updated.")

if __name__ == '__main__':
    process_all_wk4()
