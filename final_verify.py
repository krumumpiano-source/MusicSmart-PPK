from docx import Document
import glob

def clean(t):
    for z in [chr(0x200B),chr(0x200C),chr(0x200D)]: t=t.replace(z,'')
    return t.strip()

def read_full(path):
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        t = clean(p.text)
        if t: parts.append(t)
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [clean(c.text) for c in row.cells if clean(c.text)]
            if cells: parts.append(' | '.join(cells))
    return '\n'.join(parts)

print('='*65)
print('  FINAL 3-TIER ALIGNMENT VERIFICATION REPORT')
print('='*65)

# WK3 check
wk3_text = read_full('วก3_คศ4_InventAndTransform_โครงสร้างรายวิชา.docx')
checks_wk3 = {
    'MusicSmart': 'MusicSmart' in wk3_text,
    'Backward Design': 'Backward Design' in wk3_text,
    'Digital Portfolio': 'Digital Portfolio' in wk3_text,
    'Mini-Recital': 'Mini-Recital' in wk3_text,
    'Podcast': 'Podcast' in wk3_text,
    'Playlist': 'Playlist' in wk3_text,
    'T-PACE': 'T-PACE' in wk3_text,
    'CAR/Action Research': 'Action Research' in wk3_text,
    'Constructivism': 'Constructivism' in wk3_text,
    'Self-Directed Learning': 'Self-Directed' in wk3_text,
}
wk3_pass = sum(checks_wk3.values())
print(f'\n[WK3] วก3_คศ4_InventAndTransform: {wk3_pass}/{len(checks_wk3)} checks')
for k,v in checks_wk3.items():
    icon = 'OK' if v else 'XX'
    print(f'  [{icon}] {k}')

# WK4 check
wk4_files = sorted(glob.glob('วก4_คศ4_InventAndTransform_*.docx'))
print(f'\n[WK4] Upgraded files ({len(wk4_files)} units):')
for path in wk4_files:
    text = read_full(path)
    has_eq = 'Essential Question' in text
    has_c = 'Constructivism' in text or 'Construct Knowledge' in text
    has_sdl = 'Self-Directed' in text
    cap = any(x in text for x in ['Digital Portfolio','Mini-Recital','Podcast','Playlist'])
    name = path.split('_')[-1].replace('.docx','')
    status = 'PASS' if (has_eq and has_sdl and cap) else 'WARN'
    print(f'  [{status}] {name}: EQ={has_eq} SDL={has_sdl} Capstone={cap}')

# WK5 final alignment
wk5_files = sorted(glob.glob('วก5_คศ4_InventAndTransform_แผนที่_*.docx'))
plan_unit = {}
for p in range(1,5): plan_unit[p]=1
for p in range(5,12): plan_unit[p]=2
for p in range(12,16): plan_unit[p]=3
for p in range(16,19): plan_unit[p]=4
must_have = {1:['ศ 2.1','Construct','Goal'],2:['ศ 2.1','Mini-Recital','Peer'],3:['ศ 2.2','Podcast'],4:['ศ 2.1','Playlist']}
print(f'\n[WK5] Final Alignment ({len(wk5_files)} plans):')
issues = 0
for i, path in enumerate(wk5_files):
    pnum = i+1
    text = read_full(path)
    unit = plan_unit.get(pnum, 0)
    must = must_have.get(unit, [])
    missing = [m for m in must if m not in text]
    if missing:
        issues += 1
        print(f'  [WARN] Plan {pnum:02d} (Unit {unit}): missing {missing}')
if issues == 0:
    print('  [OK] All 18 plans PASS alignment - zero issues!')

print(f'\n{"="*65}')
print(f'  SUMMARY: WK3={wk3_pass}/{len(checks_wk3)} | WK4={"OK" if len(wk4_files)==4 else "WARN"} | WK5 Issues={issues}/18')
if wk3_pass == len(checks_wk3) and issues == 0:
    print('  RESULT: 3-TIER ALIGNMENT COMPLETE - READY FOR KOR SOR 4 SUBMISSION!')
else:
    print(f'  RESULT: Partial - WK3 gaps={len(checks_wk3)-wk3_pass}, WK5 issues={issues}')
print('='*65)
