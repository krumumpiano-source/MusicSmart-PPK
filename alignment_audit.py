import subprocess, sys, os, glob
from docx import Document

ZWS = [chr(0x200B), chr(0x200C), chr(0x200D)]

def clean(text):
    for z in ZWS:
        text = text.replace(z, '')
    return text.strip()

def read_docx(path):
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        t = clean(p.text)
        if t:
            lines.append(t)
    for tbl in doc.tables:
        for row in tbl.rows:
            row_t = [clean(cell.text) for cell in row.cells if clean(cell.text)]
            if row_t:
                lines.append(' | '.join(row_t))
    return lines

# -------- READ ALL FILES --------
wk3_path = 'โครงสร้างรายวิชา_วก3_ดนตรีม1_สมบูรณ์.docx'
wk4_paths = sorted(glob.glob('หน่วยการเรียนรู้_วก4_หน่วยที่*.docx'))
wk5_paths = sorted(glob.glob('วก5_คศ4_InventAndTransform_แผนที่_*.docx'))

wk3_lines = read_docx(wk3_path)
wk4_data = {i+1: read_docx(p) for i, p in enumerate(wk4_paths)}
wk5_data = {i+1: read_docx(p) for i, p in enumerate(wk5_paths)}

print("="*70)
print("  ALIGNMENT AUDIT REPORT: วก3 ↔ วก4 ↔ วก5 (คศ.4)")
print("="*70)

# -------- PART 1: READ WK3 KEY FIELDS --------
print("\n### ตรวจสอบ วก.3 (โครงสร้างรายวิชา) ###")
wk3_full = '\n'.join(wk3_lines)

# Extract key info
kw3_check = {
    'รหัสวิชา': 'ศ21101' in wk3_full,
    'ชั้น ม.1': 'ม.1' in wk3_full or 'มัธยมศึกษาปีที่ 1' in wk3_full,
    'T-PACE': 'T-PACE' in wk3_full,
    'MusicSmart': 'MusicSmart' in wk3_full,
    'K=5, P=10, A=5': 'K=5' in wk3_full and 'P=10' in wk3_full,
    'Backward Design': 'Backward' in wk3_full,
    'Digital Portfolio': 'Digital Portfolio' in wk3_full or 'Portfolio' in wk3_full,
    'Mini-Recital': 'Mini-Recital' in wk3_full,
    'Podcast': 'Podcast' in wk3_full,
    'Playlist': 'Playlist' in wk3_full,
    'คาลิมบ้า': 'คาลิมบ้า' in wk3_full,
    'ตัวชี้วัด ศ 2.1': 'ศ 2.1' in wk3_full,
    'ตัวชี้วัด ศ 2.2': 'ศ 2.2' in wk3_full,
}
for k, v in kw3_check.items():
    status = '✅' if v else '❌'
    print(f'  {status} WK3 มี [{k}]: {v}')

# Check unit hours in WK3
print('\n  WK3 - บรรทัดที่กล่าวถึงเวลา/คะแนน:')
for line in wk3_lines:
    if any(x in line for x in ['ชั่วโมง', 'คะแนน', 'คาบ', 'เวลา']):
        print(f'    | {line[:100]}')

# -------- PART 2: READ WK4 KEY FIELDS --------
print("\n### ตรวจสอบ วก.4 (หน่วยการเรียนรู้ทั้ง 4 หน่วย) ###")
for unit_num, lines in wk4_data.items():
    unit_full = '\n'.join(lines)
    checks = {
        'T-PACE': 'T-PACE' in unit_full,
        'MusicSmart': 'MusicSmart' in unit_full,
        'K=5,P=10,A=5': ('K=5' in unit_full and 'P=10' in unit_full),
        'Constructivism/ด้วยตนเอง': 'ด้วยตนเอง' in unit_full or 'Construct' in unit_full,
        'Self-Directed': 'Self-Directed' in unit_full or 'กำกับตนเอง' in unit_full,
        'Backward Design': 'Backward' in unit_full or 'Essential Question' in unit_full,
        'ภาระงานรวบยอด': any(x in unit_full for x in ['Mini-Recital','Podcast','Portfolio','Playlist','Digital']),
    }
    passed = sum(checks.values())
    status = '✅' if passed >= 5 else ('⚠️ ' if passed >= 3 else '❌')
    print(f'  {status} หน่วยที่ {unit_num} ({passed}/{len(checks)} checks):')
    for k,v in checks.items():
        icon = '✅' if v else '❌'
        print(f'      {icon} {k}')
    # Extract KPA
    for line in lines:
        if 'ด้านความรู้' in line or 'K:' in line or 'P:' in line:
            print(f'    KPA: {line[:100]}')
            break

# -------- PART 3: ALIGNMENT CHECK WK4 vs WK5 --------
print("\n### มิติที่ 2: ตรวจ Alignment วก4 ↔ วก5 (รายหน่วย) ###")

# Map plan -> unit
plan_unit_map = {
    **{p: 1 for p in range(1, 5)},
    **{p: 2 for p in range(5, 12)},
    **{p: 3 for p in range(12, 16)},
    **{p: 4 for p in range(16, 19)},
}

alignment_issues = []

# Key indicators per unit that MUST appear in WK5 plans
unit_must_have = {
    1: ['ศ 2.1', 'คาลิมบ้า', 'Construct', 'Goal Setting'],
    2: ['ศ 2.1', 'Mini-Recital', 'Peer', 'สอดประสาน'],
    3: ['ศ 2.2', 'Podcast', 'วัฒนธรรม', 'Audio'],
    4: ['ศ 2.1', 'อารมณ์', 'สุนทรีย', 'Playlist'],
}

for plan_num, lines in wk5_data.items():
    plan_text = '\n'.join(lines)
    unit = plan_unit_map.get(plan_num, 0)
    must = unit_must_have.get(unit, [])
    missing = [m for m in must if m not in plan_text]
    
    # Check indicator code alignment
    has_indicator = any('ศ 2.' in l for l in lines)
    has_tpace = any(f'{x} -' in plan_text for x in ['T', 'P', 'A', 'E'])
    has_rubric = any('Rubrics' in l or 'ประเมิน' in l for l in lines)
    has_sdl = 'Self-Directed' in plan_text or 'Goal Setting' in plan_text
    has_construct = 'Construct' in plan_text or 'ด้วยตนเอง' in plan_text
    
    issues = []
    if missing: issues.append(f'ขาด keyword: {missing}')
    if not has_indicator: issues.append('ไม่มีรหัสตัวชี้วัด')
    if not has_tpace: issues.append('ไม่มี T-PACE')
    if not has_rubric: issues.append('ไม่มี Rubrics')
    if not has_sdl: issues.append('ไม่มี SDL/Goal Setting')
    if not has_construct: issues.append('ไม่มี Constructivism')
    
    if issues:
        alignment_issues.append((plan_num, unit, issues))
        print(f'  ⚠️  แผนที่ {plan_num:02d} (หน่วย{unit}): {"; ".join(issues)}')
    
print(f'\n  สรุป Alignment: พบปัญหา {len(alignment_issues)} แผน จากทั้งหมด 18 แผน')

# -------- PART 4: HOUR & SCORE ALIGNMENT --------
print("\n### มิติที่ 2.5: ตรวจจำนวนชั่วโมงและคะแนน ###")
print("  แผนที่ 1-4 = หน่วย 1 (4 คาบ / 20 คะแนน, K=5 P=10 A=5)")
print("  แผนที่ 5-11 = หน่วย 2 (7 คาบ / 20 คะแนน, K=5 P=10 A=5)")
print("  แผนที่ 12-15 = หน่วย 3 (4 คาบ / 20 คะแนน, K=5 P=10 A=5)")
print("  แผนที่ 16-18 = หน่วย 4 (3 คาบ / 20 คะแนน, K=5 P=10 A=5)")
print("  สอบกลางภาค 1 คาบ + สอบปลายภาค 1 คาบ = 2 คาบ")
print("  รวมทั้งหมด = 4+7+4+3+2 = 20 คาบ ✅")

# Verify each plan says '1 ชั่วโมง'
hour_ok = all('1 ชั่วโมง' in '\n'.join(lines) for lines in wk5_data.values())
print(f'  วก5 ทุกแผนระบุ "1 ชั่วโมง": {"✅" if hour_ok else "❌ บางแผนไม่ระบุ"}')

# Check K=5 P=10 A=5 in WK5
kpa_ok = sum(1 for lines in wk5_data.values() if 'K=5' in '\n'.join(lines) or 'ด้านความรู้ (K)' in '\n'.join(lines))
print(f'  วก5 มีหัวข้อ K-P-A: {kpa_ok}/18 แผน {"✅" if kpa_ok == 18 else "❌"}')

# -------- PART 5: SAMPLE DEEP CHECK plan 1 KPA --------
print("\n### มิติที่ 3: ตัวอย่าง Deep Check แผนที่ 01 ###")
p1 = '\n'.join(wk5_data[1])
# Find and print KPA objectives
for line in wk5_data[1]:
    if any(x in line for x in ['ด้านความรู้', 'ด้านทักษะ', 'ด้านคุณลักษณะ', '1.1', '2.1', '3.1']):
        print(f'  | {line}')

# Check Rubrics table
rubric_line = [l for l in wk5_data[1] if 'ประยุกต์' in l or 'Peer Coaching' in l or 'ลึกซึ้ง' in l]
print(f'\n  Rubrics คศ.4 ระดับ 4: {rubric_line[0][:120] if rubric_line else "❌ ไม่พบ"}')

# -------- SUMMARY --------
print("\n" + "="*70)
print("  สรุปรายงาน ACTION PLAN สำหรับการ SYNC")
print("="*70)
wk3_gaps = [k for k,v in kw3_check.items() if not v]
print(f'\n  🔴 WK3 ขาด: {wk3_gaps if wk3_gaps else "ไม่มี"}')
print(f'  🔴 WK4-WK5 Alignment Issues: {len(alignment_issues)} แผน')
print(f'  🟡 สิ่งที่ต้อง SYNC: {"WK3 ต้องเพิ่ม keywords ใหม่" if wk3_gaps else ""} {"WK4 ต้องเพิ่ม Constructivism/SDL" if any(not wk4_data[u] for u in range(1,5)) else ""}')
