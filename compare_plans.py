import os
import docx

root = r"D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101"
p1 = os.path.join(root, "วก5", "วก5_แผนการจัดการเรียนรู้_แผนที่_01.docx")
p2 = os.path.join(root, "แผนการสอนพร้อมใช้", "แผนการจัดการเรียนรู้_แผนที่_01.docx")

d1 = docx.Document(p1)
d2 = docx.Document(p2)

lines = []
lines.append(f"วก5/แผนที่ 01: paras = {len(d1.paragraphs)}, tables = {len(d1.tables)}")
lines.append(f"แผนการสอนพร้อมใช้/แผนที่ 01: paras = {len(d2.paragraphs)}, tables = {len(d2.tables)}")

lines.append("\n[วก5 header]")
for p in d1.paragraphs[:6]:
    lines.append(f"  {p.text.strip()}")

lines.append("\n[แผนการสอนพร้อมใช้ header]")
for p in d2.paragraphs[:6]:
    lines.append(f"  {p.text.strip()}")

with open("scratch_compare.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print("Compared.")
