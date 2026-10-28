# -*- coding: utf-8 -*-
"""
ซิงค์ข้อมูลหลักสูตรและเอกสารวิชาการทั้ง 79 ไฟล์ของ ม.2 (ศ22102 ดนตรี 2)
เข้าสู่โปรแกรมเว็บแอป MusicSmart (app/data/academic_docs.js และ app/data/lessons.js)
พร้อมคงข้อมูล ม.1 (79 ไฟล์) ไว้ครบถ้วน 100%
"""
import os
import re
import json
import sys
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding='utf-8')

ROOT_M2 = r"D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม. ดนตรี ศ22102"

def clean_txt(s):
    if not s:
        return ""
    s = s.replace("\u200b", "").replace("\ufeff", "")
    return re.sub(r"\s+", " ", s).strip()

def parse_docx(filepath, stop_at_appendix=False):
    doc = docx.Document(filepath)
    paragraphs = []
    tables_data = []

    if stop_at_appendix:
        for child in doc.element.body:
            if child.tag.endswith('p'):
                p = Paragraph(child, doc)
                txt = p.text.strip()
                if txt:
                    if txt.startswith("ภาคผนวก"):
                        paragraphs.append("📎 หมายเหตุ: เอกสารฉบับเต็มในไฟล์ Word (.docx) ได้ผนวกรวมภาคผนวก (ใบกิจกรรม, เครื่องมือวัดผลรูบริก และแบบทดสอบ) ไว้ท้ายเล่มเรียบร้อยแล้ว (สามารถเปิดดูแยกหมวดได้ในแท็บ สื่อ & ใบกิจกรรม / เครื่องมือวัดผล / พิมพ์เขียว & ข้อสอบ ด้านบน)")
                        break
                    paragraphs.append(txt)
            elif child.tag.endswith('tbl'):
                t = Table(child, doc)
                t_rows = []
                for row in t.rows:
                    r_cells = [c.text.strip() for c in row.cells]
                    t_rows.append(r_cells)
                if t_rows:
                    tables_data.append(t_rows)
    else:
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            t_rows = []
            for row in t.rows:
                r_cells = [c.text.strip() for c in row.cells]
                t_rows.append(r_cells)
            if t_rows:
                tables_data.append(t_rows)

    return {
        "filename": os.path.basename(filepath),
        "paragraphs": [clean_txt(p) for p in paragraphs],
        "tables": [[[clean_txt(c) for c in row] for row in tbl] for tbl in tables_data]
    }

def build_m2_academic_data():
    wk3_file = os.path.join(ROOT_M2, "01_เล่มเอกสารส่งวิชาการ_วก", "วก3_โครงสร้างรายวิชา", "วก3_โครงสร้างรายวิชา_ศ22102_ดนตรี2.docx")
    wk3_raw = parse_docx(wk3_file)

    m2_data = {
        "wk3": {
            "title": "วก.3 โครงสร้างรายวิชา ศ22102 ดนตรี 2 (ชั้น ม.2)",
            "code": "ศ22102",
            "name": "ดนตรี 2",
            "grade": "มัธยมศึกษาปีที่ 2",
            "term": "ภาคเรียนที่ 2/2569",
            "periods": 20,
            "credits": 0.5,
            "paragraphs": wk3_raw["paragraphs"],
            "tables": wk3_raw["tables"]
        },
        "wk4": [],
        "wk5": [],
        "manuals": [],
        "media": [],
        "eval": [],
        "supp": [],
        "tests": []
    }

    # 1. wk4 (4 units)
    wk4_dir = os.path.join(ROOT_M2, "01_เล่มเอกสารส่งวิชาการ_วก", "วก4_หน่วยการเรียนรู้")
    for f in sorted(os.listdir(wk4_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            u = parse_docx(os.path.join(wk4_dir, f), stop_at_appendix=True)
            unit_match = re.search(r"หน่วยการเรียนรู้ที่_(\d+)_(.+)\.docx", f)
            unit_no = int(unit_match.group(1)) if unit_match else len(m2_data["wk4"]) + 1
            unit_name = unit_match.group(2).replace("_", " ") if unit_match else f
            m2_data["wk4"].append({
                "id": f"unit_{unit_no}",
                "filename": f,
                "unit_no": unit_no,
                "unit_name": unit_name,
                "paragraphs": u["paragraphs"],
                "tables": u["tables"]
            })

    # 2. wk5 (18 plans)
    wk5_dir = os.path.join(ROOT_M2, "01_เล่มเอกสารส่งวิชาการ_วก", "วก5_แผนการจัดการเรียนรู้")
    for f in sorted(os.listdir(wk5_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            p = parse_docx(os.path.join(wk5_dir, f), stop_at_appendix=True)
            p_match = re.search(r"แผนที่_(\d+)", f)
            p_num = int(p_match.group(1)) if p_match else len(m2_data["wk5"]) + 1
            topic = f"แผนการจัดการเรียนรู้ที่ {p_num}"
            for line in p["paragraphs"][:10]:
                if line.startswith("เรื่อง"):
                    topic = re.sub(r"^เรื่อง\s*", "", line).strip()
                    if "เวลา" in topic:
                        topic = topic.split("เวลา")[0].strip()
                    break
            m2_data["wk5"].append({
                "period": p_num,
                "filename": f,
                "topic": topic,
                "paragraphs": p["paragraphs"],
                "tables": p["tables"]
            })

    # 3. manuals (22 files)
    m_wk4_dir = os.path.join(ROOT_M2, "02_คู่มือการสอนสำหรับครู", "วก4_คู่มือครู_หน่วยการเรียนรู้")
    for f in sorted(os.listdir(m_wk4_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            m = parse_docx(os.path.join(m_wk4_dir, f), stop_at_appendix=True)
            m2_data["manuals"].append({
                "idx": len(m2_data["manuals"]),
                "filename": f,
                "title": f.replace(".docx", "").replace("_", " "),
                "paragraphs": m["paragraphs"],
                "tables": m["tables"]
            })
    m_wk5_dir = os.path.join(ROOT_M2, "02_คู่มือการสอนสำหรับครู", "วก5_คู่มือครู_แผนการสอน")
    for f in sorted(os.listdir(m_wk5_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            m = parse_docx(os.path.join(m_wk5_dir, f), stop_at_appendix=True)
            m2_data["manuals"].append({
                "idx": len(m2_data["manuals"]),
                "filename": f,
                "title": f.replace(".docx", "").replace("_", " "),
                "paragraphs": m["paragraphs"],
                "tables": m["tables"]
            })

    # 4. media (18 files)
    media_dir = os.path.join(ROOT_M2, "03_ใบกิจกรรมและสื่อการสอน_แจกนักเรียน")
    for f in sorted(os.listdir(media_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            m = parse_docx(os.path.join(media_dir, f))
            m_match = re.search(r"ผ(\d+)_(.+)\.docx", f)
            p_num = int(m_match.group(1)) if m_match else len(m2_data["media"]) + 1
            title = m_match.group(2).replace("_", " ") if m_match else f
            m2_data["media"].append({
                "period": p_num,
                "filename": f,
                "title": title,
                "paragraphs": m["paragraphs"],
                "tables": m["tables"]
            })

    # 5. eval (8 files)
    eval_dir = os.path.join(ROOT_M2, "04_เครื่องมือวัดผลและข้อสอบ", "แบบประเมินและรูบริก")
    for idx, f in enumerate(sorted(os.listdir(eval_dir))):
        if f.endswith(".docx") and not f.startswith("~$"):
            e = parse_docx(os.path.join(eval_dir, f))
            num_match = re.search(r"ฉบับที่_(\d+)_(.+)\.docx", f)
            doc_no = int(num_match.group(1)) if num_match else idx + 1
            title = num_match.group(2).replace("_", " ") if num_match else f
            m2_data["eval"].append({
                "doc_no": doc_no,
                "filename": f,
                "title": title,
                "paragraphs": e["paragraphs"],
                "tables": e["tables"]
            })

    # 6. supp (3 files)
    supp_dir = os.path.join(ROOT_M2, "04_เครื่องมือวัดผลและข้อสอบ", "เอกสารประกอบ")
    for f in sorted(os.listdir(supp_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            s = parse_docx(os.path.join(supp_dir, f))
            m2_data["supp"].append({
                "filename": f,
                "title": f.replace(".docx", "").replace("_", " "),
                "paragraphs": s["paragraphs"],
                "tables": s["tables"]
            })

    # 7. tests (5 files)
    tests_dir = os.path.join(ROOT_M2, "04_เครื่องมือวัดผลและข้อสอบ", "แบบทดสอบท้ายหน่วยและพิมพ์เขียว")
    for f in sorted(os.listdir(tests_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            t = parse_docx(os.path.join(tests_dir, f))
            m2_data["tests"].append({
                "filename": f,
                "title": f.replace(".docx", "").replace("_", " "),
                "paragraphs": t["paragraphs"],
                "tables": t["tables"]
            })

    return m2_data

def extract_kpa(paragraphs):
    k, p, a = "", "", ""
    for i, para in enumerate(paragraphs):
        if "ด้านความรู้ (K)" in para and i + 1 < len(paragraphs) and paragraphs[i + 1].startswith("1.1"):
            k = re.sub(r"^1\.1\s*", "", paragraphs[i + 1]).strip()
        elif "ด้านทักษะกระบวนการ (P)" in para and i + 1 < len(paragraphs) and paragraphs[i + 1].startswith("2.1"):
            p = re.sub(r"^2\.1\s*", "", paragraphs[i + 1]).strip()
        elif "ด้านคุณลักษณะอันพึงประสงค์ (A)" in para and i + 1 < len(paragraphs) and paragraphs[i + 1].startswith("3.1"):
            a = re.sub(r"^3\.1\s*", "", paragraphs[i + 1]).strip()
    return {
        "k": k if k else "เข้าใจหลักการและองค์ประกอบทางดนตรีตามบทเรียน",
        "p": p if p else "ปฏิบัติทักษะทางดนตรีและกระบวนการตามลำดับขั้นตอน",
        "a": a if a else "มีความรับผิดชอบ มุ่งมั่น และร่วมมือในการเรียนรู้"
    }

def extract_steps(paragraphs):
    steps = []
    step_labels = {
        "T": "ขั้นที่ 1: T - Trigger & Tuning",
        "P": "ขั้นที่ 2: P - Practice & Precision (3-Tier Mastery)",
        "A": "ขั้นที่ 3: A - Active Collaboration & Analysis",
        "C": "ขั้นที่ 4: C - Creative Challenge (คิดค้นและสร้างสรรค์)",
        "E": "ขั้นที่ 5: E - Evaluation & Reflection"
    }
    for para in paragraphs:
        m = re.match(r"^([TPACE])\s*-\s*(\d+)\.\s*\((\d+)\s*นาที\)\s*(.+)$", para)
        if m:
            letter = m.group(1)
            dur = int(m.group(3))
            desc = m.group(4).strip()
            desc = re.sub(r"^\[[^\]]+\]\s*", "", desc).strip()
            steps.append({
                "phase": f"{step_labels.get(letter, letter)} ({dur} นาที)",
                "duration": dur,
                "desc": desc
            })
    return steps

def extract_mentor_tips(paragraphs):
    mode_txt = "Adaptive T-PACE Framework ควบคู่กับวงจร 3-Tier Mastery Loop (Bronze / Silver / Gold)"
    plan_b_txt = "เตรียมใบกิจกรรมฉบับกระดาษ (Hardcopy) ผังคีย์บอร์ดจำลอง และหลอดเทียบเสียงมาตรฐานกรณีออฟไลน์"
    for para in paragraphs:
        if "Adaptive T-PACE Framework" in para and "Mode" in para:
            m = re.search(r"\[(Mode [^\]]+)\]", para)
            if m:
                mode_txt = f"รูปแบบ Adaptive T-PACE ({m.group(1)}) ร่วมกับวงจร 3-Tier Mastery Loop (Bronze / Silver / Gold)"
        elif para.startswith("4. Plan B"):
            plan_b_txt = para.split(":", 1)[1].strip() if ":" in para else para
    return f"💡 กลยุทธ์ระดับเชี่ยวชาญ (คศ.4): ใช้ {mode_txt} คาลิมบ้าเล่นในคีย์ C Major ล้วน • แผนสำรอง (Plan B): {plan_b_txt}"

def build_m2_lessons(m2_academic):
    m2_lessons = []
    for p in m2_academic["wk5"]:
        p_num = p["period"]
        unit_no = 1
        unit_name = "ทบทวนคาลิมบ้าและเสียงประสาน"
        if 5 <= p_num <= 11:
            unit_no = 2
            unit_name = "ปฏิบัติรวมวงคาลิมบ้าเพลงวน"
        elif 12 <= p_num <= 15:
            unit_no = 3
            unit_name = "รอบรู้ดนตรีข้ามวัฒนธรรมและอาชีพดนตรี"
        elif 16 <= p_num <= 18:
            unit_no = 4
            unit_name = "ดนตรีกับสังคมและการวิจารณ์เชิงประจักษ์"

        kpa = extract_kpa(p["paragraphs"])
        steps = extract_steps(p["paragraphs"])
        mentor_tips = extract_mentor_tips(p["paragraphs"])
        matching_media = next((m for m in m2_academic["media"] if m["period"] == p_num), None)
        media_items = [
            {
                "title": f"เอกสาร วก.5 แผนการจัดการเรียนรู้ที่ {p_num:02d} ({p['topic']})",
                "type": "doc",
                "url": f"academic/m2/wk5/{p['filename']}"
            }
        ]
        if matching_media:
            media_items.append({
                "title": f"ใบกิจกรรมประกอบแผน: {matching_media['title']}",
                "type": "worksheet",
                "url": f"academic/m2/media/{matching_media['filename']}"
            })

        rubrics = []
        if p["tables"] and len(p["tables"]) > 0:
            tbl = p["tables"][0]
            for idx_r, row in enumerate(tbl[1:4]):
                if len(row) >= 4:
                    aspect = row[0]
                    weight = row[1] if len(row) >= 5 else ""
                    desc4 = row[2] if len(row) >= 5 else row[1]
                    label_str = f"{aspect} ({weight})" if weight else aspect
                    rubrics.append({
                        "score": 3 - idx_r if idx_r < 3 else 1,
                        "label": label_str,
                        "desc": desc4
                    })

        m2_lessons.append({
            "lesson_id": f"M2-P{p_num:02d}",
            "grade": "M2",
            "period": p_num,
            "unit_no": unit_no,
            "unit_name": unit_name,
            "topic": p["topic"],
            "kpa": kpa,
            "mentor_tips": mentor_tips,
            "duration_min": 50,
            "steps": steps,
            "media": media_items,
            "assessment": {
                "has_score": True,
                "max_score": 3,
                "type": "rubric",
                "rubrics": rubrics[:3]
            }
        })

    # Period 19: Midterm Exam (10 pts: K:2, P:8)
    m2_lessons.append({
        "lesson_id": "M2-P19",
        "grade": "M2",
        "period": 19,
        "unit_no": 2,
        "unit_name": "การสอบวัดผลกลางภาคเรียน (10 คะแนน)",
        "topic": "การสอบวัดผลกลางภาคเรียน (Midterm Examination — K:2, P:8)",
        "kpa": {
            "k": "ทดสอบความรู้ทฤษฎีคู่เสียงประสาน โน้ตแปลงเสียง (#, b, ♮) การย้ายคีย์สู่คีย์ C และหลักการเพลงวน (2 คะแนน)",
            "p": "สอบปฏิบัติคาลิมบ้าคีย์ C Major (คู่เสียงประสาน Dyads, คอร์ด C-Am-F-G และเพลงวน Canon) (8 คะแนน)",
            "a": "มีความซื่อสัตย์สุจริต มีสมาธิ และตรงต่อเวลาในการสอบวัดผล"
        },
        "mentor_tips": "💡 กลยุทธ์บริหารเวลา 50 นาที: 10 นาทีแรกสอบปรนัย (K:2) พร้อมกันทั้งชั้น • 35 นาทีถัดมาสอบปฏิบัติเป็นสถานีกลุ่มย่อย (Ensemble Station ครั้งละ 4 คน × 10 รอบ รอบละ 3.5 นาที) • 5 นาทีสุดท้ายตรวจสอบคะแนน",
        "duration_min": 50,
        "steps": [
            {"phase": "ขั้นที่ 1: สอบภาคทฤษฎี K:2 (10 นาที)", "duration": 10, "desc": "นักเรียนทำแบบทดสอบวัดผลกลางภาคภาคทฤษฎี (โน้ตแปลงเสียง การย้ายคีย์สู่ C Major และโครงสร้างคอร์ด)"},
            {"phase": "ขั้นที่ 2: สอบปฏิบัติคาลิมบ้าคีย์ C — P:8 (35 นาที)", "duration": 35, "desc": "สอบปฏิบัติแบบสถานีกลุ่มย่อย (ครั้งละ 4 คน รวม 10 รอบ รอบละ 3.5 นาที) บรรเลงคอร์ด C-Am-F-G และเพลงวน Canon ในคีย์ C Major"},
            {"phase": "ขั้นที่ 3: สรุปและยืนยันคะแนน (5 นาที)", "duration": 5, "desc": "ตรวจเช็กความครบถ้วนของคะแนนสอบกลางภาค (10 คะแนน) รายบุคคล"}
        ],
        "media": [
            {"title": "พิมพ์เขียวข้อสอบ Test Blueprint ศ22102 (40 ข้อ)", "type": "doc", "url": "academic/m2/tests/พิมพ์เขียวข้อสอบ_Test_Blueprint_ศ22102_ดนตรี2.docx"}
        ],
        "assessment": {
            "has_score": True,
            "max_score": 10,
            "type": "exam",
            "rubrics": [
                {"score": 3, "label": "ดีเยี่ยม (8-10 คะแนน)", "desc": "ทำข้อสอบทฤษฎีและปฏิบัติคาลิมบ้าคีย์ C ได้แม่นยำ 8–10 คะแนน"},
                {"score": 2, "label": "ดี (6-7 คะแนน)", "desc": "ได้คะแนนรวม 6–7 คะแนน"},
                {"score": 1, "label": "ผ่านเกณฑ์ขั้นต่ำ (5 คะแนน)", "desc": "ได้คะแนนรวม 5 คะแนน"}
            ]
        }
    })

    # Period 20: Final Exam (10 pts: K:2, P:8)
    m2_lessons.append({
        "lesson_id": "M2-P20",
        "grade": "M2",
        "period": 20,
        "unit_no": 4,
        "unit_name": "การสอบวัดผลปลายภาคเรียน (10 คะแนน)",
        "topic": "การสอบวัดผลปลายภาคเรียนและนิทรรศการภัณฑารักษ์ดนตรี (Final Examination — K:2, P:8)",
        "kpa": {
            "k": "ทดสอบความรู้รวบยอดเรื่องดนตรีข้ามวัฒนธรรม อาชีพดนตรี และดนตรีกับสังคมตามพิมพ์เขียวข้อสอบ (2 คะแนน)",
            "p": "นำเสนอผลงานภัณฑารักษ์ดนตรีเพื่อสังคม (Social Soundscape Pitching) และแฟ้มสะสมงานพัฒนาการตนเอง (8 คะแนน)",
            "a": "มีความภาคภูมิใจในพัฒนาการทางดนตรีของตนเองและเห็นคุณค่าของดนตรีต่อสังคม"
        },
        "mentor_tips": "💡 กลยุทธ์บริหารเวลา 50 นาที: ใช้ระบบ Gallery Walk & Curator Pitching (30 นาที) ประเมินปฏิบัติปลายภาค (P:8) ควบคู่กับการทำข้อสอบปรนัยปลายภาค (K:2) ช่วง 15 นาทีท้ายคาบ",
        "duration_min": 50,
        "steps": [
            {"phase": "ขั้นที่ 1: Gallery Walk & Curator Pitching — P:8 (30 นาที)", "duration": 30, "desc": "นำเสนอผลงานเพลย์ลิสต์เพื่อสังคมและผลงานคาลิมบ้าคีย์ C แบบหมุนเวียนฐานนิทรรศการ (8 กลุ่ม กลุ่มละ 3.5 นาที)"},
            {"phase": "ขั้นที่ 2: สอบภาคทฤษฎีปลายภาค — K:2 (15 นาที)", "duration": 15, "desc": "นักเรียนทำแบบทดสอบวัดผลสัมฤทธิ์ปลายภาคเรียนตามพิมพ์เขียวข้อสอบ (Test Blueprint ศ22102)"},
            {"phase": "ขั้นที่ 3: สรุปผลสัมฤทธิ์ตลอดภาคเรียน (5 นาที)", "duration": 5, "desc": "ครูสรุปคะแนนรวม 100 คะแนน (4 หน่วย 80 + กลางภาค 10 + ปลายภาค 10) และชื่นชมความสำเร็จของผู้เรียน"}
        ],
        "media": [
            {"title": "พิมพ์เขียวข้อสอบ Test Blueprint ศ22102 (40 ข้อ)", "type": "doc", "url": "academic/m2/tests/พิมพ์เขียวข้อสอบ_Test_Blueprint_ศ22102_ดนตรี2.docx"}
        ],
        "assessment": {
            "has_score": True,
            "max_score": 10,
            "type": "exam",
            "rubrics": [
                {"score": 3, "label": "ดีเยี่ยม (8-10 คะแนน)", "desc": "นำเสนอผลงานและทำข้อสอบปลายภาคได้ 8–10 คะแนน"},
                {"score": 2, "label": "ดี (6-7 คะแนน)", "desc": "ได้คะแนนรวม 6–7 คะแนน"},
                {"score": 1, "label": "ผ่านเกณฑ์ขั้นต่ำ (5 คะแนน)", "desc": "ได้คะแนนรวม 5 คะแนน"}
            ]
        }
    })
    return m2_lessons

if __name__ == "__main__":
    print("🔄 กำลังอ่านและแปลงไฟล์ Word ม.2 ทั้ง 79 ไฟล์เข้าสู่ฐานข้อมูลเว็บแอป MusicSmart...")
    m2_academic = build_m2_academic_data()

    # Load existing M1 academic_docs.js
    with open("app/data/academic_docs.js", "r", encoding="utf-8") as f:
        js_txt = f.read()
    if "const ACADEMIC_DOCS_BY_GRADE =" in js_txt:
        json_part = js_txt.split("const ACADEMIC_DOCS_BY_GRADE = ")[1].split(";\nlet ACADEMIC_DOCS")[0]
        by_grade = json.loads(json_part)
        m1_academic = by_grade["M1"]
    else:
        json_part = js_txt.split("const ACADEMIC_DOCS = ")[1].split(";\n\nif (typeof module")[0]
        m1_academic = json.loads(json_part)

    combined_by_grade = {
        "M1": m1_academic,
        "M2": m2_academic
    }

    out_academic_js = (
        "// ==========================================================================\n"
        "// Phayaophitthayakhom School - Academic Documents Knowledge Base (วก. 1-7)\n"
        "// รองรับทั้ง ชั้น ม.1 (ศ21102 ดนตรี 1 - 79 ไฟล์) และ ชั้น ม.2 (ศ22102 ดนตรี 2 - 79 ไฟล์)\n"
        "// ==========================================================================\n\n"
        "const ACADEMIC_DOCS_BY_GRADE = " + json.dumps(combined_by_grade, ensure_ascii=False, indent=2) + ";\n"
        "let ACADEMIC_DOCS = ACADEMIC_DOCS_BY_GRADE.M1;\n\n"
        "if (typeof module !== 'undefined' && module.exports) {\n"
        "  module.exports = { ACADEMIC_DOCS, ACADEMIC_DOCS_BY_GRADE };\n"
        "}\n"
    )
    with open("app/data/academic_docs.js", "w", encoding="utf-8") as f:
        f.write(out_academic_js)
    print(f"✅ อัปเดต app/data/academic_docs.js สำเร็จ! (รองรับ M1 79 ไฟล์ + M2 79 ไฟล์ รวม 158 ไฟล์)")

    # Update app/data/lessons.js with both M1 and M2
    with open("app/data/lessons.js", "r", encoding="utf-8") as f:
        lessons_txt = f.read()
    m1_json_str = lessons_txt.split('"M1": ')[1].split(',\n  "M2":')[0]
    m1_lessons = json.loads(m1_json_str)
    m2_lessons = build_m2_lessons(m2_academic)

    new_lessons_js = (
        "// ==========================================================================\n"
        "// Generated Lesson Plans for Music Education (PPK)\n"
        "// บูรณาการหลักสูตรจริง วก.3 - วก.5 โรงเรียนพะเยาพิทยาคม (M.1 ศ21102 & M.2 ศ22102)\n"
        "// ==========================================================================\n\n"
        "const LESSONS_DATA = {\n"
        '  "M1": ' + json.dumps(m1_lessons, ensure_ascii=False, indent=4) + ",\n"
        '  "M2": ' + json.dumps(m2_lessons, ensure_ascii=False, indent=4) + "\n"
        "};\n\n"
        "if (typeof module !== 'undefined' && module.exports) {\n"
        "  module.exports = LESSONS_DATA;\n"
        "}\n"
    )
    with open("app/data/lessons.js", "w", encoding="utf-8") as f:
        f.write(new_lessons_js)
    print(f"✅ อัปเดต app/data/lessons.js สำเร็จ! (M1: {len(m1_lessons)} คาบ, M2: {len(m2_lessons)} คาบ)")
