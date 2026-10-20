import os
import sys
import re
import docx
from docxcompose.composer import Composer

sys.stdout.reconfigure(encoding='utf-8')

base = r'D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101'
media_dir = os.path.join(base, 'สื่อการสอน')
eval_dir = os.path.join(base, 'เครื่องมือวัดผล')
test_dir = os.path.join(base, 'แบบทดสอบและพิมพ์เขียว')
supp_dir = os.path.join(base, 'เอกสารประกอบ')

units_info = {
    1: {'title': 'พื้นฐานคาลิมบ้า', 'plans': range(1, 5)},
    2: {'title': 'ปฏิบัติรวมวงคาลิมบ้า', 'plans': range(5, 12)},
    3: {'title': 'รอบรู้วงดนตรีและวัฒนธรรม', 'plans': range(12, 16)},
    4: {'title': 'สุนทรียภาพแห่งการฟัง', 'plans': range(16, 19)}
}

def get_unit_info(plan_num):
    for u_num, data in units_info.items():
        if plan_num in data['plans']:
            return u_num, data['title']
    return 1, 'พื้นฐานคาลิมบ้า'

plan_topics = {
    1: 'กายวิภาคและการก้องกังวาน (Acoustics & Posture)',
    2: 'รหัสนามธรรมสู่ระบบเสียงและการขับร้อง (Tab, 5-Line Staff & Solfege Singing)',
    3: 'จังหวะ ทิศทางเสียง และการขับร้องบันไดเสียง (C Major Scale, Singing & Metronome)',
    4: 'การบรรเลงเดี่ยวเชิงประยุกต์และการขับร้องประกอบ (Solo Performance, Singing & Care)',
    5: 'พลวัตของการรวมวง (Ensemble Dynamics)',
    6: 'นวัตกรรมเครื่องประกอบจังหวะ (Rhythm Invention & Upcycling)',
    7: 'การกำกับตนเองในการซ้อมกลุ่มย่อย (Sectional Rehearsal 1)',
    8: 'การกำกับตนเองในการซ้อมกลุ่มย่อย (Sectional Rehearsal 2 & Tempo Mastery)',
    9: 'สมดุลแห่งเสียง (Balance & Run-through)',
    10: 'Mini-Recital: ประเมินเพื่อพัฒนา ตอนที่ 1 (กลุ่ม 1-4)',
    11: 'Mini-Recital: ประเมินเพื่อพัฒนา ตอนที่ 2 (กลุ่ม 5-8 และสรุปบทเรียน)',
    12: 'การจัดกลุ่มเสียงวงดนตรีไทย (Thai Ensemble Classification)',
    13: 'สุนทรียะและวิวัฒนาการดนตรีสากล (Western Ensemble & Evolution)',
    14: 'การสื่อสารข้ามวัฒนธรรม (Cross-Cultural Scriptwriting)',
    15: 'นวัตกรรมสื่อเสียง (Podcast Production & Digital Audio Guide)',
    16: 'ถอดรหัสอารมณ์เพลง (Music & Emotion Analytics)',
    17: "สุนทรียศาสตร์และการวิจารณ์ (Feldman's Art Criticism Method)",
    18: 'การรังสรรค์เพลย์ลิสต์ส่วนบุคคล (Digital Playlist Curation & Reflection)'
}

def get_media_path(p_num):
    prefix = f"ผ{p_num:02d}"
    for f in os.listdir(media_dir):
        if prefix in f and f.endswith('.docx') and not f.startswith('~$'):
            return os.path.join(media_dir, f), f
    return None, None

def get_eval_paths(p_num):
    results = []
    tool_map = {
        1: [1, 3],
        2: [1, 3],
        3: [1, 3],
        4: [1, 3],
        5: [2, 3],
        6: [2, 3],
        7: [2, 3],
        8: [2, 3],
        9: [2, 3],
        10: [2, 4, 3],
        11: [2, 3],
        12: [5, 8],
        13: [5, 8],
        14: [5, 8],
        15: [5, 8],
        16: [6, 7, 8],
        17: [7, 6, 8],
        18: [6, 7, 8],
    }
    target_tools = tool_map.get(p_num, [])
    for t_id in target_tools:
        prefix = f"ฉบับที่_{t_id}_"
        for f in os.listdir(eval_dir):
            if prefix in f and f.endswith('.docx') and not f.startswith('~$'):
                results.append((os.path.join(eval_dir, f), f))
                break
    return results

def get_test_path(p_num):
    unit_test_map = {
        4: 'แบบทดสอบท้ายหน่วย_หน่วยที่_1_พื้นฐานคาลิมบ้า.docx',
        11: 'แบบทดสอบท้ายหน่วย_หน่วยที่_2_ปฏิบัติรวมวงคาลิมบ้า.docx',
        15: 'แบบทดสอบท้ายหน่วย_หน่วยที่_3_รอบรู้วงดนตรีและวัฒนธรรม.docx',
        18: 'แบบทดสอบท้ายหน่วย_หน่วยที่_4_สุนทรียภาพแห่งการฟัง.docx'
    }
    if p_num in unit_test_map:
        fn = unit_test_map[p_num]
        return os.path.join(test_dir, fn), fn
    return None, None

def get_supp_paths(p_num):
    supps = []
    if p_num == 1:
        fn = 'เอกสารวิชาการ_กรอบแนวคิดการจัดการเรียนรู้_T_PACE_Framework.docx'
        supps.append((os.path.join(supp_dir, fn), fn))
    if p_num in [10, 15]:
        fn = 'หนังสือยินยอมการคุ้มครองข้อมูลส่วนบุคคล_PDPA_Consent_Form.docx'
        supps.append((os.path.join(supp_dir, fn), fn))
    if p_num == 11:
        fn = 'วก7_รายงานการวิจัยในชั้นเรียนอย่างย่อ_ศ21101.docx'
        supps.append((os.path.join(supp_dir, fn), fn))
    return supps

def add_divider_page(doc, title="ภาคผนวก", subtitle="", items=None, letter="", item_title=""):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = docx.shared.Pt(26)
    r.font.color.rgb = docx.shared.RGBColor(0x1A, 0x36, 0x5D)
    
    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.size = docx.shared.Pt(16)
        r_sub.bold = True
        r_sub.font.color.rgb = docx.shared.RGBColor(0x2B, 0x6C, 0xB0)
        
    p_c = doc.add_paragraph()
    p_c.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run("รายวิชา ดนตรี 1 (ศ21101) ชั้นมัธยมศึกษาปีที่ 1 ภาคเรียนที่ 1 โรงเรียนพะเยาพิทยาคม")
    r_c.font.size = docx.shared.Pt(13)
    r_c.font.color.rgb = docx.shared.RGBColor(0x71, 0x80, 0x96)
    
    p_sep = doc.add_paragraph()
    p_sep.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r_sep = p_sep.add_run("―" * 45)
    r_sep.font.color.rgb = docx.shared.RGBColor(0xCB, 0xD5, 0xE0)
    
    if letter and item_title:
        p_sec = doc.add_paragraph()
        p_sec.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        r_sec = p_sec.add_run(f"ภาคผนวก {letter}: {item_title}")
        r_sec.bold = True
        r_sec.font.size = docx.shared.Pt(18)
        r_sec.font.color.rgb = docx.shared.RGBColor(0x2B, 0x6C, 0xB0)
        
    if items:
        p_hdr = doc.add_paragraph()
        p_hdr.paragraph_format.left_indent = docx.shared.Inches(0.8)
        r_hdr = p_hdr.add_run("สารบัญเอกสารแนบท้าย (ภาคผนวก):")
        r_hdr.bold = True
        r_hdr.font.size = docx.shared.Pt(14)
        
        for it in items:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = docx.shared.Inches(1.1)
            r_b = p_it.add_run("• ")
            r_b.bold = True
            r_b.font.color.rgb = docx.shared.RGBColor(0x2B, 0x6C, 0xB0)
            r_t = p_it.add_run(it)
            r_t.font.size = docx.shared.Pt(12)
            
    p_sep2 = doc.add_paragraph()
    p_sep2.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r_sep2 = p_sep2.add_run("―" * 45)
    r_sep2.font.color.rgb = docx.shared.RGBColor(0xCB, 0xD5, 0xE0)
    doc.add_page_break()

def add_teacher_operational_guide(doc, p_num, p_topic, u_num, u_title):
    # Check if Section 7 already added
    for p in doc.paragraphs:
        if "หมวดที่ 7: คู่มือปฏิบัติการจัดการเรียนรู้" in p.text:
            return

    p_h1 = doc.add_paragraph()
    r_h1 = p_h1.add_run("หมวดที่ 7: คู่มือปฏิบัติการจัดการเรียนรู้ สื่อการสอน และเครื่องมือวัดผลสำหรับครู (Teacher's Operational Manual)")
    r_h1.bold = True
    r_h1.font.size = docx.shared.Pt(14)
    r_h1.font.color.rgb = docx.shared.RGBColor(0x1A, 0x36, 0x5D)

    # 7.1 In-class media execution
    p_71 = doc.add_paragraph()
    r_71 = p_71.add_run("7.1 แนวทางปฏิบัติการใช้สื่อการเรียนรู้และใบกิจกรรมในชั้นเรียน:")
    r_71.bold = True
    r_71.font.size = docx.shared.Pt(12)

    doc.add_paragraph(
        f"1. วัตถุประสงค์ของสื่อหลักประจำแผน: สื่อการเรียนรู้และใบกิจกรรมประจำแผนที่ {p_num} ออกแบบมาเพื่อให้นักเรียนได้ลงมือปฏิบัติจริง (Hands-on Experience) "
        f"สอดคล้องกับหัวข้อเรื่อง {p_topic} โดยเน้นการสร้างองค์ความรู้ด้วยตนเองผ่านกระบวนการกลุ่มบัดดี้ (Buddy Coaching: 1 คนปฏิบัติ 1 คนสังเกตและให้ข้อเสนอแนะ)"
    )
    doc.add_paragraph(
        "2. ขั้นตอนการจัดกิจกรรมและบทบาทครู: ครูต้องชี้แจงกติกาและเป้าหมายย่อย (Micro-goal) ให้ชัดเจนก่อนแจกอุปกรณ์ เมื่อนักเรียนเริ่มทำกิจกรรม ครูเดินสังเกตเชิงรุก (Active Facilitation) "
        "คอยกระตุ้นคำถามปลายเปิด (Scaffolding Questioning) แทนการบอกคำตอบโดยตรง เพื่อให้นักเรียนค้นพบระดับเสียงและทักษะด้วยตนเอง"
    )
    doc.add_paragraph(
        "3. แผนสำรองกรณีฉุกเฉิน (Plan B Offline Protocol): หากระบบดิจิทัลหรือสัญญาณเครือข่ายอินเทอร์เน็ตขัดข้อง ให้ครูแจกใบกิจกรรมฉบับกระดาษ (Hardcopy) ที่แนบในภาคผนวก ก "
        "และใช้เครื่องมือเคาะจังหวะกายภาพ (ส้อมเสียง C-523.3Hz / Metronome เครื่องกล) ดำเนินการเรียนรู้ตามขั้นตอนเดิมได้ทันที 100% โดยไม่ทำให้กระบวนการเรียนรู้หยุดชะงัก"
    )

    # 7.2 Assessment protocol
    p_72 = doc.add_paragraph()
    r_72 = p_72.add_run("7.2 แนวทางปฏิบัติการใช้เครื่องมือวัดและประเมินผล (Rubrics & Observation Protocol):")
    r_72.bold = True
    r_72.font.size = docx.shared.Pt(12)

    doc.add_paragraph(
        "1. การประเมินตามสภาพจริง (Authentic Assessment): ใช้เกณฑ์รูบริก (Rubrics) ที่แนบในภาคผนวก ข ในการประเมินทักษะกระบวนการ (P) และสังเกตคุณลักษณะอันพึงประสงค์ (A) "
        "โดยครูต้องสังเกตพฤติกรรมเชิงประจักษ์อย่างต่อเนื่องตลอด 50 นาที ไม่ตัดสินเพียงผลงานสุดท้าย"
    )
    doc.add_paragraph(
        "2. เกณฑ์การตัดสิน 4 ระดับคุณภาพ: ระดับ 4 (ดีเยี่ยม: ทำได้ถูกต้องคล่องแคล่ว มีความคิดสร้างสรรค์ ช่วยเหลือเพื่อน), ระดับ 3 (ดี: ปฏิบัติได้ถูกต้องตามเกณฑ์มาตรฐาน), "
        "ระดับ 2 (พอใช้: ปฏิบัติตามคำแนะนำได้บางส่วน), และระดับ 1 (ปรับปรุง: ต้องได้รับการโค้ชชิ่งซ้ำ)"
    )
    doc.add_paragraph(
        "3. การให้ข้อมูลย้อนกลับเพื่อพัฒนา (Formative Actionable Feedback): หลังเสร็จสิ้นกิจกรรม ครูต้องให้ข้อมูลป้อนกลับทันทีแก่นักเรียนรายบุคคลหรือกลุ่มบัดดี้ "
        "โดยระบุจุดเด่น 1 ข้อ และจุดที่ต้องพัฒนาเพิ่มเติม 1 ข้อ พร้อมบันทึกคะแนนสะสมลงในระบบดิจิทัล MusicSmart เพื่อนำไปตัดเกรดตามโครงสร้าง วก.3"
    )

    # 7.3 Test & CAR protocol
    p_73 = doc.add_paragraph()
    r_73 = p_73.add_run("7.3 คำแนะนำการบริหารจัดการแบบทดสอบ เอกสารประกอบ และการวิจัยในชั้นเรียน (CAR):")
    r_73.bold = True
    r_73.font.size = docx.shared.Pt(12)

    doc.add_paragraph(
        "1. การบริหารแบบทดสอบและเฉลย: ในแผนที่มีการวัดผลสัมฤทธิ์ท้ายหน่วย ครูต้องจัดห้องเรียนให้โปร่ง สุ่มสลับตัวเลือกและข้อสอบ (Randomization) เพื่อวัดความรู้ความเข้าใจที่แท้จริง "
        "และนำผลการสอบมาวิเคราะห์ความยากง่ายและอำนาจจำแนกเพื่อจัดกลุ่มสอนซ่อมเสริมแก่นักเรียนกลุ่มเสี่ยง"
    )
    doc.add_paragraph(
        "2. การปฏิบัติตาม พ.ร.บ. คุ้มครองข้อมูลส่วนบุคคล (PDPA): ในคาบที่มีการบันทึกภาพนิ่ง คลิปวิดีโอการแสดง หรือไฟล์เสียงพอดคาสต์เพื่อการเรียนการสอนและการประเมิน วPA "
        "ครูต้องจัดเก็บหนังสือยินยอม (PDPA Consent Form) จากผู้ปกครองและนักเรียนไว้เป็นหลักฐานทุกครั้ง"
    )
    doc.add_paragraph(
        "3. การต่อยอดสู่การวิจัยในชั้นเรียน (CAR) ระดับ คศ.4: ครูนำบันทึกหลังสอนในหมวดที่ 6 และสถิติจาก Rubrics/แอป MusicSmart มาวิเคราะห์ปัญหาและสร้างนวัตกรรมแก้ปัญหาเชิงประจักษ์อย่างต่อเนื่อง "
        "สอดคล้องกับรายงานการวิจัยในชั้นเรียน วก.7 เพื่อเตรียมความพร้อมสำหรับประเมินเลื่อนวิทยฐานะเชี่ยวชาญ"
    )

def assemble_plan_document(target_path, p_num, is_teacher_manual=False):
    u_num, u_title = get_unit_info(p_num)
    p_topic = plan_topics[p_num]
    
    master_doc = docx.Document(target_path)
    
    # If teacher manual, add the operational guide first!
    if is_teacher_manual:
        add_teacher_operational_guide(master_doc, p_num, p_topic, u_num, u_title)
        
    # Check if appendix already exists in doc
    has_appendix = any("สารบัญเอกสารแนบท้าย (ภาคผนวก)" in p.text for p in master_doc.paragraphs)
    if has_appendix:
        print(f"Skipping {os.path.basename(target_path)}: Appendix already assembled.")
        return
        
    # Gather materials for this plan
    media_path, media_fn = get_media_path(p_num)
    eval_paths = get_eval_paths(p_num)
    test_path, test_fn = get_test_path(p_num)
    supp_paths = get_supp_paths(p_num)
    
    appendix_items = []
    if media_fn:
        appendix_items.append(f"ภาคผนวก ก: สื่อการเรียนรู้และใบกิจกรรมประจำแผน ({media_fn})")
    if eval_paths:
        eval_names = ", ".join(fn for _, fn in eval_paths)
        appendix_items.append(f"ภาคผนวก ข: เครื่องมือวัดและประเมินผล ({eval_names})")
    if test_fn:
        appendix_items.append(f"ภาคผนวก ค: แบบทดสอบวัดผลสัมฤทธิ์ทางการเรียนท้ายหน่วย ({test_fn})")
    if supp_paths:
        supp_names = ", ".join(fn for _, fn in supp_paths)
        appendix_items.append(f"ภาคผนวก ง: เอกสารประกอบการจัดการเรียนรู้ ({supp_names})")
        
    # Add Master Appendix Divider
    add_divider_page(
        master_doc, 
        title="ภาคผนวก", 
        subtitle=f"แผนการจัดการเรียนรู้ที่ {p_num}: {p_topic}", 
        items=appendix_items
    )
    
    composer = Composer(master_doc)
    
    # Append Media (ภาคผนวก ก)
    if media_path and os.path.exists(media_path):
        m_doc = docx.Document(media_path)
        composer.append(m_doc)
        
    # Append Evaluation Tools (ภาคผนวก ข)
    for ep, efn in eval_paths:
        if os.path.exists(ep):
            e_doc = docx.Document(ep)
            composer.append(e_doc)
            
    # Append Unit Test (ภาคผนวก ค)
    if test_path and os.path.exists(test_path):
        t_doc = docx.Document(test_path)
        composer.append(t_doc)
        
    # Append Supplementary Docs (ภาคผนวก ง)
    for sp, sfn in supp_paths:
        if os.path.exists(sp):
            s_doc = docx.Document(sp)
            composer.append(s_doc)
            
    composer.save(target_path)
    print(f"✅ Successfully assembled: {os.path.basename(target_path)}")

# --------------------------------------------------------------------------
# Execute Assembly for วก5, แผนการสอนพร้อมใช้, and คู่มือครู วก5
# --------------------------------------------------------------------------
print("====================================================================")
print("ASSEMBLING วก5 (18 Plans)")
print("====================================================================")
wk5_dir = os.path.join(base, 'วก5')
for f in sorted(os.listdir(wk5_dir)):
    if f.endswith('.docx') and not f.startswith('~$'):
        m = re.search(r'แผนที่_(\d+)', f)
        if m:
            p_num = int(m.group(1))
            assemble_plan_document(os.path.join(wk5_dir, f), p_num, is_teacher_manual=False)

print("\n====================================================================")
print("ASSEMBLING แผนการสอนพร้อมใช้ (18 Plans)")
print("====================================================================")
ready_dir = os.path.join(base, 'แผนการสอนพร้อมใช้')
for f in sorted(os.listdir(ready_dir)):
    if f.endswith('.docx') and not f.startswith('~$'):
        m = re.search(r'แผนที่_(\d+)', f)
        if m:
            p_num = int(m.group(1))
            assemble_plan_document(os.path.join(ready_dir, f), p_num, is_teacher_manual=False)

print("\n====================================================================")
print("ASSEMBLING คู่มือการสอนสำหรับครู วก5 (18 Plans)")
print("====================================================================")
tm_wk5_dir = os.path.join(base, 'คู่มือการสอนสำหรับครู', 'คู่มือครู ม.1 ดนตรี1 ศ21101', 'วก5')
for f in sorted(os.listdir(tm_wk5_dir)):
    if f.endswith('.docx') and not f.startswith('~$'):
        m = re.search(r'แผนที่_(\d+)', f)
        if m:
            p_num = int(m.group(1))
            assemble_plan_document(os.path.join(tm_wk5_dir, f), p_num, is_teacher_manual=True)

print("\nAssembly of 18 plans across all 3 sets complete!")
