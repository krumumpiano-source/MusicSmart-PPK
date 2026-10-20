import os
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
import json

root_dir = r"D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101"

def parse_docx(filepath, stop_at_appendix=False):
    doc = docx.Document(filepath)
    paragraphs = []
    tables_data = []

    if stop_at_appendix:
        for child in doc.element.body:
            if child.tag.endswith('p'):
                p = Paragraph(child, doc)
                txt = p.text.strip()
                if txt:
                    if txt.startswith("ภาคผนวก"):
                        paragraphs.append("📎 หมายเหตุ: เอกสารฉบับเต็มในไฟล์ Word (.docx) ได้ผนวกรวมภาคผนวก (ใบกิจกรรม, เครื่องมือวัดผลรูบริก และแบบทดสอบ) ไว้ท้ายเล่มเรียบร้อยแล้ว (สามารถเปิดดูแยกหมวดได้ในแท็บ สื่อ & ใบกิจกรรม / เครื่องมือวัดผล / พิมพ์เขียว & ข้อสอบ ด้านบน)")
                        break
                    paragraphs.append(txt)
            elif child.tag.endswith('tbl'):
                t = Table(child, doc)
                t_rows = []
                for row in t.rows:
                    r_cells = [c.text.strip() for c in row.cells]
                    t_rows.append(r_cells)
                if t_rows:
                    tables_data.append(t_rows)
    else:
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            t_rows = []
            for row in t.rows:
                r_cells = [c.text.strip() for c in row.cells]
                t_rows.append(r_cells)
            if t_rows:
                tables_data.append(t_rows)

    return {
        "filename": os.path.basename(filepath),
        "paragraphs": paragraphs,
        "tables": tables_data
    }

# 1. Parse วก3 (1 file)
wk3_file = os.path.join(root_dir, "01_เล่มเอกสารส่งวิชาการ_วก", "วก3_โครงสร้างรายวิชา", "วก3_โครงสร้างรายวิชา_ศ21101_ดนตรี1.docx")
wk3_data = parse_docx(wk3_file) if os.path.exists(wk3_file) else None

# 2. Parse วก4 (4 units)
wk4_dir = os.path.join(root_dir, "01_เล่มเอกสารส่งวิชาการ_วก", "วก4_หน่วยการเรียนรู้")
wk4_data = []
if os.path.exists(wk4_dir):
    for f in sorted(os.listdir(wk4_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            wk4_data.append(parse_docx(os.path.join(wk4_dir, f), stop_at_appendix=True))

# 3. Parse วก5 (18 plans)
wk5_dir = os.path.join(root_dir, "01_เล่มเอกสารส่งวิชาการ_วก", "วก5_แผนการจัดการเรียนรู้")
wk5_data = []
if os.path.exists(wk5_dir):
    for f in sorted(os.listdir(wk5_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            wk5_data.append(parse_docx(os.path.join(wk5_dir, f), stop_at_appendix=True))

# 3.5 Parse คู่มือการสอนสำหรับครู (4 unit manuals + 18 lesson manuals = 22 files)
manuals_data = []
manual_wk4_dir = os.path.join(root_dir, "02_คู่มือการสอนสำหรับครู", "วก4_คู่มือครู_หน่วยการเรียนรู้")
if os.path.exists(manual_wk4_dir):
    for f in sorted(os.listdir(manual_wk4_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            manuals_data.append(parse_docx(os.path.join(manual_wk4_dir, f)))
manual_wk5_dir = os.path.join(root_dir, "02_คู่มือการสอนสำหรับครู", "วก5_คู่มือครู_แผนการสอน")
if os.path.exists(manual_wk5_dir):
    for f in sorted(os.listdir(manual_wk5_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            manuals_data.append(parse_docx(os.path.join(manual_wk5_dir, f)))

# 4. Parse สื่อการสอน (18 media/worksheets)
media_dir = os.path.join(root_dir, "03_ใบกิจกรรมและสื่อการสอน_แจกนักเรียน")
media_data = []
if os.path.exists(media_dir):
    for f in sorted(os.listdir(media_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            media_data.append(parse_docx(os.path.join(media_dir, f)))

# 5. Parse เครื่องมือวัดผล (8 rubrics)
eval_dir = os.path.join(root_dir, "04_เครื่องมือวัดผลและข้อสอบ", "แบบประเมินและรูบริก")
eval_data = []
if os.path.exists(eval_dir):
    for f in sorted(os.listdir(eval_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            eval_data.append(parse_docx(os.path.join(eval_dir, f)))

# 6. Parse เอกสารประกอบ (วก7, PDPA, T-PACE)
supp_dir = os.path.join(root_dir, "04_เครื่องมือวัดผลและข้อสอบ", "เอกสารประกอบ")
supp_data = []
if os.path.exists(supp_dir):
    for f in sorted(os.listdir(supp_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            supp_data.append(parse_docx(os.path.join(supp_dir, f)))

# 7. Parse แบบทดสอบและพิมพ์เขียว (5 files)
test_dir = os.path.join(root_dir, "04_เครื่องมือวัดผลและข้อสอบ", "แบบทดสอบท้ายหน่วยและพิมพ์เขียว")
test_data = []
if os.path.exists(test_dir):
    for f in sorted(os.listdir(test_dir)):
        if f.endswith(".docx") and not f.startswith("~$"):
            test_data.append(parse_docx(os.path.join(test_dir, f)))

full_package = {
    "wk3": wk3_data,
    "wk4": wk4_data,
    "wk5": wk5_data,
    "manuals": manuals_data,
    "media": media_data,
    "eval": eval_data,
    "supp": supp_data,
    "tests": test_data
}

output_path = "m1_full_curriculum_extracted.json"
with open(output_path, "w", encoding="utf-8") as f:
    json.dump(full_package, f, ensure_ascii=False, indent=2)

total_docs = (1 if wk3_data else 0) + len(wk4_data) + len(wk5_data) + len(manuals_data) + len(media_data) + len(eval_data) + len(supp_data) + len(test_data)
print(f"Full curriculum parsed ({total_docs} .docx files) and saved to {output_path}. Size: {os.path.getsize(output_path):,} bytes")
