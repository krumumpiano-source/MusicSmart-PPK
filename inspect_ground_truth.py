import os
import sys
import docx
import re

sys.stdout.reconfigure(encoding='utf-8')
base = r'D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101'

print("=== 1. วก3 โครงสร้างรายวิชา ===")
doc3_path = os.path.join(base, 'วก3', 'วก3_โครงสร้างรายวิชา_ศ21101_ดนตรี1.docx')
doc3 = docx.Document(doc3_path)
print("Paragraphs in วก3 (first 10):")
for p in doc3.paragraphs[:10]:
    if p.text.strip():
        print(" ", p.text.strip())

print("\nTables in วก3:")
for t_idx, table in enumerate(doc3.tables):
    print(f"Table {t_idx+1}: {len(table.rows)} rows, {len(table.columns)} cols")
    for r in table.rows:
        row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
        if any('หน่วยที่' in c for c in row_txt) or any('รวม' in c for c in row_txt):
            print("  Row:", row_txt[:4])

print("\n=== 2. วก4 หน่วยการเรียนรู้ 4 หน่วย ===")
for i in range(1, 5):
    f_match = [f for f in os.listdir(os.path.join(base, 'วก4')) if f'หน่วยที่_{i}' in f]
    if f_match:
        fn = f_match[0]
        d = docx.Document(os.path.join(base, 'วก4', fn))
        first_few = [p.text.strip() for p in d.paragraphs[:5] if p.text.strip()]
        print(f"File: {fn}")
        print("  Headings:", first_few[:2])

print("\n=== 3. วก5 แผนการจัดการเรียนรู้ 18 แผน ===")
for i in range(1, 19):
    fn = f"วก5_แผนการจัดการเรียนรู้_แผนที่_{i:02d}.docx"
    fp = os.path.join(base, 'วก5', fn)
    if os.path.exists(fp):
        d = docx.Document(fp)
        first_few = [p.text.strip() for p in d.paragraphs[:6] if p.text.strip()]
        print(f"Plan {i:02d} ({fn}):")
        for line in first_few[:3]:
            print("   ", line)
