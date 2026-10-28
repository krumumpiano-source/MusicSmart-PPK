import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pythainlp.tokenize import word_tokenize

# Target Directory
target_dir = r"D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101\สื่อการสอน"
os.makedirs(target_dir, exist_ok=True)

ZWSP = '\u200B'

def insert_zwsp(text):
    if not text: return ""
    tokens = word_tokenize(text, engine="newmm")
    return ZWSP.join(tokens)

def set_font(run, font_name="TH Sarabun New", size=16):
    run.font.name = font_name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set('ascii', font_name)
    rFonts.set('hAnsi', font_name)
    rFonts.set('cs', font_name)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(insert_zwsp(text))
        set_font(run, size=20)
        run.bold = True
    else:
        run = p.add_run(insert_zwsp(text))
        set_font(run, size=18)
        run.bold = True

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(insert_zwsp(text))
    set_font(run, size=16)
    run.bold = bold

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(insert_zwsp(text))
    set_font(run, size=16)

def add_table(doc, headers, rows_count, col_widths=None):
    table = doc.add_table(rows=1 + rows_count, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(insert_zwsp(h))
        set_font(run, size=16)
        run.bold = True
    
    # Fill empty rows to apply font
    for row in table.rows[1:]:
        for cell in row.cells:
            p = cell.paragraphs[0]
            run = p.add_run(" ")
            set_font(run, size=16)
    return table

def create_doc(filename, title, content_builder):
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    
    add_heading(doc, title, level=1)
    doc.add_paragraph() # spacer
    
    content_builder(doc)
    
    doc.save(os.path.join(target_dir, filename))

# --- Content Builders for each plan ---

def b_n1_p01(doc):
    add_heading(doc, "ส่วนที่ 1: เช็คลิสต์การดูแลรักษา (Kalimba Care Checklist)", level=2)
    add_bullet(doc, "ทำความสะอาดลิ่มเสียง (Tines) ด้วยผ้าไมโครไฟเบอร์หลังเล่นเสร็จ")
    add_bullet(doc, "เก็บใส่กล่องหรือถุงผ้าเพื่อป้องกันฝุ่นและความชื้น")
    add_bullet(doc, "หลีกเลี่ยงการวางไว้ในที่ที่มีอุณหภูมิร้อนจัดหรือโดนแสงแดดโดยตรง")
    doc.add_paragraph()
    add_heading(doc, "ส่วนที่ 2: บันทึกการเทียบเสียง (Tuning Log)", level=2)
    add_paragraph(doc, "คำชี้แจง: ให้นักเรียนใช้แอปพลิเคชันเทียบเสียง (Tuner) ตรวจสอบเสียงแต่ละลิ่มและทำเครื่องหมาย")
    add_table(doc, ["ลิ่มเสียง (คีย์)", "สถานะเสียงก่อนจูน (สูงไป/ต่ำไป/ตรง)", "การปรับแก้ (เคาะขึ้น/เคาะลง)", "ผู้ตรวจสอบ (Peer Check)"], 5)

def b_n1_p02(doc):
    add_paragraph(doc, "คำชี้แจง: ให้นักเรียนแปลงสัญลักษณ์โน้ตสากลในโจทย์ให้กลายเป็น Tablature ตัวเลขสำหรับคาลิมบ้า พร้อมระบุตำแหน่งมือ (ซ้าย/ขวา)")
    add_table(doc, ["ห้องเพลงที่", "โน้ตสากล (ระบบ CDEFG)", "Tablature (ระบบตัวเลข)", "ตำแหน่งนิ้ว (ซ้าย/ขวา)"], 6)
    doc.add_paragraph()
    add_heading(doc, "ความท้าทาย (Challenge):", level=2)
    add_paragraph(doc, "จงแต่งทำนองสั้นๆ ความยาว 4 ห้องเพลง โดยใช้โน้ตอย่างน้อย 5 เสียง และเขียนเป็น Tablature ลงในช่องว่างด้านล่าง")
    add_table(doc, ["ห้องที่ 1", "ห้องที่ 2", "ห้องที่ 3", "ห้องที่ 4"], 2)

def b_n1_p03(doc):
    add_paragraph(doc, "ภารกิจ: ฝึกซ้อมไล่สเกล (C Major) และเล่นแพทเทิร์นจังหวะตามที่กำหนดให้ผ่านด่าน โดยต้องได้รับการลงนามรับรองจากเพื่อนคู่ซ้อม (Peer Coaching)")
    add_table(doc, ["ด่านที่ (Level)", "ภารกิจ (Mission)", "ความเร็ว (BPM)", "ลายเซ็นเพื่อนผู้ประเมิน", "ข้อเสนอแนะจากเพื่อน"], 4)

def b_n1_p04(doc):
    add_paragraph(doc, "คำชี้แจง: หลังจากบันทึกวิดีโอการแสดงเดี่ยวของตนเองแล้ว ให้นักเรียนรับชมและวิเคราะห์ข้อบกพร่องตามความเป็นจริง")
    add_heading(doc, "1. วิเคราะห์ตนเอง (Self-Reflection)", level=2)
    add_table(doc, ["ประเด็น", "สิ่งที่ทำได้ดีแล้ว", "สิ่งที่ต้องปรับปรุงด่วน", "แนวทางแก้ไขในการซ้อมครั้งต่อไป"], 3)
    doc.add_paragraph()
    add_paragraph(doc, "เป้าหมายในการพัฒนาตนเองในสัปดาห์หน้า: ______________________________________________________")

def b_n2_p05(doc):
    add_paragraph(doc, "ชื่อวง: _______________________ เพลงที่ใช้บรรเลง: _______________________")
    add_heading(doc, "1. การแบ่งพาร์ท (Part Assignment)", level=2)
    add_table(doc, ["ชื่อ-สกุลสมาชิก", "หน้าที่/พาร์ท (เช่น ทำนองหลัก, คอร์ด, เบส)", "ความรับผิดชอบหลัก"], 4)
    doc.add_paragraph()
    add_heading(doc, "2. เป้าหมายการซ้อมย่อย (Sectional Goals)", level=2)
    add_table(doc, ["เป้าหมายการซ้อมวันนี้", "ผลการซ้อม (บรรลุ/ไม่บรรลุ)", "ปัญหาที่พบ", "วิธีแก้ปัญหา"], 3)

def b_n2_p06(doc):
    add_paragraph(doc, "ภารกิจ: นำสิ่งของเหลือใช้มาประดิษฐ์เป็นเครื่องประกอบจังหวะ และออกแบบแพทเทิร์นจังหวะเพื่อใช้ร่วมกับวงคาลิมบ้า")
    add_heading(doc, "1. ออกแบบเสียง (Sound Design)", level=2)
    add_paragraph(doc, "วัสดุที่ใช้: ___________________________ ลักษณะเสียงที่ได้ (เช่น ทุ้ม, แหลม, สั้น): ___________________________")
    doc.add_paragraph()
    add_heading(doc, "2. ออกแบบแพทเทิร์น (Rhythm Pattern)", level=2)
    add_table(doc, ["จังหวะตก (Downbeat)", "จังหวะยก (Upbeat)", "เครื่องประกอบจังหวะที่ 1", "เครื่องประกอบจังหวะที่ 2"], 4)

def b_n2_p07(doc):
    add_paragraph(doc, "บันทึกการซ้อมกลุ่มประจำสัปดาห์ เพื่อเตรียมความพร้อมสำหรับการรวมวงใหญ่")
    add_table(doc, ["วันที่ซ้อม", "ท่อนเพลงที่ซ้อม", "ความผิดพลาดทางเทคนิคที่พบ (เช่น คร่อมจังหวะ)", "กลยุทธ์ที่ใช้แก้ไข (เช่น เปิดเครื่องเคาะจังหวะ)"], 3)

def b_n2_p08(doc):
    add_paragraph(doc, "ภารกิจ: เพิ่มความเร็ว (Tempo) ของวงทีละระดับ โดยยังคงความแม่นยำและความกลมกลืนของเสียงไว้")
    add_table(doc, ["ความเร็วเป้าหมาย (BPM)", "ผลลัพธ์ (ผ่าน/ไม่ผ่าน)", "จุดที่วงเกิดอาการแกว่ง/ล่ม", "ใครคือคนที่ช่วยดึงจังหวะกลับมาได้?"], 4)

def b_n2_p09(doc):
    add_paragraph(doc, "คำชี้แจง: ให้หัวหน้าวงหรือตัวแทนผู้ฟัง ประเมินสมดุลของวงใหญ่ (Full Ensemble)")
    add_table(doc, ["จุดเช็ค (Checkpoints)", "ระดับความสมดุล (ดี/ปรับปรุง)", "ข้อสังเกตเพิ่มเติม (เช่น เสียงทำนองเบาไป, คอร์ดดังกลบ)"], 4)

def b_n2_p10(doc):
    add_paragraph(doc, "คำชี้แจง: จงใช้ความรู้ทางดนตรีเพื่อวิเคราะห์และให้คำแนะนำแก่วงเพื่อนอย่างสร้างสรรค์ (Constructive Feedback)")
    add_heading(doc, "ส่วนที่ 1: การวิเคราะห์โครงสร้างและการรวมวง", level=2)
    add_bullet(doc, "จุดเด่นที่เห็นได้ชัดที่สุดในเรื่องของความกลมกลืน (Ensemble) คืออะไร?")
    add_bullet(doc, "มีพาร์ทไหนที่เสียงโดดออกมาหรือจมหายไปหรือไม่? (Balance & Blend)")
    doc.add_paragraph()
    add_heading(doc, "ส่วนที่ 2: ข้อเสนอแนะเชิงปฏิบัติ (Actionable Feedback)", level=2)
    add_table(doc, ["ปัญหาที่ตรวจพบ (ระบุท่อน/จุด)", "สาเหตุของปัญหาทางเทคนิค", "วิธีแก้ไข/คำแนะนำในการซ้อม"], 3)

def b_n2_p11(doc):
    add_paragraph(doc, "หลังจากเสร็จสิ้นการแสดง Mini-Recital ให้นักเรียนทบทวนประสบการณ์และสะท้อนคิด")
    add_table(doc, ["คำถามสะท้อนคิด (Reflection Questions)", "คำตอบของนักเรียน"], 4)

def b_n3_p12(doc):
    add_paragraph(doc, "คำชี้แจง: ฟังคลิปเสียงเครื่องดนตรี แล้วระบุประเภท พร้อมให้เหตุผลสนับสนุนตามหลักนิเวศวิทยาดนตรี")
    add_table(doc, ["ข้อที่", "จัดอยู่ในกลุ่มใด (เช่น เครื่องสายตะวันตก, ปี่พาทย์)", "ลักษณะเสียงที่ได้ยิน (Tone Color)", "เหตุผลที่จำแนกในกลุ่มนี้"], 5)

def b_n3_p13(doc):
    add_paragraph(doc, "ภารกิจ: สร้างแผนผังความคิด (Mind Mapping) จัดหมวดหมู่เครื่องดนตรีข้ามวัฒนธรรมที่มีบทบาทคล้ายกัน")
    add_table(doc, ["บทบาทในวง", "เครื่องดนตรีไทย", "เครื่องดนตรีสากล", "เครื่องดนตรีพื้นบ้าน (พะเยา/ล้านนา)", "จุดร่วมทางสุนทรียภาพ"], 4)

def b_n3_p14(doc):
    add_paragraph(doc, "โครงร่างสำหรับเตรียมการอัด Podcast ความยาว 3-5 นาที")
    add_table(doc, ["ส่วนประกอบ", "เนื้อหา (Script)", "เสียงประกอบ/BGM (Cue)", "ผู้รับผิดชอบเสียงพูด"], 4)

def b_n3_p15(doc):
    add_paragraph(doc, "เช็คลิสต์ตรวจสอบคุณภาพก่อนส่งงาน Podcast (Quality Assurance)")
    add_table(doc, ["รายการตรวจสอบ", "ผ่าน", "ต้องแก้", "ผู้ตรวจสอบระบุปัญหา"], 4)

def b_n4_p16(doc):
    add_paragraph(doc, "คำชี้แจง: ฟังบทเพลงที่กำหนด แล้วบันทึกการเปลี่ยนแปลงของอารมณ์เพลงตามช่วงเวลา (Timeline)")
    add_table(doc, ["เวลา (นาที:วินาที)", "อารมณ์ที่รู้สึก (เช่น ตื่นเต้น, เหงา)", "องค์ประกอบดนตรีที่ทำให้เกิดอารมณ์นั้น (เช่น จังหวะเร็วขึ้น, เครื่องเป่าเข้า)"], 5)

def b_n4_p17(doc):
    add_paragraph(doc, "การวิจารณ์เชิงสุนทรียศาสตร์ตามหลักการของ Feldman (Feldman's Method of Art Criticism)")
    add_table(doc, ["ขั้นตอน", "คำถามนำการคิด (Guiding Questions)", "พื้นที่สำหรับเขียนร่าง (Drafting)"], 4)

def b_n4_p18(doc):
    add_paragraph(doc, "ภารกิจ: สร้างกราฟอารมณ์เพลงและเลือกเพลงลงใน Digital Playlist ให้สอดคล้องกับกราฟที่ออกแบบไว้")
    add_heading(doc, "1. ธีมของ Playlist (Theme): _________________________________", level=2)
    doc.add_paragraph()
    add_heading(doc, "2. ออกแบบกราฟอารมณ์และรายชื่อเพลง", level=2)
    add_table(doc, ["ลำดับ", "ชื่อเพลง / ศิลปิน", "จังหวะ (Tempo) / โทนเสียง", "ตำแหน่งในกราฟอารมณ์ (เช่น Intro, Build-up, Climax)", "เหตุผลทฤษฎีดนตรีที่ร้อยเรียงเพลงนี้ต่อกัน"], 5)


files = [
    ("สื่อ_น1_ผ01_คู่มือจูนเสียงและดูแลรักษาคาลิมบ้า.docx", "ใบกิจกรรม: คู่มือจูนเสียงและดูแลรักษาคาลิมบ้า", b_n1_p01),
    ("สื่อ_น1_ผ02_ใบตารางฝึกอ่านแท็บคาลิมบ้าขั้นพื้นฐาน.docx", "ใบกิจกรรม: ถอดรหัส Tablature คาลิมบ้า", b_n1_p02),
    ("สื่อ_น1_ผ03_Challengeไล่สเกลและจังหวะ.docx", "ใบกิจกรรม: Challenge Mission ไล่สเกลคาลิมบ้า", b_n1_p03),
    ("สื่อ_น1_ผ04_ฟอร์มวิเคราะห์ข้อบกพร่องการปฏิบัติเดี่ยว.docx", "ใบสะท้อนผล: การวิเคราะห์ข้อบกพร่องทางเทคนิค (Self-Reflection)", b_n1_p04),
    
    ("สื่อ_น2_ผ05_แผนผังแบ่งพาร์ทรวมวงและเป้าหมายซ้อมย่อย.docx", "ใบกิจกรรม: บริหารจัดการวงและเป้าหมายการซ้อม (Sectional Goals)", b_n2_p05),
    ("สื่อ_น2_ผ06_ใบกิจกรรมออกแบบเครื่องประกอบจังหวะ.docx", "ใบกิจกรรม: นักประดิษฐ์จังหวะ (Rhythm & Sound Design)", b_n2_p06),
    ("สื่อ_น2_ผ07_บันทึกการซ้อมและสะท้อนผลกลุ่มย่อย.docx", "ใบกิจกรรม: บันทึกการซ้อมและแก้ปัญหาทางดนตรี", b_n2_p07),
    ("สื่อ_น2_ผ08_ความท้าทายเพิ่มความเร็วTempo.docx", "ใบกิจกรรม: Speed Challenge (ทลายขีดจำกัด Tempo)", b_n2_p08),
    ("สื่อ_น2_ผ09_เช็คลิสต์สมดุลเสียงวงใหญ่.docx", "ใบประเมิน: เช็คลิสต์สมดุลและความกลมกลืน (Balance & Blend)", b_n2_p09),
    ("สื่อ_น2_ผ10_คู่มือPeerCritiqueการแสดง.docx", "คู่มือการวิจารณ์เพื่อน: Actionable Peer Critique", b_n2_p10),
    ("สื่อ_น2_ผ11_ใบสะท้อนผลการแสดงMiniRecital.docx", "ใบสะท้อนผล: ถอดบทเรียนจากเวที (Post-Performance Reflection)", b_n2_p11),
    
    ("สื่อ_น3_ผ12_กระดาษคำตอบAudioQuizจำแนกเสียง.docx", "ใบกิจกรรม: นักสืบเสียงดนตรี (Audio Quiz & Analysis)", b_n3_p12),
    ("สื่อ_น3_ผ13_ใบงานMappingจัดหมวดหมู่เครื่องดนตรี.docx", "ใบกิจกรรม: แผนผังเชื่อมโยงสุนทรียภาพเครื่องดนตรีข้ามวัฒนธรรม", b_n3_p13),
    ("สื่อ_น3_ผ14_โครงร่างสคริปต์Podcastข้ามวัฒนธรรม.docx", "ใบกิจกรรม: ออกแบบโครงร่างสคริปต์ Podcast (Script Outline)", b_n3_p14),
    ("สื่อ_น3_ผ15_เช็คลิสต์ตรวจสอบคุณภาพPodcastก่อนเผยแพร่.docx", "ใบประเมิน: QA Checklist ตรวจสอบคุณภาพ Podcast", b_n3_p15),
    
    ("สื่อ_น4_ผ16_ใบวิเคราะห์อารมณ์เพลงRealTime.docx", "ใบกิจกรรม: ถอดรหัสอารมณ์เพลง (Real-Time Emotion Mapping)", b_n4_p16),
    ("สื่อ_น4_ผ17_โครงร่างการเขียนบทวิจารณ์Feldman.docx", "ใบกิจกรรม: ร่างบทวิจารณ์สุนทรียศาสตร์ (Feldman's Method)", b_n4_p17),
    ("สื่อ_น4_ผ18_แบบฟอร์มออกแบบกราฟอารมณ์DigitalPlaylist.docx", "ใบกิจกรรม: สถาปนิกออกแบบอารมณ์ (Digital Playlist Curation)", b_n4_p18)
]

for filename, title, builder in files:
    create_doc(filename, title, builder)
