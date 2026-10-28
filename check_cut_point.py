import docx, os, glob, sys
sys.stdout.reconfigure(encoding='utf-8')

for p_num in [1, 2, 4, 11, 18]:
    fn = f'เอกสารแผนการสอน/ม.1 ดนตรี1 ศ21101/วก5/วก5_แผนการจัดการเรียนรู้_แผนที่_{p_num:02d}.docx'
    doc = docx.Document(fn)
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t == 'ภาคผนวก':
            print(f'{os.path.basename(fn)}: Appendix starts at paragraph {i}')
            for j in range(max(0, i-3), i):
                print(f'   Before: {j}: {doc.paragraphs[j].text[:60]}')
            break
