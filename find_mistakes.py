import os
import docx

folder = r"d:\AI CURSER\PLANE MUSIC ED"
docx_files = [f for f in os.listdir(folder) if f.endswith(".docx")]

with open(os.path.join(folder, "mistakes.txt"), "w", encoding="utf-8") as out:
    for f in docx_files:
        path = os.path.join(folder, f)
        try:
            doc = docx.Document(path)
            found = False
            for p in doc.paragraphs:
                text = p.text.replace('\u200b', '')
                if "ตั้งสาย" in text:
                    if not found:
                        out.write(f"--- FILE: {f} ---\n")
                        found = True
                    out.write(f"PARAGRAPH: {text.strip()}\n")
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.replace('\u200b', '')
                        if "ตั้งสาย" in text:
                            if not found:
                                out.write(f"--- FILE: {f} ---\n")
                                found = True
                            out.write(f"TABLE CELL: {text.strip().replace(chr(10), ' ')}\n")
        except Exception as e:
            out.write(f"Error reading {f}: {e}\n")
