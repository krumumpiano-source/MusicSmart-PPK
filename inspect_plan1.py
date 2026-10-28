import os
import docx

root = r"D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101"
p = os.path.join(root, "แผนการสอนพร้อมใช้", "แผนการจัดการเรียนรู้_แผนที่_01.docx")
doc = docx.Document(p)

lines = []
lines.append("=== Paragraphs in แผนการสอนพร้อมใช้ (แผนที่ 01) ===")
for i, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    if t and len(t) < 100:
        lines.append(f"[{i}] {t}")

lines.append("\n=== Tables in แผนการสอนพร้อมใช้ (แผนที่ 01) ===")
for idx, tbl in enumerate(doc.tables):
    lines.append(f"\n-- Table {idx+1} ({len(tbl.rows)} rows, {len(tbl.columns)} cols) --")
    for r in tbl.rows[:3]:
        lines.append("  | " + " | ".join(c.text.strip().replace("\n", " ") for c in r.cells))

with open("scratch_plan1_full.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print("Plan 1 inspected.")
