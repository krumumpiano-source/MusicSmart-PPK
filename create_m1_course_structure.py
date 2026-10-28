import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pythainlp.tokenize import word_tokenize
from docx.oxml.ns import qn
import os

def insert_zwsp(text):
    if not text: return text
    words = word_tokenize(text, engine='newmm')
    return '\u200b'.join(words)

doc = docx.Document()

# Set margins
sections = doc.sections
for section in sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'TH SarabunPSK'
font.size = Pt(16)

def add_para(text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(insert_zwsp(text))
    run.font.name = 'TH SarabunPSK'
    run.font.size = Pt(16)
    run.bold = bold
    return p

# Header
add_para("คำอธิบายรายวิชา", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para("รหัสวิชา ศ21101 ชื่อวิชา ดนตรี 1 กลุ่มสาระการเรียนรู้ ศิลปะ")
add_para("ชั้นมัธยมศึกษาปีที่ 1 ภาคเรียนที่ 1 เวลา 20 ชั่วโมง จำนวน 0.5 หน่วยกิต")
add_para("")

# Course Description
add_para("คำอธิบายรายวิชา", bold=True)
desc_k = "ศึกษาและปฏิบัติเกี่ยวกับการอ่าน เขียน และร้องโน้ตสากลเบื้องต้น การเปรียบเทียบเสียงร้องและเสียงของเครื่องดนตรีต่างวัฒนธรรม การใช้และบำรุงรักษาเครื่องดนตรีคาลิมบ้าอย่างระมัดระวังและปลอดภัย ปฏิบัติการร้องเพลงและบรรเลงคาลิมบ้าประกอบเพลงเดี่ยวและรวมวงที่หลากหลายรูปแบบ วิเคราะห์จัดประเภทของวงดนตรีไทยและสากล อธิบายบทบาทความสัมพันธ์และอิทธิพลของดนตรีที่มีต่อสังคมไทย รวมถึงระบุความหลากหลายขององค์ประกอบดนตรีในวัฒนธรรมต่างกัน"
add_para("     " + desc_k)

desc_p = "โดยใช้กระบวนการทางดนตรี การฝึกปฏิบัติ การสืบเสาะหาความรู้ การวิเคราะห์ การวิพากษ์วิจารณ์ และการทำงานกลุ่ม เพื่อให้เกิดทักษะและความเข้าใจที่ลึกซึ้ง"
add_para("     " + desc_p)

desc_a = "เพื่อให้ผู้เรียนเกิดสุนทรียภาพ เห็นคุณค่าของดนตรี สามารถแสดงความคิดเห็นต่ออารมณ์เพลง เปรียบเทียบความรู้สึกในการฟัง นำเสนอและวิจารณ์เพลงที่ชื่นชอบ ใช้เกณฑ์ประเมินคุณภาพงานดนตรีได้อย่างสร้างสรรค์ ตลอดจนนำไปประยุกต์ใช้ในชีวิตประจำวัน ปฏิบัติตนตามหลักปรัชญาเศรษฐกิจพอเพียงและหลักวิถีพุทธ"
add_para("     " + desc_a)
add_para("")

# Indicators
add_para("รหัสตัวชี้วัด", bold=True)
indicators = [
    "ศ 2.1 ม.1/1, ม.1/2, ม.1/3, ม.1/4, ม.1/5, ม.1/6, ม.1/7, ม.1/8, ม.1/9",
    "ศ 2.2 ม.1/1, ม.1/2"
]
for ind in indicators:
    add_para("     " + ind)
add_para("รวมทั้งหมด 11 ตัวชี้วัด", bold=True)
doc.add_page_break()

# Course Structure
add_para("โครงสร้างรายวิชา ดนตรี 1 (รหัสวิชา ศ21101)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para("")

table = doc.add_table(rows=1, cols=6)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
headers = ["ลำดับที่", "ชื่อหน่วยการเรียนรู้", "มาตรฐาน/ตัวชี้วัด", "สาระสำคัญ", "เวลา(ชั่วโมง)", "น้ำหนักคะแนน"]
for i, header in enumerate(headers):
    run = hdr_cells[i].paragraphs[0].add_run(insert_zwsp(header))
    run.font.name = 'TH SarabunPSK'
    run.font.size = Pt(16)
    run.bold = True
    hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

data = [
    ("1", "พื้นฐานคาลิมบ้า", "ศ 2.1 ม.1/1, 1/9", "การอ่าน เขียน ร้องโน้ตสากลเบื้องต้น และการดูแลรักษาเครื่องดนตรีคาลิมบ้าอย่างถูกวิธี", "4", "20"),
    ("2", "ปฏิบัติรวมวงคาลิมบา", "ศ 2.1 ม.1/3", "การบรรเลงเครื่องดนตรีคาลิมบ้าประกอบการร้องเพลงและการเล่นรวมวงในบทเพลงที่หลากหลาย", "7", "20"),
    ("3", "รอบรู้วงดนตรีและวัฒนธรรม", "ศ 2.1 ม.1/2, 1/4\nศ 2.2 ม.1/1, 1/2", "การจัดประเภทวงดนตรีไทยและสากล การเปรียบเทียบเสียงเครื่องดนตรีต่างวัฒนธรรม และอิทธิพลดนตรีต่อสังคม", "4", "20"),
    ("4", "สุนทรียภาพแห่งการฟัง", "ศ 2.1 ม.1/5, 1/6, 1/7, 1/8", "การฟังและวิจารณ์อารมณ์เพลง การนำเสนอเพลงที่ชื่นชอบ และการใช้เกณฑ์ประเมินคุณภาพงานดนตรี", "3", "20"),
    ("", "สอบกลางภาค", "ทดสอบหน่วยที่ 1-2", "", "1", "10"),
    ("", "สอบปลายภาค", "ทดสอบหน่วยที่ 3-4", "", "1", "10"),
    ("", "รวม", "รวม 11 ตัวชี้วัด", "", "20", "100")
]

for row_data in data:
    row_cells = table.add_row().cells
    for i, item in enumerate(row_data):
        p = row_cells[i].paragraphs[0]
        run = p.add_run(insert_zwsp(item))
        run.font.name = 'TH SarabunPSK'
        run.font.size = Pt(16)
        if i in [0, 4, 5]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

output_path = os.path.join(os.getcwd(), 'วก_3-รายวิชาดนตรี_ม1.docx')
doc.save(output_path)
print("Document generated successfully at:", output_path)
