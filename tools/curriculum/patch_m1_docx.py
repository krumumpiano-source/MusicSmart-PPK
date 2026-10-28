# -*- coding: utf-8 -*-
"""patch_m1_docx.py — แก้เอกสาร วก.3 / วก.4 / วก.5 ของ ม.1 ศ2110x ให้ตรงกัน (ใช้ python-docx)

ลำดับงาน: วก.3 -> วก.4 -> วก.5  แล้วจึง `python tools/curriculum/extract_wk5_docx.py` และ `node tools/curriculum/build_wk5_data.js`

* สำรองโฟลเดอร์เอกสารทั้งหมดก่อนแก้ (scratch/docx_backup_<เวลา>/)  ถ้าสำรองไม่สำเร็จจะไม่แก้
* บันทึก ก่อน -> หลัง ทุกจุดที่แก้ลง tools/curriculum/logs/patch_m1_docx_<เวลา>.json
* รันซ้ำได้ (idempotent): จุดที่แก้แล้วจะไม่ตรงกฎอีก
* ภาคเรียน/ปีการศึกษา/รหัสวิชา อ่านจาก app/data/term-config.js (แหล่งเดียว) — ไม่ฝังเลขในสคริปต์
  ใช้ --term 2569/2 เพื่อกำหนดเอง  --dry-run เพื่อดูโดยไม่เขียน

ข้อควรรู้ของ Word: ข้อความมี ZWSP (U+200B) แทรกระหว่างคำ และถูกแบ่งเป็นหลาย run — ฟังก์ชัน sub() จับคู่บนข้อความที่ตัด ZWSP แล้ว
แล้วเขียนกลับตามขอบเขต run เดิม (รักษารูปแบบของ run แรกที่ถูกแก้)
"""
import argparse, datetime, glob, json, os, re, shutil, sys
from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
def _m1_dir():
    base = os.path.join(ROOT, 'เอกสารแผนการสอน')
    c = sorted(glob.glob(os.path.join(base, 'ม.1 ดนตรี1 *')))
    return c[0] if c else os.path.join(base, 'ม.1 ดนตรี1')  # ไม่ผูกกับรหัสวิชาในชื่อโฟลเดอร์


DOC_DIR = _m1_dir()
ZW = '\u200b\u200c\u200d\ufeff'

# ชื่อหน่วยตาม วก.3 (ตารางที่ 3 ของโครงสร้างรายวิชา) — แหล่งความจริงของชื่อหน่วย
UNIT_NAMES = {
    1: 'พื้นฐานโน้ตและทักษะปฏิบัติคาลิมบ้า',
    2: 'บทเพลงและการประพันธ์เนื้อร้องสร้างสรรค์',
    3: 'วงดนตรีไทยและวงดนตรีสากล',
    4: 'การวิเคราะห์วิจารณ์และอิทธิพลของดนตรีต่อสังคม',
}
# น้ำหนักคะแนนรายคาบ (K, P, A) — รวมต่อหน่วย 20 + กลางภาค K10 + ปลายภาค K10 = 100
WEIGHTS = {
    3: (1, 2, 1), 4: (1, 2, 1), 5: (1, 2, 1), 6: (1, 2, 1), 7: (0, 4, 0),   # หน่วย 1 = K4/P12/A4
    9: (1, 7, 2), 10: (1, 7, 2),                                              # หน่วย 2 = K2/P14/A4 (+กลางภาค K10)
    13: (2, 5, 3), 15: (2, 5, 3),                                             # หน่วย 3 = K4/P10/A6
    17: (2, 5, 3), 19: (2, 5, 3),                                             # หน่วย 4 = K4/P10/A6 (+ปลายภาค K10)
    11: (10, 0, 0), 20: (10, 0, 0),
}
EXAM_PERIODS = (11, 20)


# ---------------------------------------------------------------- term (อ่านจาก term-config.js)
def load_term_config():
    src = open(os.path.join(ROOT, 'app', 'data', 'term-config.js'), encoding='utf-8').read()
    m = re.search(r'/\*TERM_CONFIG_JSON\*/(.*?)/\*END_TERM_CONFIG_JSON\*/', src, re.S)
    return json.loads(m.group(1))


def resolve_term(cfg, today=None, forced=None):
    if forced:
        y, t = forced.split('/')
        return int(y), int(t)
    today = today or (datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(hours=7)).strftime('%Y-%m-%d')
    terms = sorted(cfg['terms'], key=lambda t: t['start'])
    ly = max(t['year'] for t in terms)
    out = [dict(t) for t in terms]
    for k in range(1, cfg.get('projectAheadYears', 0) + 1):
        for t in [t for t in terms if t['year'] == ly]:
            out.append({'year': t['year'] + k, 'term': t['term'], 'start': str(int(t['start'][:4]) + k) + t['start'][4:]})
    out.sort(key=lambda t: t['start'])
    cur = out[0]
    for t in out:
        if t['start'] <= today:
            cur = t
    return cur['year'], cur['term']


# ---------------------------------------------------------------- helpers
def iter_pars(doc):
    seen = set()

    def from_container(c):
        for p in c.paragraphs:
            if p._p not in seen:
                seen.add(p._p); yield p
        for t in c.tables:
            yield from from_table(t)

    def from_table(t):
        for row in t.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                yield from from_container(cell)

    yield from from_container(doc)
    for s in doc.sections:
        for part in (s.header, s.footer):
            try:
                yield from from_container(part)
            except Exception:
                pass


def nmap(p):
    full = ''.join(r.text for r in p.runs)
    idx = [i for i, ch in enumerate(full) if ch not in ZW]
    return full, ''.join(full[i] for i in idx), idx


def ntext(p):
    return nmap(p)[1]


def replace_span(p, start, end, text):
    pos, first = 0, True
    for r in p.runs:
        t = r.text
        rs, re_ = pos, pos + len(t)
        pos = re_
        if re_ <= start or rs >= end:
            if not (first is True and rs <= start < re_ and start == end):
                continue
        ls, le = max(start, rs) - rs, min(end, re_) - rs
        if first:
            r.text = t[:ls] + text + t[le:]
            first = False
        else:
            r.text = t[:ls] + t[le:]


def sub(p, pattern, repl):
    """แทนที่บนข้อความที่ตัด ZWSP แล้ว คืนจำนวนที่แก้ [(old,new)]"""
    rx = pattern if hasattr(pattern, 'finditer') else re.compile(pattern)
    full, norm, idx = nmap(p)
    ms = list(rx.finditer(norm))
    changes = []
    for m in reversed(ms):
        if m.end() == m.start():
            continue
        new = repl(m) if callable(repl) else m.expand(repl)
        if new == m.group(0):
            continue
        replace_span(p, idx[m.start()], idx[m.end() - 1] + 1, new)
        changes.append((m.group(0), new))
        # idx/ข้อความเปลี่ยนแล้ว ต้องคำนวณใหม่ก่อนจับคู่ถัดไป (ย้อนจากท้ายจึงใช้ idx เดิมได้)
    return list(reversed(changes))


def set_par_text(p, text):
    old = ntext(p)
    if old == text:
        return None
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)
    return old or '(ว่าง)'


def drop_par(p):
    el = p._p
    el.getparent().remove(el)


class Log:
    def __init__(self):
        self.rows = []

    def add(self, doc, where, rule, old, new):
        self.rows.append({'file': doc, 'where': where, 'rule': rule, 'old': old, 'new': new})


# ---------------------------------------------------------------- rule engine
def apply_global(doc, name, log, code, term, year):
    for p in iter_pars(doc):
        for pat, repl, rule in (
            (r'ศ2110\d', code, 'รหัสวิชา → ตามภาคเรียนปัจจุบัน'),
            (r'ภาคเรียนที่\s*[123](?!\d)', f'ภาคเรียนที่ {term}', 'ภาคเรียน → ตาม term-config'),
            (r'ปีการศึกษา\s*25\d\d', f'ปีการศึกษา {year}', 'ปีการศึกษา → ตาม term-config'),
            (r'เข้าไลน์กลุ่มวิชาดนตรี\s*และ\s*', '', 'ตัดเรื่องไลน์'),
            (r'/ไลน์กลุ่ม', '', 'ตัดเรื่องไลน์'),
        ):
            for old, new in sub(p, pat, repl):
                log.add(name, ntext(p)[:60], rule, old, new)


def drop_numbered(doc, name, log, find_rx, renum_rx, group_idx_fmt, log_rule):
    """ลบย่อหน้าที่ตรง find_rx แล้วเรียงเลขย่อหน้าถัดไปที่เป็นรายการเดียวกันใหม่ (เลข-1)"""
    pars = [p for p in iter_pars(doc) if p._p.getparent() is not None]
    for i, p in enumerate(pars):
        t = ntext(p)
        if re.search(find_rx, t):
            log.add(name, t[:60], log_rule, t, '(ลบย่อหน้า)')
            drop_par(p)
            for q in pars[i + 1:]:
                qt = ntext(q)
                m = re.match(renum_rx, qt)
                if not m:
                    break
                for old, new in sub(q, renum_rx, group_idx_fmt(m)):
                    log.add(name, qt[:40], 'เลขลำดับ -1', old, new)
            return True
    return False


def patch_wk3(path, name, log, code, term, year):
    doc = Document(path)
    apply_global(doc, name, log, code, term, year)
    for p in iter_pars(doc):
        t = ntext(p)
        for pat, repl, rule in (
            (r'ระหว่างภาค\s*80\s*:\s*ปลายภาค\s*20', 'ระหว่างภาค 80 : กลางภาค 10 : ปลายภาค 10', 'สัดส่วน 80/10/10'),
            (r'กำหนดที่\s*K\s*:\s*30,?\s*P\s*:\s*50,?\s*A\s*:\s*20',
             'กำหนดที่ประมาณ K : P : A ≈ 34 : 46 : 20 (นับข้อสอบกลางภาคและปลายภาคเป็นด้าน K)', 'K:P:A ตามตารางหน่วย'),
            (r'K\s*:\s*30\s*P\s*:\s*50\s*A\s*:\s*20',
             'K : P : A ≈ 34 : 46 : 20 (นับข้อสอบกลางภาคและปลายภาคเป็นด้าน K)', 'K:P:A ตามตารางหน่วย'),
        ):
            for old, new in sub(p, pat, repl):
                log.add(name, t[:60], rule, old, new)
    doc.save(path) if not ARGS.dry_run else None


def patch_wk4(path, name, log, code, term, year):
    doc = Document(path)
    apply_global(doc, name, log, code, term, year)
    # หน่วยที่ 1: เรื่อง LINE / QR
    rules = [
        (r'การเข้า\s*LINE\s*กลุ่มวิชาดนตรี\s*และการลงทะเบียนในแอปพลิเคชัน', 'การลงทะเบียนและเข้าสู่ระบบในแอปพลิเคชัน', 'ตัด LINE'),
        (r'การเข้ากลุ่ม\s*LINE\s*และการใช้งานแอปพลิเคชัน', 'และการใช้งานแอปพลิเคชัน', 'ตัด LINE'),
        (r'เข้า\s*LINE\s*กลุ่ม\s*แนะนำเมนู', 'แนะนำเมนู', 'ตัด LINE'),
        (r'P:\s*เข้า\s*LINE\s*&\s*แอป\s*MusicSmart', 'P: เข้าใช้แอป MusicSmart', 'ตัด LINE'),
        (r'เข้า\s*LINE\s*กลุ่ม\s*&\s*ลงทะเบียนแอป', 'ลงทะเบียนแอป', 'ตัด LINE'),
        (r'นำเข้า\s*LINE\s*&\s*สอนเมนูแอป', 'สอนเมนูแอป', 'ตัด LINE'),
    ]
    for p in iter_pars(doc):
        t = ntext(p)
        for pat, repl, rule in rules:
            for old, new in sub(p, pat, repl):
                log.add(name, t[:60], rule, old, new)
    drop_numbered(doc, name, log, r'^9\.\d+\s*QR\s*Code\s*LINE', r'^(9\.)(\d+)(\s)',
                  lambda m: (lambda mm: f'{mm.group(1)}{int(mm.group(2)) - 1}{mm.group(3)}'), 'ตัด QR LINE')
    # ย่อหน้า QR ในเซลล์ตาราง (ถ้ามี)
    for p in list(iter_pars(doc)):
        t = ntext(p)
        if re.fullmatch(r'\s*QR\s*Code\s*LINE.*', t):
            log.add(name, t[:60], 'ตัด QR LINE', t, '(ลบย่อหน้า)')
            drop_par(p)
    for p in list(iter_pars(doc)):
        t = ntext(p)
        for old, new in sub(p, r'\s*QR\s*Code\s*LINE(?=\s*คู่มือ)', ''):
            log.add(name, t[:60], 'ตัด QR LINE ในรายการสื่อ', old, new or '(ลบ)')
    leftover = [ntext(p)[:80] for p in iter_pars(doc) if re.search(r'LINE|ไลน์|QR', ntext(p))]
    if leftover:
        log.add(name, '-', 'ยังพบ LINE/QR (ต้องตรวจมือ)', '; '.join(leftover), '')
    doc.save(path) if not ARGS.dry_run else None


def period_of(fn):
    m = re.search(r'ที่_?(\d+)', fn)
    return int(m.group(1))


def patch_wk5(path, name, log, code, term, year):
    n = period_of(os.path.basename(path))
    doc = Document(path)
    apply_global(doc, name, log, code, term, year)

    def rule(pat, repl, why, only=None):
        for p in iter_pars(doc):
            t = ntext(p)
            for old, new in sub(p, pat, repl):
                log.add(name, t[:60], why, old, new)

    # ชื่อหน่วยตาม วก.3 (หัวกระดาษ "ภาคเรียนที่ … หน่วยการเรียนรู้ที่ N เรื่อง …" อยู่หลังตัวแบ่งบรรทัดในย่อหน้าเดียวกับรหัสวิชา)
    hdr = re.compile(r'^ภาคเรียนที่[^\n]*?หน่วยการเรียนรู้ที่\s*(\d)\s*เรื่อง\s*([^\n]+)$', re.M)
    for p in iter_pars(doc):
        t = ntext(p)
        m = hdr.search(t)
        if m:
            u = int(m.group(1))
            pat = re.compile(r'(?<=เรื่อง )\s*' + re.escape(m.group(2)) + r'$', re.M)
            for old, new in sub(p, pat, UNIT_NAMES[u]):
                log.add(name, t[:60].replace('\n', ' '), 'ชื่อหน่วยตาม วก.3', old, new)

    rule(r'\(เก็บระหว่างภาค\s*80\s*\+\s*สอบปลายภาค\s*20\)', '(เก็บระหว่างภาค 80 + สอบกลางภาค 10 + สอบปลายภาค 10)', 'สัดส่วน 80/10/10')
    rule(r'\(K:\s*30,\s*P:\s*50,\s*A:\s*20\)',
         '(เก็บระหว่างภาค 80 + สอบกลางภาค 10 + สอบปลายภาค 10) และสัดส่วนด้าน K:P:A ≈ 34:46:20 (นับข้อสอบเป็น K)', 'K:P:A ตาม วก.3')

    if n == 1:
        rule(r'สแกน QR Code เข้า LINE กลุ่มวิชาดนตรีประจำห้องเรียนเพื่อใช้ติดต่อสื่อสาร,\s*', '', 'ตัดเรื่องไลน์/QR')
        rule(r'นักเรียนสแกน QR Code เพื่อเชื่อมต่อฐานข้อมูลรายบุคคลเข้าสู่ LINE Official และแอปพลิเคชัน MusicSmart',
             'นักเรียนลงทะเบียนและเข้าสู่ระบบแอปพลิเคชัน MusicSmart', 'ตัดเรื่องไลน์/QR')
        rule(r'การเข้ากลุ่ม LINE\s*&\s*แอป', 'การเข้าใช้งานแอป MusicSmart', 'ตัดเรื่องไลน์')
        rule(r'พบความพร้อมสมบูรณ์\s*100%', 'พบความพร้อมร้อยละ ......', 'ข้อความผลที่ยังไม่ยืนยัน → เว้นว่างให้ครูกรอกหลังสอน')
        drop_numbered(doc, name, log, r'^\d+\.\s*QR\s*Code\s*เข้า\s*LINE', r'^(\d+)(\.\s)',
                      lambda m: (lambda mm: f'{int(mm.group(1)) - 1}{mm.group(2)}'), 'ตัดสื่อ QR LINE')
    if n == 7:
        rule(r'พบผู้เรียนผ่านเกณฑ์ระดับดีเยี่ยมร้อยละ\s*90\s*ขึ้นไป', 'พบผู้เรียนผ่านเกณฑ์ระดับ 3 ขึ้นไปร้อยละ ......',
             'ข้อความผลที่ยังไม่ยืนยัน → เว้นว่างให้ครูกรอกหลังสอน')
    if n == 11:
        rule(r'\(คาบ\s*1\s*[-–]\s*10\)', '(คาบ 2-9)', 'ขอบเขตข้อสอบกลางภาค = คาบ 2-9')
    if n == 20:
        rule(r'(ศ\s*2\.1\s*ม\.1/2,\s*)(ม\.1/4)', r'\1ม.1/3, \2', 'เพิ่มตัวชี้วัด ม.1/3 (ตาม วก.3)')
        rule(r'\(การสอบกลางภาค,\s*', '(', 'ตัดคำว่าการสอบกลางภาคออกจากเนื้อหาข้อสอบปลายภาค')

    # ---------------- ตารางเกณฑ์ (tables[0]): น้ำหนัก K/P/A ต่อคาบ + ตัดคะแนนที่ฝังในระดับ
    t0 = doc.tables[0]
    w = WEIGHTS.get(n, (0, 0, 0))
    strip = re.compile(r'\s*\(\s*(?:ต่ำกว่า\s*)?\d+(?:\.\d+)?(?:\s*[-–]\s*\d+(?:\.\d+)?)?\s*คะแนน\s*\)')
    for ri in (1, 2, 3):
        row = t0.rows[ri]
        cells = row.cells
        wt = w[ri - 1]
        want = str(wt) if wt else '-'
        o = set_par_text(cells[1].paragraphs[0], want)
        if o is not None:
            log.add(name, f'ตารางเกณฑ์ แถว {ri} น้ำหนัก', 'น้ำหนักคะแนนตามตารางหน่วย/คาบ', o, want)
        if n in EXAM_PERIODS:
            continue
        for ci in range(2, 6):
            for p in cells[ci].paragraphs:
                t = ntext(p)
                for old, new in sub(p, strip, ''):
                    log.add(name, f'ตารางเกณฑ์ แถว {ri} ระดับ {6 - ci}', 'ตัดคะแนนที่ฝังในระดับ (ระดับ=ผลการประเมิน, คะแนน=น้ำหนัก×ระดับ/4)', old, '')
    if n in EXAM_PERIODS:
        scope = 'คาบ 2-9' if n == 11 else 'คาบ 11-19'
        bands = [
            f'ตอบแบบทดสอบถูกต้อง 16-20 ข้อ (ร้อยละ 80-100 = 8-10 คะแนน) แสดงความรู้ความเข้าใจเนื้อหา{scope}ได้แม่นยำครบถ้วน',
            f'ตอบแบบทดสอบถูกต้อง 14-15 ข้อ (ร้อยละ 70-79 = 7-7.5 คะแนน) มีความรู้ความเข้าใจในระดับดี',
            f'ตอบแบบทดสอบถูกต้อง 10-13 ข้อ (ร้อยละ 50-69 = 5-6.5 คะแนน) ผ่านเกณฑ์ขั้นต่ำ',
            f'ตอบแบบทดสอบถูกต้องน้อยกว่า 10 ข้อ (ต่ำกว่าร้อยละ 50 = ต่ำกว่า 5 คะแนน) ต้องได้รับการสอนเสริมและสอบแก้ตัว',
        ]
        for k, ci in enumerate(range(2, 6)):
            o = set_par_text(t0.rows[1].cells[ci].paragraphs[0], bands[k])
            if o is not None:
                log.add(name, f'ตารางเกณฑ์ แถว 1 ระดับ {4 - k}', 'ข้อสอบ: ระดับตามจำนวนข้อที่ถูก', o, bands[k])
        for ri in (2, 3):
            for ci in range(2, 6):
                o = set_par_text(t0.rows[ri].cells[ci].paragraphs[0], 'ไม่ประเมินในคาบสอบ (ข้อสอบวัดด้านความรู้ K)')
                if o is not None:
                    log.add(name, f'ตารางเกณฑ์ แถว {ri} ระดับ {6 - ci}', 'คาบสอบไม่ประเมิน P/A', o, 'ไม่ประเมินในคาบสอบ (ข้อสอบวัดด้านความรู้ K)')
    doc.save(path) if not ARGS.dry_run else None


def main():
    global ARGS
    ap = argparse.ArgumentParser()
    ap.add_argument('--term', help='เช่น 2569/2 (ค่าเริ่มต้น: คำนวณจาก term-config.js และวันนี้)')
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--only', choices=['wk3', 'wk4', 'wk5'])
    ARGS = ap.parse_args()
    cfg = load_term_config()
    year, term = resolve_term(cfg, forced=ARGS.term)
    prefix = cfg['subjectPrefixByGrade']['ม.1']
    code = f'{prefix}{term}'
    print(f'ภาคเรียน {term}/{year}  รหัสวิชา {code}')

    stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    if not ARGS.dry_run:
        bk = os.path.join(ROOT, 'scratch', f'docx_backup_{stamp}')
        shutil.copytree(DOC_DIR, bk)
        print('สำรองไว้ที่', bk)

    log = Log()
    jobs = []
    for dp, _, fs in os.walk(DOC_DIR):
        for f in sorted(fs):
            if not f.endswith('.docx') or f.startswith('~$'):
                continue
            full = os.path.join(dp, f)
            if f.startswith('วก3'):
                jobs.append(('wk3', full, f, patch_wk3))
            elif f.startswith('วก4'):
                jobs.append(('wk4', full, f, patch_wk4))
            elif f.startswith('วก5'):
                jobs.append(('wk5', full, f, patch_wk5))
    order = {'wk3': 0, 'wk4': 1, 'wk5': 2}
    jobs.sort(key=lambda j: (order[j[0]], j[2]))
    for kind, full, f, fn in jobs:
        if ARGS.only and ARGS.only != kind:
            continue
        before = len(log.rows)
        fn(full, f, log, code, term, year)
        print(f'{kind} {f}: แก้ {len(log.rows) - before} จุด')

    os.makedirs(os.path.join(ROOT, 'tools', 'curriculum', 'logs'), exist_ok=True)
    lp = os.path.join(ROOT, 'tools', 'curriculum', 'logs', f'patch_m1_docx_{stamp}{"_dry" if ARGS.dry_run else ""}.json')
    json.dump({'term': term, 'year': year, 'code': code, 'changes': log.rows}, open(lp, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('บันทึกก่อน→หลังที่', lp, f'(รวม {len(log.rows)} จุด)')


if __name__ == '__main__':
    main()
