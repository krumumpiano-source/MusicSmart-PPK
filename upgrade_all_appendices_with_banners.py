import os
import sys
import docx
from docx.oxml.ns import qn
from docxcompose.composer import Composer

sys.stdout.reconfigure(encoding='utf-8')

base = r'D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101'
media_dir = os.path.join(base, 'สื่อการสอน')
eval_dir = os.path.join(base, 'เครื่องมือวัดผล')
test_dir = os.path.join(base, 'แบบทดสอบและพิมพ์เขียว')
supp_dir = os.path.join(base, 'เอกสารประกอบ')

units_info = {
    1: {'title': 'พื้นฐานคาลิมบ้า', 'plans': range(1, 5)},
    2: {'title': 'ปฏิบัติรวมวงคาลิมบ้า', 'plans': range(5, 12)},
    3: {'title': 'รอบรู้วงดนตรีและวัฒนธรรม', 'plans': range(12, 16)},
    4: {'title': 'สุนทรียภาพแห่งการฟัง', 'plans': range(16, 19)}
}

def get_unit_info(plan_num):
    for u_num, data in units_info.items():
        if plan_num in data['plans']:
            return u_num, data['title']
    return 1, 'พื้นฐานคาลิมบ้า'

plan_topics = {
    1: 'กายวิภาคและการก้องกังวาน (Acoustics & Posture)',
    2: 'รหัสนามธรรมสู่ระบบเสียงและการขับร้อง (Tab, 5-Line Staff & Solfege Singing)',
    3: 'จังหวะ ทิศทางเสียง และการขับร้องบันไดเสียง (C Major Scale, Singing & Metronome)',
    4: 'การบรรเลงเดี่ยวเชิงประยุกต์และการขับร้องประกอบ (Solo Performance, Singing & Care)',
    5: 'พลวัตของการรวมวง (Ensemble Dynamics)',
    6: 'นวัตกรรมเครื่องประกอบจังหวะ (Rhythm Invention & Upcycling)',
    7: 'การกำกับตนเองในการซ้อมกลุ่มย่อย (Sectional Rehearsal 1)',
    8: 'การกำกับตนเองในการซ้อมกลุ่มย่อย (Sectional Rehearsal 2 & Tempo Mastery)',
    9: 'สมดุลแห่งเสียง (Balance & Run-through)',
    10: 'Mini-Recital: ประเมินเพื่อพัฒนา ตอนที่ 1 (กลุ่ม 1-4)',
    11: 'Mini-Recital: ประเมินเพื่อพัฒนา ตอนที่ 2 (กลุ่ม 5-8 และสรุปบทเรียน)',
    12: 'การจัดกลุ่มเสียงวงดนตรีไทย (Thai Ensemble Classification)',
    13: 'สุนทรียะและวิวัฒนาการดนตรีสากล (Western Ensemble & Evolution)',
    14: 'การสื่อสารข้ามวัฒนธรรม (Cross-Cultural Scriptwriting)',
    15: 'นวัตกรรมสื่อเสียง (Podcast Production & Digital Audio Guide)',
    16: 'ถอดรหัสอารมณ์เพลง (Music & Emotion Analytics)',
    17: "สุนทรียศาสตร์และการวิจารณ์ (Feldman's Art Criticism Method)",
    18: 'การรังสรรค์เพลย์ลิสต์ส่วนบุคคล (Digital Playlist Curation & Reflection)'
}

units_data_wk4 = {
    1: {
        'title': 'พื้นฐานคาลิมบ้า',
        'plans': [1, 2, 3, 4],
        'eval_ids': [1, 3],
        'test_fn': 'แบบทดสอบท้ายหน่วย_หน่วยที่_1_พื้นฐานคาลิมบ้า.docx',
        'supp_fns': ['เอกสารวิชาการ_กรอบแนวคิดการจัดการเรียนรู้_T_PACE_Framework.docx']
    },
    2: {
        'title': 'ปฏิบัติรวมวงคาลิมบ้า',
        'plans': [5, 6, 7, 8, 9, 10, 11],
        'eval_ids': [2, 4, 3],
        'test_fn': 'แบบทดสอบท้ายหน่วย_หน่วยที่_2_ปฏิบัติรวมวงคาลิมบ้า.docx',
        'supp_fns': ['วก7_รายงานการวิจัยในชั้นเรียนอย่างย่อ_ศ21101.docx', 'หนังสือยินยอมการคุ้มครองข้อมูลส่วนบุคคล_PDPA_Consent_Form.docx']
    },
    3: {
        'title': 'รอบรู้วงดนตรีและวัฒนธรรม',
        'plans': [12, 13, 14, 15],
        'eval_ids': [5, 8],
        'test_fn': 'แบบทดสอบท้ายหน่วย_หน่วยที่_3_รอบรู้วงดนตรีและวัฒนธรรม.docx',
        'supp_fns': ['หนังสือยินยอมการคุ้มครองข้อมูลส่วนบุคคล_PDPA_Consent_Form.docx']
    },
    4: {
        'title': 'สุนทรียภาพแห่งการฟัง',
        'plans': [16, 17, 18],
        'eval_ids': [6, 7, 8],
        'test_fn': 'แบบทดสอบท้ายหน่วย_หน่วยที่_4_สุนทรียภาพแห่งการฟัง.docx',
        'supp_fns': ['เอกสารวิชาการ_กรอบแนวคิดการจัดการเรียนรู้_T_PACE_Framework.docx']
    }
}

def set_font_thai(run, name='TH Sarabun New'):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:cs'), name)
    rFonts.set(qn('w:eastAsia'), name)

def strip_appendix(doc):
    body = doc._body._element
    elements = list(body)
    app_idx = -1
    for i, el in enumerate(elements):
        if el.tag.endswith('p') and el.text and el.text.strip() == 'ภาคผนวก':
            app_idx = i
            break
    if app_idx != -1:
        to_remove = [el for el in elements[app_idx:] if not el.tag.endswith('sectPr')]
        for el in to_remove:
            body.remove(el)
        return True
    return False

def add_master_appendix_page(doc, title="ภาคผนวก", subtitle="", toc_items=None):
    doc.add_page_break()
    p_title = doc.add_paragraph()
    p_title.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run(title)
    r_t.bold = True
    r_t.font.size = docx.shared.Pt(26)
    r_t.font.color.rgb = docx.shared.RGBColor(0x1A, 0x36, 0x5D)
    set_font_thai(r_t)
    
    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        r_sub = p_sub.add_run(subtitle)
        r_sub.bold = True
        r_sub.font.size = docx.shared.Pt(16)
        r_sub.font.color.rgb = docx.shared.RGBColor(0x2B, 0x6C, 0xB0)
        set_font_thai(r_sub)
        
    p_sch = doc.add_paragraph()
    p_sch.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r_sch = p_sch.add_run("รายวิชา ดนตรี 1 (ศ21101) ชั้นมัธยมศึกษาปีที่ 1 ภาคเรียนที่ 2 โรงเรียนพะเยาพิทยาคม")
    r_sch.font.size = docx.shared.Pt(13)
    r_sch.font.color.rgb = docx.shared.RGBColor(0x71, 0x80, 0x96)
    set_font_thai(r_sch)
    
    p_sep = doc.add_paragraph()
    p_sep.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r_sep = p_sep.add_run("―" * 50)
    r_sep.font.color.rgb = docx.shared.RGBColor(0xCB, 0xD5, 0xE0)
    
    if toc_items:
        p_hdr = doc.add_paragraph()
        p_hdr.paragraph_format.left_indent = docx.shared.Inches(0.8)
        r_hdr = p_hdr.add_run("สารบัญเอกสารแนบท้าย (ภาคผนวก):")
        r_hdr.bold = True
        r_hdr.font.size = docx.shared.Pt(14)
        set_font_thai(r_hdr)
        
        for item in toc_items:
            p_it = doc.add_paragraph()
            p_it.paragraph_format.left_indent = docx.shared.Inches(1.1)
            r_it = p_it.add_run(item)
            r_it.font.size = docx.shared.Pt(13)
            set_font_thai(r_it)
            
    p_sep2 = doc.add_paragraph()
    p_sep2.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    r_sep2 = p_sep2.add_run("―" * 50)
    r_sep2.font.color.rgb = docx.shared.RGBColor(0xCB, 0xD5, 0xE0)

def prepend_document_banner(doc, letter, category, filename, detail_label=""):
    p_line = doc.paragraphs[0].insert_paragraph_before()
    r_line = p_line.add_run("―" * 55)
    r_line.font.color.rgb = docx.shared.RGBColor(0xCB, 0xD5, 0xE0)
    
    p_file = doc.paragraphs[0].insert_paragraph_before()
    r_file = p_file.add_run(f"เอกสารแนบ: {filename}")
    r_file.bold = True
    r_file.font.size = docx.shared.Pt(13)
    r_file.font.color.rgb = docx.shared.RGBColor(0x2B, 0x6C, 0xB0)
    set_font_thai(r_file)
    
    if detail_label:
        r_det = p_file.add_run(f" ({detail_label})")
        r_det.font.size = docx.shared.Pt(12)
        r_det.font.color.rgb = docx.shared.RGBColor(0x71, 0x80, 0x96)
        set_font_thai(r_det)

    p_head = doc.paragraphs[0].insert_paragraph_before()
    p_head.paragraph_format.page_break_before = True
    r_head = p_head.add_run(f"ภาคผนวก {letter}: {category}")
    r_head.bold = True
    r_head.font.size = docx.shared.Pt(18)
    r_head.font.color.rgb = docx.shared.RGBColor(0x1A, 0x36, 0x5D)
    set_font_thai(r_head)

def get_media_path(p_num):
    prefix = f"ผ{p_num:02d}"
    for f in os.listdir(media_dir):
        if prefix in f and f.endswith('.docx') and not f.startswith('~$'):
            return os.path.join(media_dir, f), f
    return None, None

def get_eval_paths(p_num):
    results = []
    tool_map = {
        1: [1, 3],
        2: [1, 3],
        3: [1, 3],
        4: [1, 3],
        5: [2, 3],
        6: [2, 3],
        7: [2, 3],
        8: [2, 3],
        9: [2, 3],
        10: [2, 4, 3],
        11: [2, 3],
        12: [5, 8],
        13: [5, 8],
        14: [5, 8],
        15: [5, 8],
        16: [6, 7, 8],
        17: [7, 6, 8],
        18: [6, 7, 8],
    }
    target_tools = tool_map.get(p_num, [])
    for t_id in target_tools:
        prefix = f"ฉบับที่_{t_id}_"
        for f in os.listdir(eval_dir):
            if prefix in f and f.endswith('.docx') and not f.startswith('~$'):
                results.append((os.path.join(eval_dir, f), f, t_id))
                break
    return results

def get_test_path(p_num):
    unit_test_map = {
        4: 'แบบทดสอบท้ายหน่วย_หน่วยที่_1_พื้นฐานคาลิมบ้า.docx',
        11: 'แบบทดสอบท้ายหน่วย_หน่วยที่_2_ปฏิบัติรวมวงคาลิมบ้า.docx',
        15: 'แบบทดสอบท้ายหน่วย_หน่วยที่_3_รอบรู้วงดนตรีและวัฒนธรรม.docx',
        18: 'แบบทดสอบท้ายหน่วย_หน่วยที่_4_สุนทรียภาพแห่งการฟัง.docx'
    }
    if p_num in unit_test_map:
        fn = unit_test_map[p_num]
        return os.path.join(test_dir, fn), fn
    return None, None

def get_supp_paths(p_num):
    supps = []
    if p_num == 1:
        fn = 'เอกสารวิชาการ_กรอบแนวคิดการจัดการเรียนรู้_T_PACE_Framework.docx'
        supps.append((os.path.join(supp_dir, fn), fn))
    if p_num in [10, 15]:
        fn = 'หนังสือยินยอมการคุ้มครองข้อมูลส่วนบุคคล_PDPA_Consent_Form.docx'
        supps.append((os.path.join(supp_dir, fn), fn))
    if p_num == 11:
        fn = 'วก7_รายงานการวิจัยในชั้นเรียนอย่างย่อ_ศ21101.docx'
        supps.append((os.path.join(supp_dir, fn), fn))
    return supps

def process_plan_file(target_path, p_num):
    if not os.path.exists(target_path):
        return
    
    u_num, u_title = get_unit_info(p_num)
    p_topic = plan_topics[p_num]
    
    master_doc = docx.Document(target_path)
    strip_appendix(master_doc)
    
    media_path, media_fn = get_media_path(p_num)
    eval_paths = get_eval_paths(p_num)
    test_path, test_fn = get_test_path(p_num)
    supp_paths = get_supp_paths(p_num)
    
    appendix_items = []
    if media_fn:
        appendix_items.append(f"• ภาคผนวก ก: สื่อการเรียนรู้และใบกิจกรรมประจำแผน ({media_fn})")
    if eval_paths:
        eval_names = ", ".join(fn for _, fn, _ in eval_paths)
        appendix_items.append(f"• ภาคผนวก ข: เครื่องมือวัดและประเมินผล ({eval_names})")
    if test_fn:
        appendix_items.append(f"• ภาคผนวก ค: แบบทดสอบวัดผลสัมฤทธิ์ทางการเรียนท้ายหน่วย ({test_fn})")
    if supp_paths:
        supp_names = ", ".join(fn for _, fn in supp_paths)
        appendix_items.append(f"• ภาคผนวก ง: เอกสารประกอบการจัดการเรียนรู้ ({supp_names})")
        
    add_master_appendix_page(
        master_doc,
        title="ภาคผนวก",
        subtitle=f"แผนการจัดการเรียนรู้ที่ {p_num}: {p_topic}",
        toc_items=appendix_items
    )
    
    composer = Composer(master_doc)
    
    # 1. ภาคผนวก ก: Media
    if media_path and os.path.exists(media_path):
        m_doc = docx.Document(media_path)
        prepend_document_banner(m_doc, "ก", "สื่อการเรียนรู้และใบกิจกรรมประจำแผน", media_fn)
        composer.append(m_doc)
        
    # 2. ภาคผนวก ข: Eval tools
    for idx, (ep, efn, tid) in enumerate(eval_paths):
        if os.path.exists(ep):
            e_doc = docx.Document(ep)
            letter = "ข" if idx == 0 else f"ข (ฉบับที่ {tid})"
            prepend_document_banner(e_doc, letter, "เครื่องมือวัดและประเมินผล", efn)
            composer.append(e_doc)
            
    # 3. ภาคผนวก ค: Unit test
    if test_path and os.path.exists(test_path):
        t_doc = docx.Document(test_path)
        prepend_document_banner(t_doc, "ค", "แบบทดสอบวัดผลสัมฤทธิ์ทางการเรียนท้ายหน่วย", test_fn)
        composer.append(t_doc)
        
    # 4. ภาคผนวก ง: Supp docs
    for idx, (sp, sfn) in enumerate(supp_paths):
        if os.path.exists(sp):
            s_doc = docx.Document(sp)
            letter = "ง" if idx == 0 else f"ง ({idx+1})"
            prepend_document_banner(s_doc, letter, "เอกสารประกอบการจัดการเรียนรู้", sfn)
            composer.append(s_doc)
            
    composer.save(target_path)
    print(f"✅ Reassembled Plan with Banners: {os.path.basename(target_path)}")

def process_unit_file(target_path, u_num):
    if not os.path.exists(target_path):
        return
        
    u_info = units_data_wk4[u_num]
    u_title = u_info['title']
    plan_list = u_info['plans']
    
    master_doc = docx.Document(target_path)
    strip_appendix(master_doc)
    
    # Media files for plans in unit
    unit_media_files = []
    for p in plan_list:
        mp, mfn = get_media_path(p)
        if mp and os.path.exists(mp):
            unit_media_files.append((mp, mfn, p))
            
    # Eval tools
    unit_eval_files = []
    for t_id in u_info['eval_ids']:
        prefix = f"ฉบับที่_{t_id}_"
        for f in sorted(os.listdir(eval_dir)):
            if prefix in f and f.endswith('.docx') and not f.startswith('~$'):
                unit_eval_files.append((os.path.join(eval_dir, f), f, t_id))
                break
                
    # Test & Blueprint
    unit_test_files = []
    t_path = os.path.join(test_dir, u_info['test_fn'])
    if os.path.exists(t_path):
        unit_test_files.append((t_path, u_info['test_fn'], 'แบบทดสอบประจำหน่วย 10 ข้อ'))
    bp_path = os.path.join(test_dir, 'พิมพ์เขียวข้อสอบ_Test_Blueprint_ศ21101_ดนตรี1.docx')
    if os.path.exists(bp_path):
        unit_test_files.append((bp_path, 'พิมพ์เขียวข้อสอบ_Test_Blueprint_ศ21101_ดนตรี1.docx', 'พิมพ์เขียวข้อสอบ Test Blueprint'))
        
    # Supp files
    unit_supp_files = []
    for sfn in u_info['supp_fns']:
        sp = os.path.join(supp_dir, sfn)
        if os.path.exists(sp):
            unit_supp_files.append((sp, sfn))
            
    # Build Table of Contents
    toc_items = []
    if unit_media_files:
        m_list = ", ".join(f"แผน {p}" for p in plan_list)
        toc_items.append(f"• ภาคผนวก ก: ชุดสื่อการจัดการเรียนรู้และใบกิจกรรมประจำหน่วย ({m_list} รวม {len(unit_media_files)} ชุด)")
    if unit_eval_files:
        e_list = ", ".join(f"ฉบับที่ {eid}" for eid in u_info['eval_ids'])
        toc_items.append(f"• ภาคผนวก ข: ชุดเครื่องมือวัดและประเมินผลประจำหน่วย ({e_list} รวม {len(unit_eval_files)} ฉบับ)")
    if unit_test_files:
        toc_items.append(f"• ภาคผนวก ค: แบบทดสอบวัดผลสัมฤทธิ์ทางการเรียนท้ายหน่วยที่ {u_num} และพิมพ์เขียวข้อสอบ (10 ข้อ)")
    if unit_supp_files:
        s_list = ", ".join(fn for _, fn in unit_supp_files)
        toc_items.append(f"• ภาคผนวก ง: เอกสารประกอบการจัดการเรียนรู้ ({s_list})")
        
    add_master_appendix_page(
        master_doc,
        title="ภาคผนวก",
        subtitle=f"หน่วยการเรียนรู้ที่ {u_num}: {u_title}",
        toc_items=toc_items
    )
    
    composer = Composer(master_doc)
    
    # 1. Append Media
    for idx, (mp, mfn, p_no) in enumerate(unit_media_files):
        m_doc = docx.Document(mp)
        letter = f"ก (แผนการเรียนรู้ที่ {p_no})"
        prepend_document_banner(m_doc, letter, f"สื่อการเรียนรู้และใบกิจกรรมประจำหน่วย — แผนที่ {p_no}", mfn)
        composer.append(m_doc)
        
    # 2. Append Eval tools
    for idx, (ep, efn, tid) in enumerate(unit_eval_files):
        e_doc = docx.Document(ep)
        letter = f"ข (ฉบับที่ {tid})"
        prepend_document_banner(e_doc, letter, "ชุดเครื่องมือวัดและประเมินผลประจำหน่วย", efn)
        composer.append(e_doc)
        
    # 3. Append Test & Blueprint
    for idx, (tp, tfn, desc) in enumerate(unit_test_files):
        t_doc = docx.Document(tp)
        letter = "ค" if idx == 0 else "ค (ต่อ)"
        prepend_document_banner(t_doc, letter, desc, tfn)
        composer.append(t_doc)
        
    # 4. Append Supp files
    for idx, (sp, sfn) in enumerate(unit_supp_files):
        s_doc = docx.Document(sp)
        letter = "ง" if idx == 0 else f"ง ({idx+1})"
        prepend_document_banner(s_doc, letter, "เอกสารประกอบการจัดการเรียนรู้", sfn)
        composer.append(s_doc)
        
    composer.save(target_path)
    print(f"✅ Reassembled Unit with Banners: {os.path.basename(target_path)}")

# --------------------------------------------------------------------------
# Main Execution
# --------------------------------------------------------------------------
if __name__ == '__main__':
    print("=================================================================")
    print("UPGRADING ALL APPENDICES WITH DEDICATED SECTION BANNERS & PAGE BREAKS")
    print("=================================================================")
    
    # 1. วก5 in เอกสารแผนการสอน/ม.1 ดนตรี1 ศ21101/วก5
    print("\n--- 1. Processing วก5 Files ---")
    wk5_dir = os.path.join(base, 'วก5')
    for p in range(1, 19):
        fn = f"วก5_แผนการจัดการเรียนรู้_แผนที่_{p:02d}.docx"
        process_plan_file(os.path.join(wk5_dir, fn), p)
        
    # 2. แผนการสอนพร้อมใช้
    print("\n--- 2. Processing แผนการสอนพร้อมใช้ Files ---")
    ready_dir = os.path.join(base, 'แผนการสอนพร้อมใช้')
    for p in range(1, 19):
        fn = f"แผนการจัดการเรียนรู้_แผนที่_{p:02d}.docx"
        process_plan_file(os.path.join(ready_dir, fn), p)
        
    # 3. คู่มือครู วก5
    print("\n--- 3. Processing คู่มือครู วก5 Files ---")
    teacher_wk5_dir = os.path.join(base, 'คู่มือการสอนสำหรับครู', 'คู่มือครู ม.1 ดนตรี1 ศ21101', 'วก5')
    for p in range(1, 19):
        fn = f"คู่มือครู_วก5_แผนการจัดการเรียนรู้_แผนที่_{p:02d}.docx"
        process_plan_file(os.path.join(teacher_wk5_dir, fn), p)
        
    # 4. วก4 in เอกสารแผนการสอน/ม.1 ดนตรี1 ศ21101/วก4
    print("\n--- 4. Processing วก4 Files ---")
    wk4_dir = os.path.join(base, 'วก4')
    for u in range(1, 5):
        pattern = f"วก4_หน่วยการเรียนรู้ที่_{u}_"
        matches = [f for f in os.listdir(wk4_dir) if f.startswith(pattern) and f.endswith('.docx')]
        if matches:
            process_unit_file(os.path.join(wk4_dir, matches[0]), u)
            
    # 5. คู่มือครู วก4
    print("\n--- 5. Processing คู่มือครู วก4 Files ---")
    teacher_wk4_dir = os.path.join(base, 'คู่มือการสอนสำหรับครู', 'คู่มือครู ม.1 ดนตรี1 ศ21101', 'วก4')
    for u in range(1, 5):
        pattern = f"คู่มือครู_วก4_หน่วยการเรียนรู้ที่_{u}_"
        matches = [f for f in os.listdir(teacher_wk4_dir) if f.startswith(pattern) and f.endswith('.docx')]
        if matches:
            process_unit_file(os.path.join(teacher_wk4_dir, matches[0]), u)
            
    print("\n=================================================================")
    print("ALL APPENDICES UPGRADED WITH BANNERS & PAGE BREAKS SUCCESSFULLY!")
    print("=================================================================")
