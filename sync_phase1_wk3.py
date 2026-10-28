"""
PHASE 1: อัปเกรด วก.3 ให้สะท้อน คศ.4 เต็มรูปแบบ
สิ่งที่เพิ่ม:
  - เพิ่มชื่อนวัตกรรม T-PACE + MusicSmart ในคำอธิบายรายวิชา
  - เพิ่มภาระงานรวบยอดใหม่ (Digital Portfolio, Mini-Recital, Podcast, Playlist) ในตาราง
  - เพิ่ม Backward Design ในหลักการออกแบบ
  - แก้สัดส่วน K-P-A ให้ตรงกับความเป็นจริง (วิชาปฏิบัติ P ต้องมากกว่า K)
  - เพิ่มสาระสำคัญระดับ คศ.4 ในตารางหน่วยการเรียนรู้
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from copy import deepcopy

ZWS = chr(0x200B)

def clean(t):
    for z in [chr(0x200B),chr(0x200C),chr(0x200D)]: t = t.replace(z,'')
    return t.strip()

def zws(t):
    return ZWS.join(list(t))

def set_font(run, bold=False, size=16):
    run.font.name = 'TH Sarabun New'
    run.font.size = Pt(size)
    run.bold = bold
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
    r.rPr.rFonts.set(qn('w:cs'), 'TH Sarabun New')

src = 'โครงสร้างรายวิชา_วก3_ดนตรีม1_สมบูรณ์.docx'
dst = 'วก3_คศ4_InventAndTransform_โครงสร้างรายวิชา.docx'

doc = Document(src)

# --- PATCH 1: อัปเกรดคำอธิบายรายวิชา (para 007 ที่พูดถึง T-PACE) ---
for para in doc.paragraphs:
    t = clean(para.text)
    if 'T-PACE Framework' in t and 'Backward Design' not in t:
        for run in para.runs:
            run.text = ''
        new_text = zws(
            'โดยใช้การออกแบบการจัดการเรียนรู้แบบอิงมาตรฐาน (Backward Design) ด้วยรูปแบบนวัตกรรม '
            'T-PACE Framework (Technology–Performance–Active–Creative–Evaluation) ควบคู่กับระบบ '
            'นิเวศการเรียนรู้ดิจิทัล (Digital Learning Ecosystem) ผ่านแอปพลิเคชัน MusicSmart '
            'ที่สนับสนุนการจูนเสียง (Tuner) การประเมินแบบ Real-time Peer Review '
            'และการเก็บข้อมูลเชิงประจักษ์เพื่อการวิจัยในชั้นเรียน (Classroom Action Research: CAR)'
        )
        if para.runs:
            para.runs[0].text = new_text
            set_font(para.runs[0])
        else:
            r = para.add_run(new_text)
            set_font(r)
        print('✅ PATCH 1: อัปเกรดคำอธิบายรายวิชา (Backward Design + MusicSmart + CAR)')
        break

# --- PATCH 2: อัปเกรดตารางโครงสร้างรายวิชา (TABLE 2) ให้มีภาระงานรวบยอด คศ.4 ---
unit_upgrades = {
    'พื้นฐานคาลิมบ้า': {
        'สาระสำคัญ': 'การสร้างความรู้ด้วยตนเอง (Constructivism) เรื่องสรีรวิทยาการจับเครื่องดนตรีและการก้องกังวาน (Resonance) พัฒนาทักษะผ่าน Backward Design',
        'ภาระงาน': 'Digital Portfolio (คลิปวิดีโอ + Self-Reflection), Personal Goal Setting บน MusicSmart Dashboard'
    },
    'ปฏิบัติรวมวงคาลิมบา': {
        'สาระสำคัญ': 'การสอดประสานเสียงในวงดนตรีรวมวง ผ่านกระบวนการ Peer Coaching และ Sectional Rehearsal นำสู่ Mini-Recital',
        'ภาระงาน': 'Mini-Recital (เพลงพื้นบ้านพะเยาประยุกต์) + Peer Critique บน MusicSmart App'
    },
    'รอบรู้วงดนตรีและวัฒนธรรม': {
        'สาระสำคัญ': 'การวิเคราะห์อิทธิพลวัฒนธรรมข้ามชาติต่อเสียงดนตรี และการสร้างสรรค์ Digital Audio Guide (Podcast) ที่บูรณาการทักษะปฏิบัติ',
        'ภาระงาน': 'Digital Audio Guide / Podcast (บันทึกเสียงผ่าน MusicSmart) + ร่างสคริปต์ข้ามวัฒนธรรม'
    },
    'สุนทรียภาพแห่งการฟัง': {
        'สาระสำคัญ': 'การถอดรหัสอารมณ์เพลงโดยใช้ Live Poll Analytics และการ Curate Digital Playlist เพื่อสะท้อนอัตลักษณ์ทางสุนทรียศาสตร์ส่วนบุคคล',
        'ภาระงาน': 'Digital Playlist Curation + บทวิจารณ์ดนตรีเชิงสุนทรียศาสตร์ เผยแพร่บน MusicSmart Community'
    },
}

for tbl in doc.tables:
    hdr = clean(tbl.rows[0].cells[0].text) if tbl.rows else ''
    # TABLE 2: โครงสร้างรายวิชา (ลำดับ|ชื่อหน่วย|ตัวชี้วัด|สาระสำคัญ|เวลา|คะแนน)
    if 'ลำดับ' in hdr or 'ชื่อหน่วย' in hdr:
        for row in tbl.rows[1:]:
            cells = row.cells
            if len(cells) < 4: continue
            unit_name = clean(cells[1].text)
            for key, upgrade in unit_upgrades.items():
                if key in unit_name:
                    # Column 3 = สาระสำคัญ
                    cells[3].paragraphs[0].clear()
                    run = cells[3].paragraphs[0].add_run(zws(upgrade['สาระสำคัญ']))
                    set_font(run, size=14)
                    # Add ภาระงาน if column exists (col 4-5 range)
                    # Add as new paragraph in สาระสำคัญ cell
                    p2 = cells[3].add_paragraph()
                    r2 = p2.add_run(zws(f"ภาระงานรวบยอด: {upgrade['ภาระงาน']}"))
                    set_font(r2, bold=True, size=13)
                    print(f'✅ PATCH 2: อัปเกรดตารางหน่วย [{key}]')
                    break

# --- PATCH 3: อัปเกรดตาราง K-P-A (TABLE 3) ให้ตรงกับธรรมชาติวิชาปฏิบัติ ---
# วิชาดนตรีปฏิบัติ P ควรมากกว่า K
# หน่วย 1-2 (ปฏิบัติ): K=5, P=10, A=5 รวม 20
# หน่วย 3-4 (ทฤษฎี+สุนทรีย): K=8, P=7, A=5 รวม 20
kpa_map = {
    'หน่วยที่ 1': ('5', '10', '5', '20'),
    'หน่วยที่ 2': ('5', '10', '5', '20'),
    'หน่วยที่ 3': ('8', '7',  '5', '20'),
    'หน่วยที่ 4': ('8', '7',  '5', '20'),
}
for tbl in doc.tables:
    hdr = clean(tbl.rows[0].cells[0].text) if tbl.rows else ''
    if 'หน่วยการเรียนรู้' in hdr and 'K' in clean(tbl.rows[1].cells[1].text if len(tbl.rows) > 1 else ''):
        for row in tbl.rows[2:]:  # skip 2 header rows
            cells = row.cells
            if not cells: continue
            label = clean(cells[0].text)
            if label in kpa_map:
                k, p, a, total = kpa_map[label]
                for ci, val in enumerate([None, k, p, a, total], start=0):
                    if ci == 0 or val is None: continue
                    cells[ci].paragraphs[0].clear()
                    run = cells[ci].paragraphs[0].add_run(zws(val))
                    set_font(run, size=14)
                print(f'✅ PATCH 3: อัปเกรด K-P-A ตาราง [{label}]')

# --- PATCH 4: เพิ่ม paragraph แนวปรัชญา Backward Design + CAR ก่อนจะปิดเอกสาร ---
# Find the last paragraph and add a new section
last_content_para = None
for para in doc.paragraphs:
    if clean(para.text): last_content_para = para

if last_content_para:
    from docx.oxml import OxmlElement
    def add_after(para, text, bold=False):
        new_p = deepcopy(para._element)
        for r in new_p.findall(qn('w:r')):
            new_p.remove(r)
        r_elem = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        for attr in ['w:ascii','w:eastAsia','w:cs']:
            rFonts.set(qn(attr), 'TH Sarabun New')
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '32')
        szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), '32')
        rPr.append(rFonts); rPr.append(sz); rPr.append(szCs)
        if bold:
            b = OxmlElement('w:b'); rPr.append(b)
        r_elem.append(rPr)
        t_elem = OxmlElement('w:t')
        t_elem.text = ZWS.join(list(text))
        t_elem.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r_elem.append(t_elem)
        new_p.append(r_elem)
        last_content_para._element.addnext(new_p)
        return new_p

    p1 = add_after(last_content_para,
        'หลักการออกแบบหลักสูตรและการจัดการเรียนรู้ (Curriculum Design Principles)', bold=True)
    p2 = add_after(last_content_para,
        '1. Backward Design: ออกแบบจากเป้าหมายและหลักฐานการเรียนรู้ก่อน ตามด้วยการวางแผนกิจกรรม')
    p3 = add_after(last_content_para,
        '2. T-PACE Framework: นวัตกรรมกรอบการสอนดนตรีที่บูรณาการ Technology–Performance–Active–Creative–Evaluation')
    p4 = add_after(last_content_para,
        '3. Digital Learning Ecosystem: MusicSmart App ทำหน้าที่เป็น Co-Pilot ให้ Feedback Real-time และเก็บ Data สำหรับ CAR')
    p5 = add_after(last_content_para,
        '4. Constructivism & SDL: นักเรียนสร้างองค์ความรู้ด้วยตนเอง (Construct Knowledge) และกำกับการเรียนรู้ส่วนบุคคล (Self-Directed Learning)')
    print('✅ PATCH 4: เพิ่มหลักการออกแบบหลักสูตร 4 ข้อ (Backward Design, T-PACE, Ecosystem, SDL)')

doc.save(dst)
print(f'\n✅ PHASE 1 COMPLETE: บันทึกเป็น [{dst}]')
print('  ไฟล์ต้นฉบับ [โครงสร้างรายวิชา_วก3_ดนตรีม1_สมบูรณ์.docx] ไม่ถูกเขียนทับ')
