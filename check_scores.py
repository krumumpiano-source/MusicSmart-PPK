import os
import docx

folder = r"d:\AI CURSER\PLANE MUSIC ED"
target_file = "โครงสร้างรายวิชา_วก3_ดนตรีม1_สมบูรณ์.docx"
path = os.path.join(folder, target_file)

if os.path.exists(path):
    doc = docx.Document(path)
    with open(os.path.join(folder, "scores.txt"), "w", encoding="utf-8") as out:
        out.write("--- PARAGRAPHS ---\n")
        for p in doc.paragraphs:
            text = p.text.strip().replace('\u200b', '')
            if text:
                out.write(text + "\n")
                
        out.write("\n--- TABLES ---\n")
        for i, table in enumerate(doc.tables):
            out.write(f"Table {i}:\n")
            for row in table.rows:
                row_data = [cell.text.strip().replace('\n', ' ').replace('\u200b', '') for cell in row.cells]
                out.write(" | ".join(row_data) + "\n")
            out.write("\n")
