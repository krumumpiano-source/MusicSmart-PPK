import sys
from docx import Document

def dump(path, out):
    out.write(f"\n\n===== FILE: {path} =====\n")
    try:
        doc = Document(path)
    except Exception as e:
        out.write(f"[ERROR opening: {e}]\n")
        return
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            out.write(t + "\n")
    for table in doc.tables:
        out.write("[TABLE]\n")
        for row in table.rows:
            cells = [c.text.strip().replace("\n"," ") for c in row.cells]
            out.write(" | ".join(cells) + "\n")

if __name__ == "__main__":
    out_path = sys.argv[1]
    files = sys.argv[2:]
    with open(out_path, "w", encoding="utf-8") as out:
        for f in files:
            dump(f, out)
    print("done", len(files))
