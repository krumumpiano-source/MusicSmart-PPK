# -*- coding: utf-8 -*-
import os, sys, io
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PLAN_POINTS = {
    1:(3,2), 2:(2,3), 3:(3,2), 4:(2,3),
    5:(1,1), 6:(1,1), 7:(1,1), 8:(1,1), 9:(1,1), 10:(1,1), 11:(4,4),
    12:(1,3), 13:(1,3), 14:(1,3), 15:(4,4),
    16:(1,5), 17:(2,4), 18:(4,4),
}
UNITS = {1:range(1,5), 2:range(5,12), 3:range(12,16), 4:range(16,19)}

def find_rubric_tables(doc):
    found = []
    def scan(tables):
        for t in tables:
            if len(t.rows) >= 3 and len(t.columns) >= 2:
                header_cell = t.rows[0].cells[1].text.strip()
                if header_cell == 'น้ำหนักคะแนน':
                    found.append(t)
            for row in t.rows:
                for cell in row.cells:
                    if cell.tables:
                        scan(cell.tables)
    scan(doc.tables)
    return found

def get_values(path):
    doc = Document(path)
    headline = None
    for p in doc.paragraphs:
        if 'Rubrics 2 มิติ' in p.text:
            headline = p.text
    if headline is None:
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if 'Rubrics 2 มิติ' in p.text:
                            headline = p.text
    p_val = ka_val = None
    for t in find_rubric_tables(doc):
        for row in t.rows[1:]:
            label = row.cells[0].text
            if '(P)' in label:
                p_val = row.cells[1].text.strip()
            elif '(K/A)' in label:
                ka_val = row.cells[1].text.strip()
    return headline, p_val, ka_val

def main():
    all_ok = True
    for unit, plans in UNITS.items():
        psum = kasum = 0
        for n in plans:
            path = os.path.join("วก5", f"วก5_แผนการจัดการเรียนรู้_แผนที่_{n:02d}.docx")
            headline, p_val, ka_val = get_values(path)
            exp_p, exp_ka = PLAN_POINTS[n]
            ok = (p_val == str(exp_p) and ka_val == str(exp_ka))
            if not ok:
                all_ok = False
            psum += int(p_val) if p_val else 0
            kasum += int(ka_val) if ka_val else 0
            print(f"plan{n:02d} unit{unit}: P={p_val} KA={ka_val} expect P={exp_p} KA={exp_ka} OK={ok}")
            print(f"   headline: {headline}")
        total = psum + kasum
        print(f"=== UNIT {unit} TOTAL: P={psum} KA={kasum} SUM={total} (expect 20) ===\n")
        if total != 20:
            all_ok = False

    print("ALL_OK =", all_ok)

if __name__ == '__main__':
    main()
