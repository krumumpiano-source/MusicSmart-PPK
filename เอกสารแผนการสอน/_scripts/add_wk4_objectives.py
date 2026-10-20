# -*- coding: utf-8 -*-
import sys, io
from docx import Document
from docx.shared import Emu

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

FONT_NAME = "TH Sarabun New"
FONT_SIZE = Emu(203200)

CONTENT = {
    "วก4/วก4_หน่วยการเรียนรู้ที่_1_พื้นฐานคาลิมบ้า.docx": {
        "enabling": [
            "1. บอกส่วนประกอบของคาลิมบ้าและจับเครื่องดนตรีได้ถูกสรีระตามหลักกายวิภาคของมือ",
            "2. อ่านและเขียนโน้ตไทย โน้ตสากลบรรทัด 5 เส้น และ Kalimba Tablature ได้ถูกต้อง",
            "3. ขับร้องและดีดบันไดเสียง C Major ตามจังหวะ Metronome ได้",
        ],
        "terminal": "บรรเลงเดี่ยวคาลิมบ้าประกอบการขับร้องบทเพลงสั้น บันทึกเป็นคลิปวิดีโอรวม 8 ห้องเพลงลงแฟ้มสะสมงานดิจิทัล (Digital Portfolio) พร้อมเขียนบันทึก Self-Reflection วิเคราะห์พัฒนาการของตนเองได้ (สอดคล้องตัวชี้วัด ศ 2.1 ม.1/1, ม.1/9)",
    },
    "วก4/วก4_หน่วยการเรียนรู้ที่_2_ปฏิบัติรวมวงคาลิมบ้า.docx": {
        "enabling": [
            "1. วางแผนและจัดสรรบทบาทเสียงประสานในการรวมวงคาลิมบ้าได้",
            "2. ออกแบบและประดิษฐ์เครื่องประกอบจังหวะจากวัสดุเหลือใช้ได้",
            "3. ควบคุมบันไดความเร็ว (Tempo) และความสมดุลของเสียง (Balance) ในการซ้อมกลุ่มย่อยได้",
        ],
        "terminal": "ร่วมแสดง Mini-Recital บทเพลงพื้นบ้านพะเยาประยุกต์รวมวงประกอบการขับร้องเป็นกลุ่ม พร้อมให้และรับข้อเสนอแนะแบบ Peer Critique ที่นำไปปฏิบัติได้จริงได้ (สอดคล้องตัวชี้วัด ศ 2.1 ม.1/3)",
    },
    "วก4/วก4_หน่วยการเรียนรู้ที่_3_รอบรู้วงดนตรีและวัฒนธรรม.docx": {
        "enabling": [
            "1. จำแนกและเปรียบเทียบเสียงเครื่องดนตรี/วงดนตรีไทยและสากลข้ามวัฒนธรรมได้",
            "2. วิเคราะห์บทบาทและอิทธิพลของดนตรีที่มีต่อสังคมไทยและวัฒนธรรมอื่นได้",
            "3. ร่างสคริปต์บรรยายและผลิตไฟล์เสียงประกอบดนตรีบรรเลงสดได้",
        ],
        "terminal": "ผลิตชิ้นงาน Digital Audio Guide / Podcast ความยาว 2-3 นาที บรรยายเปรียบเทียบวงดนตรีข้ามวัฒนธรรมพร้อมดนตรี BGM คาลิมบ้าบรรเลงสดประกอบ และเผยแพร่บน MusicSmart Community ได้ (สอดคล้องตัวชี้วัด ศ 2.1 ม.1/2, ม.1/4 และ ศ 2.2 ม.1/1, ม.1/2)",
    },
    "วก4/วก4_หน่วยการเรียนรู้ที่_4_สุนทรียภาพแห่งการฟัง.docx": {
        "enabling": [
            "1. วิเคราะห์และแสดงความคิดเห็นต่ออารมณ์ของบทเพลงที่มาจากวัฒนธรรมต่างกันได้",
            "2. ใช้เกณฑ์ประเมินคุณภาพงานดนตรีของตนเองและผู้อื่นได้อย่างเหมาะสม",
            "3. เขียนวิเคราะห์โครงสร้างบทวิจารณ์ดนตรีตามขั้นตอนของ Feldman ได้",
        ],
        "terminal": "เขียนบทวิจารณ์ดนตรีเชิงวิชาการตามหลัก Feldman และคัดสรรรังสรรค์ Digital Playlist บำบัดอารมณ์ พร้อมบันทึกสะท้อนพัฒนาการทักษะทางดนตรีของตนเองได้ (สอดคล้องตัวชี้วัด ศ 2.1 ม.1/5, ม.1/6, ม.1/7, ม.1/8)",
    },
}

def add_run(p, text, bold=False):
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT_NAME
    r.font.size = FONT_SIZE
    return r

def main():
    log = []
    for path, data in CONTENT.items():
        doc = Document(path)
        anchor = None
        for p in doc.paragraphs:
            if p.text.lstrip().startswith("3. สาระการเรียนรู้"):
                anchor = p
                break
        if anchor is None:
            log.append(f"ANCHOR NOT FOUND: {path}")
            continue

        new_paras_text = []

        p1 = anchor.insert_paragraph_before("")
        add_run(p1, "", bold=True)
        add_run(p1, "\n")
        add_run(p1, "2.3 จุดประสงค์การเรียนรู้ (Learning Objectives)", bold=True)
        new_paras_text.append(p1.text)

        p2 = anchor.insert_paragraph_before("")
        add_run(p2, "  จุดประสงค์นำทาง (Enabling Objectives):", bold=True)
        new_paras_text.append(p2.text)

        for line in data["enabling"]:
            p3 = anchor.insert_paragraph_before("")
            add_run(p3, f"    {line}", bold=False)
            new_paras_text.append(p3.text)

        p4 = anchor.insert_paragraph_before("")
        add_run(p4, "  จุดประสงค์ปลายทาง (Terminal Objective):", bold=True)
        new_paras_text.append(p4.text)

        p5 = anchor.insert_paragraph_before("")
        add_run(p5, f"    {data['terminal']}", bold=False)
        new_paras_text.append(p5.text)

        doc.save(path)
        log.append(f"OK: {path} inserted {len(new_paras_text)} paragraphs")

    for line in log:
        print(line)
    print("DONE")

if __name__ == '__main__':
    main()
