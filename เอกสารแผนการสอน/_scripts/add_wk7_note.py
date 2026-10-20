# -*- coding: utf-8 -*-
import sys, io
from docx import Document
from docx.shared import Emu

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

FONT_NAME = "TH Sarabun New"
FONT_SIZE = Emu(203200)

PATH = "เอกสารประกอบ/วก7_รายงานการวิจัยในชั้นเรียนอย่างย่อ_ศ21101.docx"

def add_run(p, text, bold=False, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = FONT_NAME
    r.font.size = FONT_SIZE
    return r

def main():
    doc = Document(PATH)
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("ชื่อเรื่อง:"):
            anchor = p
            break
    if anchor is None:
        print("ANCHOR NOT FOUND")
        return

    note = ("หมายเหตุ: เอกสารฉบับนี้เป็นรายงานการวิจัยในชั้นเรียนอย่างย่อที่จัดทำขึ้นในภาคเรียนที่ 2 "
            "ปีการศึกษา 2567 โดยมีวัตถุประสงค์เพื่อนำเสนอเป็นหลักฐานเชิงประจักษ์ (ผลการทดลองใช้นวัตกรรม "
            "T-PACE Framework และแอปพลิเคชัน MusicSmart) ประกอบแฟ้มสะสมผลงานขอรับการประเมินเพื่อเลื่อน "
            "วิทยฐานะครูเชี่ยวชาญ (คศ.4) จึงยังคงข้อมูลปีการศึกษาและภาคเรียนเดิมไว้ตามช่วงเวลาที่ดำเนินการวิจัยจริง "
            "แตกต่างจากเอกสารแผนการจัดการเรียนรู้ชุดปัจจุบัน (วก3-วก5 และเอกสารประกอบอื่น) ซึ่งจัดทำสำหรับการใช้สอน "
            "ในภาคเรียนที่ 2 ปีการศึกษา 2569")

    new_p = anchor.insert_paragraph_before("")
    add_run(new_p, note, bold=False, italic=True)

    sep = anchor.insert_paragraph_before("")
    add_run(sep, "———————————————————————————————————————————————————————", bold=False)

    doc.save(PATH)
    print("OK inserted note, len paragraphs now:", len(doc.paragraphs))

if __name__ == '__main__':
    main()
