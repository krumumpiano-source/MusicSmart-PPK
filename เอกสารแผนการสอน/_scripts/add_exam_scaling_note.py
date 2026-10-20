# -*- coding: utf-8 -*-
import sys, io
from docx import Document
from docx.shared import Emu

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

FONT_NAME = "TH Sarabun New"
FONT_SIZE = Emu(203200)

FILES = {
    "แบบทดสอบและพิมพ์เขียว/แบบทดสอบท้ายหน่วย_หน่วยที่_1_พื้นฐานคาลิมบ้า.docx": "กลางภาค",
    "แบบทดสอบและพิมพ์เขียว/แบบทดสอบท้ายหน่วย_หน่วยที่_2_ปฏิบัติรวมวงคาลิมบ้า.docx": "กลางภาค",
    "แบบทดสอบและพิมพ์เขียว/แบบทดสอบท้ายหน่วย_หน่วยที่_3_รอบรู้วงดนตรีและวัฒนธรรม.docx": "ปลายภาค",
    "แบบทดสอบและพิมพ์เขียว/แบบทดสอบท้ายหน่วย_หน่วยที่_4_สุนทรียภาพแห่งการฟัง.docx": "ปลายภาค",
}

def add_run(p, text, bold=False, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = FONT_NAME
    r.font.size = FONT_SIZE
    return r

def main():
    log = []
    for path, exam_name in FILES.items():
        doc = Document(path)
        anchor = None
        for p in doc.paragraphs:
            if p.text.strip().startswith("คำชี้แจง:"):
                anchor = p
                break
        if anchor is None:
            log.append(f"ANCHOR NOT FOUND: {path}")
            continue

        note = (f"หมายเหตุการแปลงคะแนน: คะแนนดิบเต็ม 10 คะแนนจากแบบทดสอบนี้ จะถูกแปลงสัดส่วน "
                f"(หารด้วย 2) เป็นคะแนนสอบ{exam_name} 5 คะแนน ตามโครงสร้างคะแนนในวก3 "
                f"(เช่น ตอบถูก 10 ข้อ = 5 คะแนนจริง, ตอบถูก 6 ข้อ = 3 คะแนนจริง)")

        # find insertion point: the paragraph right after anchor (the dashed separator line),
        # insert note before that separator so it sits between คำชี้แจง and the separator
        next_p = None
        found_anchor = False
        for p in doc.paragraphs:
            if p is anchor:
                found_anchor = True
                continue
            if found_anchor:
                next_p = p
                break

        target = next_p if next_p is not None else anchor
        new_p = target.insert_paragraph_before("")
        add_run(new_p, note, bold=False, italic=True)

        doc.save(path)
        log.append(f"OK: {path} -> {new_p.text[:50]}...")

    for line in log:
        print(line)
    print("DONE")

if __name__ == '__main__':
    main()
