# -*- coding: utf-8 -*-
import os, sys, io
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PLAN_POINTS = {
    1:(3,2), 2:(2,3), 3:(3,2), 4:(2,3),
    5:(1,1), 6:(1,1), 7:(1,1), 8:(1,1), 9:(1,1), 10:(1,1), 11:(4,4),
    12:(1,3), 13:(1,3), 14:(1,3), 15:(4,4),
    16:(1,5), 17:(2,4), 18:(4,4),
}

OLD_HEADLINE = "P 5 คะแนน, K/A 5 คะแนน"

def replace_in_paragraph(p, old, new):
    full = ''.join(r.text for r in p.runs)
    if old not in full:
        return False
    new_full = full.replace(old, new)
    if p.runs:
        p.runs[0].text = new_full
        for r in p.runs[1:]:
            r.text = ''
    return True

def walk_paragraphs(doc_or_cell):
    for p in doc_or_cell.paragraphs:
        yield p
    for t in doc_or_cell.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from walk_paragraphs(cell)

def set_cell_text(cell, new_text):
    paras = cell.paragraphs
    first_done = False
    for p in paras:
        if not p.runs:
            continue
        if not first_done:
            p.runs[0].text = new_text
            for r in p.runs[1:]:
                r.text = ''
            first_done = True
        else:
            for r in p.runs:
                r.text = ''
    if not first_done and paras:
        paras[0].add_run(new_text)

def find_rubric_tables(doc):
    found = []
    def scan(tables):
        for t in tables:
            if len(t.rows) >= 3 and len(t.columns) >= 2:
                header_cell = t.rows[0].cells[1].text.strip()
                if header_cell == 'น้ำหนักคะแนน':
                    found.append(t)
            for row in t.rows:
                for cell in row.cells:
                    if cell.tables:
                        scan(cell.tables)
    scan(doc.tables)
    return found

def fix_file(path, p_pts, ka_pts, do_headline):
    doc = Document(path)
    changed = []

    if do_headline:
        new_headline = f"P {p_pts} คะแนน, K/A {ka_pts} คะแนน"
        for p in walk_paragraphs(doc):
            if replace_in_paragraph(p, OLD_HEADLINE, new_headline):
                changed.append('headline')

    tables = find_rubric_tables(doc)
    for t in tables:
        for row in t.rows[1:]:
            label = row.cells[0].text
            if '(P)' in label:
                set_cell_text(row.cells[1], str(p_pts))
                changed.append('P_weight')
            elif '(K/A)' in label:
                set_cell_text(row.cells[1], str(ka_pts))
                changed.append('KA_weight')

    doc.save(path)
    return changed

def main():
    total_log = []
    for n in range(1, 19):
        p_pts, ka_pts = PLAN_POINTS[n]
        nn = f"{n:02d}"
        files = [
            (os.path.join("วก5", f"วก5_แผนการจัดการเรียนรู้_แผนที่_{nn}.docx"), True),
            (os.path.join("แผนการสอนพร้อมใช้", f"แผนการจัดการเรียนรู้_แผนที่_{nn}.docx"), True),
            (os.path.join("คู่มือการสอนสำหรับครู", "คู่มือครู ม.1 ดนตรี1 ศ21101", "วก5", f"คู่มือครู_วก5_แผนการจัดการเรียนรู้_แผนที่_{nn}.docx"), False),
        ]
        for path, do_headline in files:
            if not os.path.exists(path):
                total_log.append(f"MISSING: {path}")
                continue
            changed = fix_file(path, p_pts, ka_pts, do_headline)
            total_log.append(f"{path}: P={p_pts} KA={ka_pts} changed={changed}")

    for line in total_log:
        print(line)
    print("DONE")

if __name__ == '__main__':
    main()
