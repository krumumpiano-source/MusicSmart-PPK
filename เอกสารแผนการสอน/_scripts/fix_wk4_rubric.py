# -*- coding: utf-8 -*-
import os, sys, io
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PLAN_POINTS = {
    1:(3,2), 2:(2,3), 3:(3,2), 4:(2,3),
    5:(1,1), 6:(1,1), 7:(1,1), 8:(1,1), 9:(1,1), 10:(1,1), 11:(4,4),
    12:(1,3), 13:(1,3), 14:(1,3), 15:(4,4),
    16:(1,5), 17:(2,4), 18:(4,4),
}

UNITS = [
    ("วก4/วก4_หน่วยการเรียนรู้ที่_1_พื้นฐานคาลิมบ้า.docx", list(range(1,5))),
    ("วก4/วก4_หน่วยการเรียนรู้ที่_2_ปฏิบัติรวมวงคาลิมบ้า.docx", list(range(5,12))),
    ("วก4/วก4_หน่วยการเรียนรู้ที่_3_รอบรู้วงดนตรีและวัฒนธรรม.docx", list(range(12,16))),
    ("วก4/วก4_หน่วยการเรียนรู้ที่_4_สุนทรียภาพแห่งการฟัง.docx", list(range(16,19))),
]

OLD_PARA_SUB = "(P 5 คะแนน, K/A 5 คะแนน)"

def replace_in_paragraph(p, old, new):
    full = ''.join(r.text for r in p.runs)
    if old not in full:
        return False
    new_full = full.replace(old, new)
    if p.runs:
        p.runs[0].text = new_full
        for r in p.runs[1:]:
            r.text = ''
    return True

def set_cell_rubric(cell, p_pts, ka_pts):
    total = p_pts + ka_pts
    new_text = (f"รูบริก 2 มิติ ({total} คะแนน)\n"
                f"P ({p_pts}): มิติที่ 1: ทักษะกระบวนการ...\n"
                f"K/A ({ka_pts}): มิติที่ 2: ความรู้ การสะท...")
    lines = new_text.split('\n')
    para = cell.paragraphs[0]
    if not para.runs:
        para.add_run('')
    # map: run0=line0, run1='\n', run2=line1, run3='\n', run4=line2
    expected_len = 5
    if len(para.runs) >= expected_len:
        para.runs[0].text = lines[0]
        para.runs[1].text = '\n'
        para.runs[2].text = lines[1]
        para.runs[3].text = '\n'
        para.runs[4].text = lines[2]
        for r in para.runs[expected_len:]:
            r.text = ''
    else:
        para.runs[0].text = new_text
        for r in para.runs[1:]:
            r.text = ''

def main():
    log = []
    for path, plans in UNITS:
        if not os.path.exists(path):
            log.append(f"MISSING: {path}")
            continue
        doc = Document(path)
        psum = sum(PLAN_POINTS[n][0] for n in plans)
        kasum = sum(PLAN_POINTS[n][1] for n in plans)
        new_sub = f"(รวมทั้งหน่วย P {psum} คะแนน, K/A {kasum} คะแนน แบ่งน้ำหนักตามแต่ละแผน)"
        para_changed = False
        for p in doc.paragraphs:
            if replace_in_paragraph(p, OLD_PARA_SUB, new_sub):
                para_changed = True

        t = doc.tables[0]
        rows = t.rows[1:]
        if len(rows) != len(plans):
            log.append(f"ROW MISMATCH in {path}: rows={len(rows)} plans={len(plans)}")
            continue
        for row, n in zip(rows, plans):
            p_pts, ka_pts = PLAN_POINTS[n]
            set_cell_rubric(row.cells[6], p_pts, ka_pts)

        doc.save(path)
        log.append(f"{path}: para_changed={para_changed} P_total={psum} KA_total={kasum} plans={plans}")

    for line in log:
        print(line)
    print("DONE")

if __name__ == '__main__':
    main()
