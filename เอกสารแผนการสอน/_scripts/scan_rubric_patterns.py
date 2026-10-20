# -*- coding: utf-8 -*-
import glob, os, sys, io, re
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

TARGET_DIRS = [
    "วก5",
    "แผนการสอนพร้อมใช้",
    os.path.join("คู่มือการสอนสำหรับครู", "คู่มือครู ม.1 ดนตรี1 ศ21101", "วก5"),
]

def iter_texts(doc):
    for p in doc.paragraphs:
        if p.text.strip():
            yield ('PARA', p.text)
    def walk_tables(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip():
                            yield ('CELL', p.text)
                    if cell.tables:
                        yield from walk_tables(cell.tables)
    yield from walk_tables(doc.tables)

def main():
    for d in TARGET_DIRS:
        files = sorted(glob.glob(os.path.join(d, "*.docx")))
        for f in files:
            doc = Document(f)
            hits = []
            for kind, text in iter_texts(doc):
                if 'คะแนน' in text and ('รูบริก' in text or 'Rubric' in text or re.search(r'\bP\b', text) or 'K/A' in text or 'น้ำหนัก' in text):
                    hits.append((kind, text[:250]))
            print(f"=== {f} ===")
            for kind, text in hits:
                print(f"  [{kind}] {text}")
    print("DONE")

if __name__ == '__main__':
    main()
