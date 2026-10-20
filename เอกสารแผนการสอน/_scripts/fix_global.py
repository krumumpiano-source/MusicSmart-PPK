# -*- coding: utf-8 -*-
import glob
import os
import sys
import io
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

ZW_CHARS = ['​', '‌', '‍', '﻿', '⁠']

ROOT = "ม.1 ดนตรี1 ศ21101"
EXCLUDE = "แผนไฟล์เก่า"

OLD_TERM = "ภาคเรียนที่ 1"
NEW_TERM = "ภาคเรียนที่ 2"


def strip_zw(text):
    for ch in ZW_CHARS:
        if ch in text:
            text = text.replace(ch, '')
    return text


def clean_run(run):
    new_text = strip_zw(run.text)
    if new_text != run.text:
        run.text = new_text
        return True
    return False


def replace_in_paragraph(paragraph, old, new):
    changed = False
    full = ''.join(r.text for r in paragraph.runs)
    if old in full:
        new_full = full.replace(old, new)
        if paragraph.runs:
            paragraph.runs[0].text = new_full
            for r in paragraph.runs[1:]:
                r.text = ''
        changed = True
    return changed


def process_paragraphs(paragraphs, stats):
    for p in paragraphs:
        for r in p.runs:
            if clean_run(r):
                stats['zw'] += 1
        if replace_in_paragraph(p, OLD_TERM, NEW_TERM):
            stats['term'] += 1


def process_tables(tables, stats):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs, stats)
                if cell.tables:
                    process_tables(cell.tables, stats)


def process_file(path):
    stats = {'zw': 0, 'term': 0}
    try:
        doc = Document(path)
    except Exception as e:
        print(f"[ERROR opening {path}: {e}]")
        return stats
    process_paragraphs(doc.paragraphs, stats)
    process_tables(doc.tables, stats)
    # headers/footers
    for section in doc.sections:
        for hf in (section.header, section.footer):
            process_paragraphs(hf.paragraphs, stats)
            process_tables(hf.tables, stats)
    if stats['zw'] or stats['term']:
        doc.save(path)
    return stats


def main():
    files = []
    for f in glob.glob(os.path.join(ROOT, '**', '*.docx'), recursive=True):
        if EXCLUDE in f:
            continue
        files.append(f)
    files.sort()
    print(f"Total files: {len(files)}")
    total_zw = 0
    total_term = 0
    for f in files:
        stats = process_file(f)
        if stats['zw'] or stats['term']:
            print(f"{f}: zw_runs_cleaned={stats['zw']} term_paras_replaced={stats['term']}")
        total_zw += stats['zw']
        total_term += stats['term']
    print(f"DONE. total_zw_runs={total_zw} total_term_replacements={total_term}")


if __name__ == '__main__':
    main()
