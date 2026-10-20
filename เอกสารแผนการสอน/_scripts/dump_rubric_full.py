# -*- coding: utf-8 -*-
import os, sys, io
from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

FILES = [
    "วก5/วก5_แผนการจัดการเรียนรู้_แผนที่_01.docx",
    "วก5/วก5_แผนการจัดการเรียนรู้_แผนที่_05.docx",
    "วก5/วก5_แผนการจัดการเรียนรู้_แผนที่_12.docx",
    "วก5/วก5_แผนการจัดการเรียนรู้_แผนที่_16.docx",
    os.path.join("คู่มือการสอนสำหรับครู", "คู่มือครู ม.1 ดนตรี1 ศ21101", "วก5", "คู่มือครู_วก5_แผนการจัดการเรียนรู้_แผนที่_01.docx"),
]

def main():
    for f in FILES:
        print(f"\n===== {f} =====")
        doc = Document(f)
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                print(f"PARA[{i}]: {p.text}")
        for ti, t in enumerate(doc.tables):
            print(f"-- TABLE {ti} ({len(t.rows)}x{len(t.columns)}) --")
            for ri, row in enumerate(t.rows):
                cells = [c.text.replace(chr(10),' | ') for c in row.cells]
                print(f"  ROW{ri}: {cells}")
    print("DONE")

if __name__ == '__main__':
    main()
