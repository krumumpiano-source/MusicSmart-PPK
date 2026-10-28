import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

DOTS = "...................................................................................................................................................."

def reset_postlesson_in_doc(doc, filename):
    start_idx = None
    for i, p in enumerate(doc.paragraphs):
        if 'บันทึกผลหลังการสอน แผนการจัดการเรียนรู้' in p.text:
            start_idx = i
            break
            
    if start_idx is None:
        print(f"  [WARN] No postlesson header found in {filename}")
        return False
        
    # Check structure
    p1 = doc.paragraphs[start_idx + 1].text
    p3 = doc.paragraphs[start_idx + 3].text
    p5 = doc.paragraphs[start_idx + 5].text
    p9 = doc.paragraphs[start_idx + 9].text
    
    if not ('ผลการเรียนรู้ของนักเรียน' in p1 and 
            'ปัญหาที่เกิดขึ้นระหว่างการจัดการเรียนรู้' in p3 and 
            'แนวทางการแก้ปัญหาหรือปรับปรุงแผนครั้งต่อไป' in p5 and 
            'ความคิดเห็นของฝ่ายบริหาร' in p9):
        print(f"  [ERROR] Structure mismatch in {filename} at line {start_idx}")
        return False

    # Reset [start_idx + 2] (ผลการเรียนรู้ของนักเรียน)
    p2 = doc.paragraphs[start_idx + 2]
    if p2.runs:
        p2.runs[0].text = DOTS
        for r in p2.runs[1:]:
            r.text = ""
    else:
        p2.text = DOTS

    # Reset [start_idx + 4] (ปัญหาที่เกิดขึ้นระหว่างการจัดการเรียนรู้)
    p4 = doc.paragraphs[start_idx + 4]
    if p4.runs:
        p4.runs[0].text = DOTS
        for r in p4.runs[1:]:
            r.text = ""
    else:
        p4.text = DOTS

    # Reset [start_idx + 6] (แนวทางการแก้ปัญหาหรือปรับปรุงแผนครั้งต่อไป)
    p6 = doc.paragraphs[start_idx + 6]
    if p6.runs:
        p6.runs[0].text = DOTS
        for r in p6.runs[1:]:
            r.text = ""
    else:
        p6.text = DOTS

    # Update [start_idx + 7] (ครูผู้สอน signature)
    p7 = doc.paragraphs[start_idx + 7]
    teacher_sig = "ลงชื่อ ....................................................... ครูผู้สอน\n(นายพงศธร โพธิแก้ว)\nตำแหน่ง ครู"
    if p7.runs:
        p7.runs[0].text = teacher_sig
        for r in p7.runs[1:]:
            r.text = ""
    else:
        p7.text = teacher_sig

    # Update [start_idx + 8] (keep empty placeholder)
    p8 = doc.paragraphs[start_idx + 8]
    if p8.runs:
        p8.runs[0].text = ""
        for r in p8.runs[1:]:
            r.text = ""
    else:
        p8.text = ""

    # Update [start_idx + 9] (Header)
    p9_elem = doc.paragraphs[start_idx + 9]
    if 'ความคิดเห็นของฝ่ายบริหาร' in p9_elem.text:
        p9_elem.text = 'ความคิดเห็นและการรับรองของฝ่ายบริหาร'

    # Reset [start_idx + 10] (ความคิดเห็นของฝ่ายบริหาร)
    p10 = doc.paragraphs[start_idx + 10]
    if p10.runs:
        p10.runs[0].text = ""
        for r in p10.runs[1:]:
            r.text = ""
    else:
        p10.text = ""

    # Update [start_idx + 11] (ฝ่ายบริหาร 3-tier signature)
    p11 = doc.paragraphs[start_idx + 11]
    admin_sig = (
        "1. ความเห็นของหัวหน้ากลุ่มสาระการเรียนรู้ศิลปะ: .................................................................................\n"
        "ลงชื่อ ....................................................... หัวหน้ากลุ่มสาระการเรียนรู้ศิลปะ\n"
        "(นายอุดม ไชยวุฒิ)\n\n"
        "2. ความเห็นของรองผู้อำนวยการกลุ่มบริหารงานวิชาการ: ............................................................................\n"
        "ลงชื่อ ....................................................... รองผู้อำนวยการกลุ่มบริหารงานวิชาการ\n"
        "(นายกฤศ วงค์เรือง)\n\n"
        "3. การพิจารณาอนุมัติของผู้อำนวยการสถานศึกษา: ...................................................................................\n"
        "ลงชื่อ ....................................................... ผู้อำนวยการโรงเรียนพะเยาพิทยาคม\n"
        "(ว่าที่ร้อยตรีญาณบดินทร์ อินเตชะ)"
    )
    if p11.runs:
        p11.runs[0].text = admin_sig
        for r in p11.runs[1:]:
            r.text = ""
    else:
        p11.text = admin_sig

    # Also update header teacher name if needed
    for hp in doc.paragraphs[:15]:
        if 'สอนวันที่' in hp.text and 'ผู้สอน' in hp.text:
            if 'นายพงศธร โพธิแก้ว' not in hp.text:
                import re
                hp.text = re.sub(r'ผู้สอน\s+.*$', 'ผู้สอน นายพงศธร โพธิแก้ว', hp.text)

    return True


def update_semester_in_doc(doc):
    count = 0
    for p in doc.paragraphs:
        if 'ภาคเรียนที่ 1' in p.text:
            # Replace in runs
            replaced = False
            for r in p.runs:
                if 'ภาคเรียนที่ 1' in r.text:
                    r.text = r.text.replace('ภาคเรียนที่ 1', 'ภาคเรียนที่ 2')
                    replaced = True
            if not replaced:
                p.text = p.text.replace('ภาคเรียนที่ 1', 'ภาคเรียนที่ 2')
            count += 1
        elif 'ภาคเรียน 1' in p.text and 'ภาคเรียน 10' not in p.text:
            replaced = False
            for r in p.runs:
                if 'ภาคเรียน 1' in r.text:
                    r.text = r.text.replace('ภาคเรียน 1', 'ภาคเรียน 2')
                    replaced = True
            if not replaced:
                p.text = p.text.replace('ภาคเรียน 1', 'ภาคเรียน 2')
            count += 1

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if 'ภาคเรียนที่ 1' in p.text:
                        replaced = False
                        for r in p.runs:
                            if 'ภาคเรียนที่ 1' in r.text:
                                r.text = r.text.replace('ภาคเรียนที่ 1', 'ภาคเรียนที่ 2')
                                replaced = True
                        if not replaced:
                            p.text = p.text.replace('ภาคเรียนที่ 1', 'ภาคเรียนที่ 2')
                        count += 1
                    elif 'ภาคเรียน 1' in p.text and 'ภาคเรียน 10' not in p.text:
                        replaced = False
                        for r in p.runs:
                            if 'ภาคเรียน 1' in r.text:
                                r.text = r.text.replace('ภาคเรียน 1', 'ภาคเรียน 2')
                                replaced = True
                        if not replaced:
                            p.text = p.text.replace('ภาคเรียน 1', 'ภาคเรียน 2')
                        count += 1
    return count


def process_all():
    base_dir = r'เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101'
    
    # 1. Post-lesson target dirs (54 plans)
    postlesson_dirs = [
        os.path.join(base_dir, 'วก5'),
        os.path.join(base_dir, r'คู่มือการสอนสำหรับครู\คู่มือครู ม.1 ดนตรี1 ศ21101\วก5'),
        os.path.join(base_dir, 'แผนการสอนพร้อมใช้'),
    ]
    
    print("=== STEP 1: RESET POST-LESSON TO OFFICIAL BLANK FORM ===")
    total_plans_reset = 0
    for d in postlesson_dirs:
        print(f"Processing directory: {d}")
        files = sorted([f for f in os.listdir(d) if f.endswith('.docx')])
        for f in files:
            fp = os.path.join(d, f)
            doc = docx.Document(fp)
            if reset_postlesson_in_doc(doc, f):
                doc.save(fp)
                total_plans_reset += 1
                
    print(f"-> Successfully reset post-lesson notes in {total_plans_reset} plans!\n")

    print("=== STEP 2: UPDATE SEMESTER TO ภาคเรียนที่ 2 (2/2569) ===")
    all_target_dirs = [
        os.path.join(base_dir, 'วก3'),
        os.path.join(base_dir, 'วก4'),
        os.path.join(base_dir, 'วก5'),
        os.path.join(base_dir, r'คู่มือการสอนสำหรับครู\คู่มือครู ม.1 ดนตรี1 ศ21101\วก4'),
        os.path.join(base_dir, r'คู่มือการสอนสำหรับครู\คู่มือครู ม.1 ดนตรี1 ศ21101\วก5'),
        os.path.join(base_dir, 'แผนการสอนพร้อมใช้'),
        os.path.join(base_dir, 'แบบทดสอบและพิมพ์เขียว'),
    ]
    
    total_files_updated = 0
    total_replacements = 0
    for d in all_target_dirs:
        print(f"Updating semester in directory: {d}")
        files = sorted([f for f in os.listdir(d) if f.endswith('.docx')])
        for f in files:
            fp = os.path.join(d, f)
            doc = docx.Document(fp)
            c = update_semester_in_doc(doc)
            if c > 0:
                doc.save(fp)
                total_files_updated += 1
                total_replacements += c
                print(f"  {f}: {c} replacements to ภาคเรียนที่ 2")
                
    print(f"\n-> Successfully updated semester in {total_files_updated} files with {total_replacements} replacements!")

if __name__ == '__main__':
    process_all()
