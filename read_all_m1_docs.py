import os
import sys
import docx
import json

root_dir = r"D:\AI CURSER\PLANE MUSIC ED\เอกสารแผนการสอน\ม.1 ดนตรี1 ศ21101"

def extract_docx_summary(filepath):
    doc = docx.Document(filepath)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables_data = []
    for t in doc.tables:
        t_rows = []
        for row in t.rows:
            r_cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            t_rows.append(r_cells)
        if t_rows:
            tables_data.append(t_rows)
    return {
        "para_count": len(paragraphs),
        "table_count": len(tables_data),
        "first_paragraphs": paragraphs[:8],
        "tables_preview": [t[:3] for t in tables_data[:2]]
    }

summary = {}

# 1. วก3
wk3_path = os.path.join(root_dir, "วก3", "วก3_โครงสร้างรายวิชา_ศ21101_ดนตรี1.docx")
if os.path.exists(wk3_path):
    summary["วก3"] = extract_docx_summary(wk3_path)

# 2. วก4 sample
wk4_files = os.listdir(os.path.join(root_dir, "วก4"))
summary["วก4_files"] = wk4_files
if wk4_files:
    summary["วก4_sample"] = extract_docx_summary(os.path.join(root_dir, "วก4", wk4_files[0]))

# 3. วก5 sample
wk5_files = sorted(os.listdir(os.path.join(root_dir, "วก5")))
summary["วก5_files_count"] = len(wk5_files)
summary["วก5_files"] = wk5_files
if wk5_files:
    summary["วก5_sample"] = extract_docx_summary(os.path.join(root_dir, "วก5", wk5_files[0]))

# 4. สื่อการสอน
media_files = sorted(os.listdir(os.path.join(root_dir, "สื่อการสอน")))
summary["media_count"] = len(media_files)
summary["media_files"] = media_files
if media_files:
    summary["media_sample"] = extract_docx_summary(os.path.join(root_dir, "สื่อการสอน", media_files[0]))

# 5. เครื่องมือวัดผล
eval_files = sorted(os.listdir(os.path.join(root_dir, "เครื่องมือวัดผล")))
summary["eval_count"] = len(eval_files)
summary["eval_files"] = eval_files
if eval_files:
    summary["eval_sample"] = extract_docx_summary(os.path.join(root_dir, "เครื่องมือวัดผล", eval_files[0]))

# 6. เอกสารประกอบ
supp_files = sorted(os.listdir(os.path.join(root_dir, "เอกสารประกอบ")))
summary["supp_files"] = supp_files
for sf in supp_files:
    summary[f"supp_{sf}"] = extract_docx_summary(os.path.join(root_dir, "เอกสารประกอบ", sf))

# 7. แบบทดสอบและพิมพ์เขียว
test_files = sorted(os.listdir(os.path.join(root_dir, "แบบทดสอบและพิมพ์เขียว")))
summary["test_files"] = test_files
if test_files:
    summary["test_sample"] = extract_docx_summary(os.path.join(root_dir, "แบบทดสอบและพิมพ์เขียว", test_files[0]))

with open("scratch_summary_docs.json", "w", encoding="utf-8") as f:
    json.dump(summary, f, ensure_ascii=False, indent=2)

print("Extraction done successfully.")
