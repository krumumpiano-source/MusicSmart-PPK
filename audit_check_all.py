import docx, os, glob, sys

sys.stdout.reconfigure(encoding='utf-8')

def check_file(path, is_wk4=False):
    doc = docx.Document(path)
    base_name = os.path.basename(path)
    
    # 1. Find appendix divider
    app_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == 'ภาคผนวก':
            app_idx = i
            break
            
    if app_idx == -1:
        return {'file': base_name, 'status': 'NO_APPENDIX', 'details': 'No paragraph exactly equals ภาคผนวก'}
        
    app_paras = doc.paragraphs[app_idx:]
    
    # 2. Extract TOC items
    toc_items = []
    for p in app_paras:
        t = p.text.strip()
        if t.startswith('• ภาคผนวก') or t.startswith('ภาคผนวก '):
            toc_items.append(t)
            
    # 3. Check what was embedded
    total_app_paras = len(app_paras)
    total_app_tables = len(doc.tables)
    
    # check for specific sections
    has_media = any('สื่อการเรียนรู้ประกอบแผน' in p.text for p in app_paras)
    has_eval = any('เครื่องมือวัดและประเมินผล' in p.text for p in app_paras)
    has_test = any('แบบทดสอบวัดผลสัมฤทธิ์' in p.text for p in app_paras)
    has_blueprint = any('พิมพ์เขียว' in p.text for p in app_paras)
    has_supp = any(any(k in p.text for k in ['รายงานการวิจัยในชั้นเรียน', 'หนังสือยินยอม', 'เอกสารวิชาการ:']) for p in app_paras)
    
    # Check if TOC promises something that isn't embedded
    toc_text = ' '.join(toc_items)
    toc_has_test = 'ภาคผนวก ค' in toc_text
    toc_has_supp = 'ภาคผนวก ง' in toc_text
    
    mismatches = []
    if toc_has_test and not has_test:
        mismatches.append('TOC has ภาคผนวก ค (แบบทดสอบ) but content not found!')
    if toc_has_supp and not has_supp:
        mismatches.append('TOC has ภาคผนวก ง (เอกสารประกอบ) but content not found!')
    if not has_media:
        mismatches.append('ภาคผนวก ก (สื่อ/ใบกิจกรรม) content NOT found!')
    if not has_eval:
        mismatches.append('ภาคผนวก ข (เครื่องมือวัด) content NOT found!')
        
    return {
        'file': base_name,
        'toc_count': len(toc_items),
        'toc_items': toc_items,
        'paras_in_app': total_app_paras,
        'total_tables': total_app_tables,
        'has_media': has_media,
        'has_eval': has_eval,
        'has_test': has_test,
        'has_blueprint': has_blueprint,
        'has_supp': has_supp,
        'mismatches': mismatches
    }

print('=== AUDITING วก4 ===')
for f in sorted(glob.glob('เอกสารแผนการสอน/ม.1 ดนตรี1 ศ21101/วก4/*.docx')):
    res = check_file(f, is_wk4=True)
    status_icon = '✅' if not res['mismatches'] else '❌'
    print(f"{status_icon} {res['file']}: paras={res['paras_in_app']}, tables={res['total_tables']}, mismatches={res['mismatches']}")

print('\n=== AUDITING วก5 ===')
for f in sorted([f for f in glob.glob('เอกสารแผนการสอน/ม.1 ดนตรี1 ศ21101/วก5/*.docx') if not os.path.basename(f).startswith('~$')]):
    res = check_file(f, is_wk4=False)
    status_icon = '✅' if not res['mismatches'] else '❌'
    print(f"{status_icon} {res['file']}: paras={res['paras_in_app']}, tables={res['total_tables']}, mismatches={res['mismatches']}")
