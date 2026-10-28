import os
import docx

folder = r"d:\AI CURSER\PLANE MUSIC ED"

files_wk4 = [f for f in os.listdir(folder) if f.startswith("หน่วยการเรียนรู้_วก4_") and f.endswith(".docx")]
files_wk5 = [f for f in os.listdir(folder) if (f.startswith("แผนการสอน_วก5_") or f.startswith("วก5_")) and f.endswith(".docx")]

with open(os.path.join(folder, "extracted_data.txt"), "w", encoding="utf-8") as out:
    for f_list in [files_wk4, files_wk5]:
        for fname in sorted(f_list):
            out.write(f"\n\n{'='*40}\nFILE: {fname}\n{'='*40}\n")
            doc = docx.Document(os.path.join(folder, fname))
            
            out.write("--- PARAGRAPHS ---\n")
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    out.write(text + "\n")
                    
            out.write("\n--- TABLES ---\n")
            for i, table in enumerate(doc.tables):
                out.write(f"Table {i}:\n")
                for row in table.rows:
                    row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    out.write(" | ".join(row_data) + "\n")
                out.write("\n")
