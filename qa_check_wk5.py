import os
import glob
from docx import Document
from docx.oxml.ns import qn

def extract_full_text(docx_path):
    """Extract all text from a docx including tables, stripping ZWS."""
    doc = Document(docx_path)
    all_text = []
    for para in doc.paragraphs:
        cleaned = para.text.replace('\u200B', '').replace('\u200C', '').replace('\u200D', '')
        if cleaned.strip():
            all_text.append(cleaned)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cleaned = cell.text.replace('\u200B', '').replace('\u200C', '').replace('\u200D', '').strip()
                if cleaned:
                    row_text.append(cleaned)
            if row_text:
                all_text.append(" | ".join(row_text))
    return '\n'.join(all_text)

def check_plan(text, plan_num):
    issues = []
    upgrades = []

    # --- CHECK 1: Constructivism markers (Ind.3) ---
    construct_markers = [
        'สร้างความรู้', 'Construct', 'ทดลองด้วยตนเอง', 'ค้นพบ', 'สมมติฐาน',
        'อิสระ', 'ด้วยตนเอง', 'คิดค้น', 'วิเคราะห์ปัญหา', 'ทดลองจับ',
        'ร่างสคริปต์', 'ค้นหา', 'Chaos'
    ]
    construct_found = [m for m in construct_markers if m in text]
    if len(construct_found) < 2:
        issues.append(f"  ❌ [Ind.3] Constructivism ยังน้อยมาก (พบคำกุญแจ: {construct_found or 'ไม่พบ'})")
    else:
        upgrades.append(f"  ✅ [Ind.3] Constructivism: {', '.join(construct_found[:3])}")

    # --- CHECK 2: Self-Directed Learning markers (Ind.8) ---
    sdl_markers = [
        'Self-Directed', 'กำกับตนเอง', 'เป้าหมาย', 'Goal', 'Action Plan',
        'Micro-goal', 'Dashboard', 'Personal Goal', 'วางแผน', 'ตั้งเป้า'
    ]
    sdl_found = [m for m in sdl_markers if m in text]
    if len(sdl_found) < 2:
        issues.append(f"  ❌ [Ind.8] Self-Directed Learning ยังน้อยมาก (พบคำกุญแจ: {sdl_found or 'ไม่พบ'})")
    else:
        upgrades.append(f"  ✅ [Ind.8] Self-Directed: {', '.join(sdl_found[:3])}")

    # --- CHECK 3: Feedback Loop (Ind.6) ---
    fb_markers = [
        'Feedback', 'Data', 'Analytics', 'ผล', 'ปรับปรุง', 'ทันที',
        'Live Poll', 'Dashboard', 'กราฟ', 'บันทึกผล', 'Action Plan'
    ]
    fb_found = [m for m in fb_markers if m in text]
    if len(fb_found) < 2:
        issues.append(f"  ❌ [Ind.6] Feedback Loop ยังน้อย (พบคำกุญแจ: {fb_found or 'ไม่พบ'})")
    else:
        upgrades.append(f"  ✅ [Ind.6] Feedback Loop: {', '.join(fb_found[:3])}")

    # --- CHECK 4: Rubrics upgraded ---
    rubric_markers = [
        'ประยุกต์', 'วิเคราะห์', 'สังเคราะห์', 'ประเมินค่า', 'Peer Coaching',
        'อธิบาย', 'สะท้อนคิด', 'แนะนำ', 'ตรรกะ', 'ลึกซึ้ง', 'เป็นรูปธรรม'
    ]
    rubric_found = [m for m in rubric_markers if m in text]
    if not rubric_found:
        issues.append(f"  ❌ [Rubrics] ยังไม่มีคำกริยาขั้นสูงใน Rubrics")
    else:
        upgrades.append(f"  ✅ [Rubrics] คำกริยาขั้นสูง: {', '.join(rubric_found[:3])}")

    # --- CHECK 5: T-PACE Framework markers ---
    tpace = ['T -', 'P -', 'A -', 'C -', 'E -']
    tpace_found = [m for m in tpace if m in text]
    if len(tpace_found) < 3:
        issues.append(f"  ⚠️  [T-PACE] ครบแค่ {len(tpace_found)}/5 ขั้น ({tpace_found})")
    else:
        upgrades.append(f"  ✅ [T-PACE] ครบ {len(tpace_found)}/5 ขั้น")

    # --- CHECK 6: Section structure completeness ---
    required_sections = ['มาตรฐานการเรียนรู้', 'สาระสำคัญ', 'สาระการเรียนรู้',
                         'จุดประสงค์', 'ชิ้นงาน', 'กิจกรรม', 'สื่อ', 'การวัดและประเมิน',
                         'เกณฑ์การประเมิน', 'บันทึกผลหลัง']
    missing = [s for s in required_sections if s not in text]
    if missing:
        issues.append(f"  ❌ [โครงสร้าง] หายไป: {missing}")
    else:
        upgrades.append(f"  ✅ [โครงสร้าง] ครบทุก Section")

    # --- CHECK 7: Teacher-centered language still present ---
    teacher_centered = ['ครูอธิบาย', 'ครูบรรยาย', 'ครูสาธิต', 'ครูบอก', 'ครูแนะนำให้']
    tc_found = [m for m in teacher_centered if m in text]
    if tc_found:
        issues.append(f"  ⚠️  [Teacher-Centered] ยังมีร่องรอย: {tc_found}")

    return issues, upgrades

# --- MAIN QA RUN ---
folder = "."
new_files = sorted(glob.glob(os.path.join(folder, "วก5_คศ4_InventAndTransform_แผนที่_*.docx")))
orig_files = sorted(glob.glob(os.path.join(folder, "แผนการสอน_วก5_แผนที่_*.docx")))

print("=" * 70)
print("  QA REPORT: วก5_คศ4_InventAndTransform (18 แผน)")
print("=" * 70)

total_issues = 0
critical_plans = []

for i, fpath in enumerate(new_files):
    plan_num = i + 1
    fname = os.path.basename(fpath)
    text = extract_full_text(fpath)
    word_count = len(text.replace(' ', ''))
    issues, upgrades = check_plan(text, plan_num)
    total_issues += len(issues)

    status = "✅ PASS" if not issues else f"⚠️  ISSUES ({len(issues)})"
    print(f"\n{'─'*60}")
    print(f"  แผนที่ {plan_num:02d} | {status} | {word_count} ตัวอักษร")
    print(f"  ไฟล์: {fname}")
    for u in upgrades:
        print(u)
    for issue in issues:
        print(issue)
        if 'ไม่พบ' in issue or 'หายไป' in issue:
            if plan_num not in critical_plans:
                critical_plans.append(plan_num)

print(f"\n{'='*70}")
print(f"  สรุป: พบปัญหารวม {total_issues} จุด ใน {len(critical_plans)} แผนที่วิกฤต")
if critical_plans:
    print(f"  แผนที่ต้องแก้ด่วน: {critical_plans}")
else:
    print("  ทุกแผนผ่านการตรวจ QA")
print("=" * 70)

# --- Content Sample Check: แผนที่ 1 ---
if new_files:
    print("\n📋 ตัวอย่างเนื้อหากิจกรรมการเรียนรู้ แผนที่ 01 (50 บรรทัดแรก):")
    text1 = extract_full_text(new_files[0])
    lines = [l for l in text1.split('\n') if l.strip()]
    for ln in lines[:50]:
        print(f"  | {ln}")
