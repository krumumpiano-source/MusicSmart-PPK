from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import glob

ZWS = chr(0x200B)

def clean(t):
    for z in [chr(0x200B), chr(0x200C), chr(0x200D)]:
        t = t.replace(z, '')
    return t.strip()

def zws(t):
    return ZWS.join(list(t))

def set_font(run, size=15):
    run.font.name = 'TH Sarabun New'
    run.font.size = Pt(size)
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
    r.rPr.rFonts.set(qn('w:cs'), 'TH Sarabun New')

def replace_run_text(para, new_text, size=15):
    for run in para.runs:
        run.text = ''
    if para.runs:
        para.runs[0].text = zws(new_text)
        set_font(para.runs[0], size=size)
    else:
        r = para.add_run(zws(new_text))
        set_font(r, size=size)

sdl_eq = {
    '1': {
        'eq': '2.2 คำถามสำคัญ (Essential Question): เสียงคาลิมบ้าที่ดีเกิดขึ้นได้อย่างไร? นักเรียนสามารถสร้างความรู้เรื่องสรีรวิทยาดนตรีด้วยตนเองได้หรือไม่?',
        'sdl': 'การจัดการเรียนรู้เชิงรุก (Active Learning) ด้วยรูปแบบ T-PACE Framework ควบคู่กับ MusicSmart App โดยใช้หลัก Constructivism: นักเรียนทดลองค้นพบ (Construct Knowledge) ก่อนครูยืนยัน และตั้งเป้าหมายส่วนบุคคล (Self-Directed Goal Setting) ใน Digital Dashboard',
    },
    '2': {
        'eq': '2.2 คำถามสำคัญ (Essential Question): เสียงหลายชิ้นจะรวมเป็นหนึ่งเดียวได้อย่างไร? กระบวนการ Peer Critique ช่วยยกระดับ Mini-Recital ได้จริงหรือ?',
        'sdl': 'การจัดการเรียนรู้เชิงรุก (Active Learning) ด้วยรูปแบบ T-PACE Framework ควบคู่กับ MusicSmart App โดยใช้หลัก Constructivism: นักเรียนทดลองรวมวงอิสระก่อน (Chaos → Construct) เพื่อค้นพบความสำคัญของการสอดประสาน และใช้ Self-Directed Goal Setting ใน Sectional Rehearsal',
    },
    '3': {
        'eq': '2.2 คำถามสำคัญ (Essential Question): ดนตรีข้ามพรมแดนวัฒนธรรมได้อย่างไร? เราจะเล่าเรื่องราวของเสียงผ่าน Podcast ที่ผลิตด้วยตนเองได้หรือไม่?',
        'sdl': 'การจัดการเรียนรู้เชิงรุก (Active Learning) ด้วยรูปแบบ T-PACE Framework ควบคู่กับ MusicSmart App โดยใช้หลัก Constructivism: นักเรียนคิดค้นเกณฑ์จัดกลุ่มวงดนตรีด้วยตนเอง (Construct Knowledge) และกำหนดกลุ่มเป้าหมายผู้ฟัง Podcast ด้วยตนเอง (Self-Directed Learning)',
    },
    '4': {
        'eq': '2.2 คำถามสำคัญ (Essential Question): ดนตรีพูดกับความรู้สึกของเราอย่างไร? การจัด Digital Playlist คือการนิยามตัวตนของเราผ่านเสียงได้จริงหรือ?',
        'sdl': 'การจัดการเรียนรู้เชิงรุก (Active Learning) ด้วยรูปแบบ T-PACE Framework ควบคู่กับ MusicSmart App โดยใช้หลัก Constructivism: นักเรียนนำ Live Poll Data สร้างสมมติฐานทางจิตวิทยาดนตรีด้วยตนเอง (Construct Knowledge) และ Curate Playlist ตาม Self-Directed Learning',
    },
}

wk4_files = sorted(glob.glob('วก4_คศ4_InventAndTransform_*.docx'))
print('=== Patching WK4 SDL + Essential Question ===')
for i, path in enumerate(wk4_files):
    ukey = str(i + 1)
    info = sdl_eq.get(ukey, {})
    doc = Document(path)
    patched = []

    for para in doc.paragraphs:
        t = clean(para.text)

        # Inject Essential Question if missing
        if 'คำถามสำคัญ' in t and 'Essential Question' not in t and info.get('eq'):
            replace_run_text(para, info['eq'])
            patched.append('EQ')

        # Inject SDL/Constructivism if T-PACE line found but SDL missing
        if 'T-PACE Framework' in t and 'Self-Directed' not in t and info.get('sdl'):
            replace_run_text(para, info['sdl'])
            patched.append('SDL+Constructivism')

    doc.save(path)
    print(f'  Unit {ukey}: Patched {patched} -> {path[-35:]}')

print('\nWK4 SDL/EQ injection complete.')
print('Running final verification...')

# Quick re-verify
for path in wk4_files:
    doc = Document(path)
    full = ' '.join([clean(p.text) for p in doc.paragraphs])
    has_eq = 'Essential Question' in full
    has_sdl = 'Self-Directed' in full
    name = path.split('_')[-1].replace('.docx', '')
    status = 'PASS' if (has_eq and has_sdl) else 'FAIL'
    print(f'  [{status}] Unit {name}: EQ={has_eq}, SDL={has_sdl}')
